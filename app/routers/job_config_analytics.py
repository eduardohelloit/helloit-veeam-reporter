"""
job_config_analytics.py — Tela de Auditoria de Configuração de Rotinas.

GET  /analytics/job-config                  — página principal (filtros + grid de rotinas)
GET  /api/job-config/analytics/snapshots    — lista de snapshots filtrada (autocomplete + grid)
GET  /api/job-config/analytics/snapshot/{id} — detalhe completo de um snapshot (com objetos)
GET  /api/job-config/analytics/summary      — cards de conformidade por cliente
GET  /api/job-config/analytics/export-excel — exporta snapshots filtrados
"""
from __future__ import annotations

import io
import json
import re
from datetime import datetime
from typing import Optional

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import HTMLResponse, JSONResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from pathlib import Path
from sqlalchemy import text
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Client, JobConfigSnapshot, JobConfigObject, JobConfigImport, UserClient
from app.services import branding_service
from app.services.classification_service import situation_badge_class

router = APIRouter()

BASE_DIR  = Path(__file__).parent.parent
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))
templates.env.globals["situation_badge_class"] = situation_badge_class
templates.env.globals["now"] = datetime.now
templates.env.globals["get_branding"] = branding_service.get_branding
templates.env.filters["format_number"] = lambda n: f"{n:,}".replace(",", ".")


def _scoped_client_ids(db: Session, request: Request):
    """None = admin (todos os clientes); set = clientes permitidos."""
    if request.session.get("is_admin"):
        return None
    uid = request.session.get("user_id")
    if not uid:
        return set()
    return {r[0] for r in db.query(UserClient.client_id)
            .filter(UserClient.user_id == uid).all()}


def _parse_json(s):
    """json.loads tolerante — nunca estoura (dados antigos podem não ser JSON válido)."""
    if not s:
        return None
    try:
        v = json.loads(s)
        return v if v else None
    except (ValueError, TypeError):
        return [s]


def _schedule_display(snap) -> str:
    """Resumo da execução: 'A cada 2 horas' (periódico) ou o horário diário."""
    unit_map = {"Hours": "horas", "Minutes": "min", "Seconds": "seg"}
    if snap.sched_periodically_enabled and snap.sched_periodically_every:
        u = unit_map.get(snap.sched_periodically_unit, snap.sched_periodically_unit or "")
        return f"A cada {snap.sched_periodically_every} {u}".strip()
    if snap.sched_daily_time:
        # o Veeam grava um datetime .NET cru ("11/22/2023 19:25:00") em que só
        # a HORA importa — a data é o dia em que o agendamento foi definido
        m = re.search(r"(\d{1,2}:\d{2})", snap.sched_daily_time)
        return f"Diário às {m.group(1)}" if m else snap.sched_daily_time
    return "—"


def _retention_display(snap) -> str:
    """
    Retenção EFETIVA conforme a política:
      • storage_type 'Days'   -> N dias  (RetainDays)
      • storage_type 'Cycles' -> N pontos (RetainCycles)
    Fallback: se houver dias usa dias; senão pontos.
    """
    st = (snap.ret_storage_type or "").lower()
    if st == "days" and snap.ret_days:
        return f"{snap.ret_days} dias"
    if st == "cycles" and (snap.ret_cycles or snap.ret_restore_points):
        return f"{snap.ret_cycles or snap.ret_restore_points} pontos"
    if snap.ret_days:
        return f"{snap.ret_days} dias"
    if snap.ret_restore_points:
        return f"{snap.ret_restore_points} pontos"
    return "—"


def _yn(v):
    if v is None:
        return "—"
    return "Sim" if v else "Não"


def _gfs_display(snap) -> str:
    if not snap.ret_gfs_enabled:
        return "Não"
    parts = []
    if snap.ret_gfs_weekly:  parts.append(f"{snap.ret_gfs_weekly} sem.")
    if snap.ret_gfs_monthly: parts.append(f"{snap.ret_gfs_monthly} mensais")
    if snap.ret_gfs_yearly:  parts.append(f"{snap.ret_gfs_yearly} anuais")
    return "Sim" + (f" ({', '.join(parts)})" if parts else "")


_DIAS_PT = {
    "monday": "seg", "tuesday": "ter", "wednesday": "qua", "thursday": "qui",
    "friday": "sex", "saturday": "sáb", "sunday": "dom",
}
_ORDINAL_PT = {
    "first": "1º", "second": "2º", "third": "3º", "fourth": "4º", "last": "último",
}


def _dias_pt(days: Optional[str]) -> str:
    """'Saturday' -> 'sáb';  'Monday Saturday' -> 'seg, sáb'."""
    if not days:
        return ""
    out = [_DIAS_PT.get(d.lower(), d) for d in days.split() if d]
    return ", ".join(out)


def _full_quando(kind: Optional[str], days: Optional[str],
                 monthly: Optional[str]) -> str:
    """
    Só o "quando": 'sáb' / 'seg, sáb' / '1º seg do mês'.

    kind='Daily'   -> vale a lista de dias da semana
    kind='Monthly' -> vale a variante mensal ('First Monday')
    """
    if (kind or "").lower() == "monthly" and monthly:
        parts = monthly.split()
        ordinal = _ORDINAL_PT.get(parts[0].lower(), parts[0]) if parts else ""
        dia = _dias_pt(parts[1]) if len(parts) > 1 else ""
        return f"{ordinal} {dia} do mês".strip()
    return _dias_pt(days)


def _full_display(enabled, kind: Optional[str], days: Optional[str],
                  monthly: Optional[str]) -> str:
    """QUANDO um tipo de full roda: 'Sim (sáb)' / 'Sim (1º seg do mês)' / 'Não'."""
    if enabled is None:
        return "—"
    if not enabled:
        return "Não"
    quando = _full_quando(kind, days, monthly)
    return f"Sim ({quando})" if quando else "Sim"


def _full_when(snap) -> str:
    """
    Resumo curto de quando a rotina GERA full, para a coluna da listagem.
    Compact full não entra: é desfragmentação do arquivo, não um full novo.
    """
    partes = []
    for rotulo, enabled, kind, days, monthly in (
        ("Sintético", snap.synth_full_enabled,  snap.synth_full_kind,
         snap.synth_full_days,  snap.synth_full_monthly),
        ("Ativo",     snap.active_full_enabled, snap.active_full_kind,
         snap.active_full_days, snap.active_full_monthly),
    ):
        if not enabled:
            continue
        quando = _full_quando(kind, days, monthly)
        partes.append(f"{rotulo}: {quando}" if quando else rotulo)
    if not partes:
        # None em ambos = rotina coletada sem esse campo
        if snap.synth_full_enabled is None and snap.active_full_enabled is None:
            return "—"
        return "Nenhum"
    return " + ".join(partes)


def _full_day_match(snap, day: str) -> bool:
    """
    A rotina gera full (sintético OU ativo) no dia da semana pedido?

    `day` em inglês minúsculo ('saturday'). Os dias ficam gravados como nomes
    .NET separados por espaço ('Monday Saturday'); no modo Monthly o dia vem
    da variante mensal ('First Monday') — que também conta, porque o full
    realmente roda numa segunda.
    """
    day = (day or "").lower()
    if not day:
        return True
    for enabled, kind, days, monthly in (
        (snap.synth_full_enabled,  snap.synth_full_kind,  snap.synth_full_days,  snap.synth_full_monthly),
        (snap.active_full_enabled, snap.active_full_kind, snap.active_full_days, snap.active_full_monthly),
    ):
        if not enabled:
            continue
        alvo = monthly if (kind or "").lower() == "monthly" else days
        if alvo and day in alvo.lower().split():
            return True
    return False


def _proxies_display(snap) -> str:
    names = _parse_json(snap.proxy_selected_names)
    if isinstance(names, list) and names:
        return ", ".join(sorted(names))
    return snap.proxy_mode or "—"


def _audited_values(snap) -> list:
    """
    Conjunto curado de campos auditáveis (label, valor de exibição) usado para
    detectar alterações de configuração entre coletas. Ordem = ordem de exibição.
    """
    return [
        ("Habilitada",          _yn(snap.is_enabled)),
        ("Execução",            _schedule_display(snap)),
        ("Dias da semana",      snap.sched_days_of_week or "—"),
        ("Retry",               str(snap.sched_retry_count) if snap.sched_retry_count is not None else "—"),
        ("Application-Aware",   _yn(snap.aap_enabled)),
        ("App-Aware (modo)",    snap.aap_mode or "—"),
        ("VMTools Quiesce",     _yn(snap.vmtools_quiesce)),
        ("SQL Log (padrão job)", _yn(snap.sql_log_backup_enabled)),
        ("Repositório",         snap.repo_sobr_name or snap.repo_name or "—"),
        ("Imutável (perf.)",    _yn(snap.repo_immutability)),
        ("Linux Hardened",      _yn(snap.repo_is_linux_hardened)),
        ("Capacity tier",       snap.repo_capacity_name or "—"),
        ("Capacity imutável",   _yn(snap.repo_capacity_immutable)),
        ("Proxy (modo)",        snap.proxy_mode or "—"),
        ("Proxies",             _proxies_display(snap)),
        ("GIP (modo)",          snap.gip_mode or "—"),
        ("Retenção",            _retention_display(snap)),
        ("GFS",                 _gfs_display(snap)),
        ("Full sintético",      _full_display(snap.synth_full_enabled, snap.synth_full_kind,
                                              snap.synth_full_days, snap.synth_full_monthly)),
        ("Full ativo",          _full_display(snap.active_full_enabled, snap.active_full_kind,
                                              snap.active_full_days, snap.active_full_monthly)),
        ("Compact full",        _full_display(snap.compact_full_enabled, snap.compact_full_kind,
                                              snap.compact_full_days, snap.compact_full_monthly)),
        ("Compressão",          str(snap.stg_compression) if snap.stg_compression is not None else "—"),
        ("Dedup",               _yn(snap.stg_dedup_enabled)),
        ("Criptografia",        _yn(snap.stg_encryption)),
    ]


# ── Helpers ────────────────────────────────────────────────────────────────────

def _latest_snapshot_ids(db: Session, allowed,
                         client_id_filter: Optional[int] = None,
                         search: Optional[str] = None,
                         job_type: Optional[str] = None) -> list[int]:
    """
    Retorna IDs do snapshot mais recente por rotina.
    allowed=None → admin (sem filtro de cliente); allowed=set → IDs permitidos.
    client_id_filter restringe adicionalmente a um único cliente.
    """
    base = """
        SELECT DISTINCT ON (client_id, COALESCE(job_id, job_name))
               id
        FROM job_config_snapshots
        WHERE 1=1
    """
    params: dict = {}
    if client_id_filter:
        base += " AND client_id = :cid"
        params["cid"] = client_id_filter
    elif allowed is not None:
        if not allowed:
            return []
        base += " AND client_id = ANY(:cids)"
        params["cids"] = list(allowed)
    if search:
        base += " AND LOWER(job_name) LIKE :search"
        params["search"] = f"%{search.lower()}%"
    if job_type:
        base += " AND job_type = :jtype"
        params["jtype"] = job_type
    base += " ORDER BY client_id, COALESCE(job_id, job_name), collected_at DESC NULLS LAST"
    return list(db.execute(text(base), params).scalars().all())


# ── Rotas ──────────────────────────────────────────────────────────────────────

@router.get("/analytics/job-config", response_class=HTMLResponse)
def analytics_page(request: Request, db: Session = Depends(get_db)):
    allowed = _scoped_client_ids(db, request)
    q = db.query(Client).filter(Client.is_active.is_(True))
    if allowed is not None:
        q = q.filter(Client.id.in_(list(allowed)))
    clients   = q.order_by(Client.name).all()
    job_types = ["Backup", "NasBackup"]
    return templates.TemplateResponse(
        "analytics_job_config.html",
        {"request": request, "clients": clients, "job_types": job_types},
    )


@router.get("/api/job-config/analytics/summary")
def api_summary(
    request: Request,
    client_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    allowed = _scoped_client_ids(db, request)
    if allowed is not None and not allowed:
        return JSONResponse({"total": 0, "aap_ok": 0, "sql_log_enabled": 0,
                             "gip_manual": 0, "sobr_immutable": 0, "gfs_enabled": 0,
                             "disabled": 0})

    cid_filter = client_id if (allowed is None or client_id in allowed) else None
    snap_ids = _latest_snapshot_ids(db, allowed, client_id_filter=cid_filter)
    if not snap_ids:
        return JSONResponse({"total": 0, "aap_ok": 0, "sql_log_enabled": 0,
                             "gip_manual": 0, "sobr_immutable": 0, "gfs_enabled": 0,
                             "disabled": 0})

    snaps = db.query(JobConfigSnapshot).filter(JobConfigSnapshot.id.in_(snap_ids)).all()

    total          = len(snaps)
    aap_ok         = sum(1 for s in snaps if s.aap_enabled)
    gip_manual     = sum(1 for s in snaps if s.gip_mode == "Selected")
    sobr_immutable = sum(1 for s in snaps if s.repo_is_sobr and s.repo_immutability)
    gfs_enabled    = sum(1 for s in snaps if s.ret_gfs_enabled)
    disabled       = sum(1 for s in snaps if s.is_enabled is False)

    # SQL log: ao menos 1 objeto com log backup habilitado
    snap_ids_set = {s.id for s in snaps}
    sql_objs = (db.query(JobConfigObject.job_config_snapshot_id)
                .filter(JobConfigObject.job_config_snapshot_id.in_(list(snap_ids_set)),
                        JobConfigObject.sql_log_backup_enabled.is_(True))
                .distinct().all())
    sql_log_enabled = len(sql_objs)

    return JSONResponse({
        "total": total,
        "aap_ok": aap_ok,
        "sql_log_enabled": sql_log_enabled,
        "gip_manual": gip_manual,
        "sobr_immutable": sobr_immutable,
        "gfs_enabled": gfs_enabled,
        "disabled": disabled,
    })


@router.get("/api/job-config/analytics/snapshots")
def api_snapshots(
    request: Request,
    client_id: Optional[int] = Query(None),
    search: Optional[str]    = Query(None),
    job_type: Optional[str]  = Query(None),
    gip_filter: Optional[str]= Query(None),
    sql_filter: Optional[str]= Query(None),
    aap_filter: Optional[str]= Query(None),
    immutable_filter: Optional[str] = Query(None),
    disabled_filter: Optional[str]  = Query(None),
    full_day: Optional[str]  = Query(None),   # monday..sunday
    limit: int = Query(200, ge=1, le=1000),
    db: Session = Depends(get_db),
):
    allowed = _scoped_client_ids(db, request)
    if allowed is not None and not allowed:
        return JSONResponse([])

    cid_filter = client_id if (allowed is None or client_id in allowed) else None
    snap_ids = _latest_snapshot_ids(db, allowed, client_id_filter=cid_filter,
                                    search=search, job_type=job_type)
    if not snap_ids:
        return JSONResponse([])

    snaps = (db.query(JobConfigSnapshot)
             .filter(JobConfigSnapshot.id.in_(snap_ids))
             .order_by(JobConfigSnapshot.job_name)
             .all())

    # Contar objetos com SQL log por snapshot
    sql_counts: dict[int, int] = {}
    obj_rows = (db.query(JobConfigObject.job_config_snapshot_id)
                .filter(JobConfigObject.job_config_snapshot_id.in_(snap_ids),
                        JobConfigObject.sql_log_backup_enabled.is_(True))
                .all())
    for (sid,) in obj_rows:
        sql_counts[sid] = sql_counts.get(sid, 0) + 1

    result = []
    for s in snaps:
        # Filtros opcionais
        if gip_filter == "selected" and s.gip_mode != "Selected":
            continue
        if gip_filter == "automatic" and s.gip_mode != "Automatic":
            continue
        sql_vms = sql_counts.get(s.id, 0)
        if sql_filter == "enabled" and sql_vms == 0:
            continue
        if sql_filter == "disabled" and sql_vms > 0:
            continue
        if aap_filter == "ok" and not s.aap_enabled:
            continue
        if aap_filter == "fail" and s.aap_enabled:
            continue
        if immutable_filter == "yes" and not (s.repo_is_sobr and s.repo_immutability):
            continue
        if immutable_filter == "no" and (s.repo_is_sobr and s.repo_immutability):
            continue
        if disabled_filter == "yes" and s.is_enabled is not False:
            continue
        if full_day and not _full_day_match(s, full_day):
            continue

        result.append({
            "id":           s.id,
            "job_name":     s.job_name,
            "job_type":     s.job_type,
            "is_enabled":   s.is_enabled,
            "collected_at": s.collected_at.strftime("%d/%m/%Y %H:%M") if s.collected_at else None,
            "repo":         s.repo_sobr_name or s.repo_name,
            "is_sobr":      s.repo_is_sobr,
            "immutable":    bool(s.repo_is_sobr and s.repo_immutability),
            "gip_mode":     s.gip_mode,
            "aap_enabled":  s.aap_enabled,
            "sql_vms":      sql_vms,
            "retention":    _retention_display(s),
            "gfs":          s.ret_gfs_enabled,
            "sched_time":   _schedule_display(s),
            "sched_retry":  s.sched_retry_count,
            "proxy_mode":   s.proxy_mode,
            "full_when":    _full_when(s),
        })
        if len(result) >= limit:
            break

    return JSONResponse(result)


@router.get("/api/job-config/analytics/snapshot/{snapshot_id}")
def api_snapshot_detail(
    snapshot_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    allowed = _scoped_client_ids(db, request)
    snap = db.query(JobConfigSnapshot).get(snapshot_id)
    if not snap or (allowed is not None and snap.client_id not in allowed):
        return JSONResponse({"error": "não encontrado"}, status_code=404)

    objs = (db.query(JobConfigObject)
            .filter(JobConfigObject.job_config_snapshot_id == snapshot_id)
            .order_by(JobConfigObject.object_name)
            .all())

    def _bool(v):
        if v is None: return None
        return bool(v)

    return JSONResponse({
        "snapshot": {
            "id": snap.id,
            "job_id": snap.job_id,
            "job_name": snap.job_name,
            "job_type": snap.job_type,
            "platform": snap.platform,
            "backup_type": snap.backup_type,
            "is_enabled": _bool(snap.is_enabled),
            "is_schedule_enabled": _bool(snap.is_schedule_enabled),
            "next_run": snap.next_run,
            "collected_at": snap.collected_at.isoformat() if snap.collected_at else None,
            "vbr_server": snap.vbr_server,
            "schedule": {
                "daily_enabled": _bool(snap.sched_daily_enabled),
                "daily_time": snap.sched_daily_time,
                "days_of_week": snap.sched_days_of_week,
                "periodically_enabled": _bool(snap.sched_periodically_enabled),
                "periodically_every": snap.sched_periodically_every,
                "periodically_unit": snap.sched_periodically_unit,
                "retry_enabled": _bool(snap.sched_retry_enabled),
                "retry_count": snap.sched_retry_count,
                "chain_job": snap.sched_chain_job_name,
            },
            "repository": {
                "name": snap.repo_name,
                "type": snap.repo_type,
                "is_sobr": _bool(snap.repo_is_sobr),
                "sobr_name": snap.repo_sobr_name,
                "immutability": _bool(snap.repo_immutability),
                "linux_hardened": _bool(snap.repo_is_linux_hardened),
                "per_vm_files": _bool(snap.repo_per_vm_files),
                "extent_count": snap.repo_extent_count,
                "extent_names": _parse_json(snap.repo_extent_names),
                "capacity_name": snap.repo_capacity_name,
                "capacity_immutable": _bool(snap.repo_capacity_immutable),
                "capacity_immut_days": snap.repo_capacity_immut_days,
            },
            "backup_proxy": {
                "mode": snap.proxy_mode,
                "automatic": _bool(snap.proxy_automatic),
                "selected_names": _parse_json(snap.proxy_selected_names),
                "transport_modes": _parse_json(snap.proxy_transport_modes),
            },
            "gip": {
                "mode": snap.gip_mode,
                "automatic": _bool(snap.gip_automatic),
            },
            "application_aware": {
                "enabled": _bool(snap.aap_enabled),
                "mode": snap.aap_mode,
                "require_success": _bool(snap.aap_require_success),
                "vmtools_quiesce": _bool(snap.vmtools_quiesce),
            },
            "sql_default": {
                "detected": _bool(snap.sql_detected),
                "processing_enabled": _bool(snap.sql_processing_enabled),
                "tx_mode": snap.sql_tx_mode,
                "log_backup_enabled": _bool(snap.sql_log_backup_enabled),
                "log_freq_min": snap.sql_log_freq_min,
                "log_retention_days": snap.sql_log_retention_days,
                "per_object_detected": _bool(snap.sql_per_object_detected),
            },
            "retention": {
                "type": snap.ret_type,
                "storage_type": snap.ret_storage_type,
                "effective": _retention_display(snap),
                "restore_points": snap.ret_restore_points,
                "cycles": snap.ret_cycles,
                "days": snap.ret_days,
                "gfs_enabled": _bool(snap.ret_gfs_enabled),
                "gfs_weekly": snap.ret_gfs_weekly,
                "gfs_monthly": snap.ret_gfs_monthly,
                "gfs_yearly": snap.ret_gfs_yearly,
            },
            "storage": {
                "compression": snap.stg_compression,
                "block_size": snap.stg_block_size,
                "dedup": _bool(snap.stg_dedup_enabled),
                "encryption": _bool(snap.stg_encryption),
                "cbt": _bool(snap.stg_cbt_enabled),
            },
            # QUANDO cada tipo de full roda (kind/days só vêm preenchidos quando
            # o tipo está habilitado — ver job_config_service._full).
            "full": {
                "algorithm": snap.full_algorithm,
                "rollbacks": _bool(snap.transform_to_rollbacks),
                "synthetic": {
                    "enabled": _bool(snap.synth_full_enabled),
                    "kind": snap.synth_full_kind,
                    "days": snap.synth_full_days,
                    "monthly": snap.synth_full_monthly,
                    "display": _full_display(snap.synth_full_enabled, snap.synth_full_kind,
                                             snap.synth_full_days, snap.synth_full_monthly),
                },
                "active": {
                    "enabled": _bool(snap.active_full_enabled),
                    "kind": snap.active_full_kind,
                    "days": snap.active_full_days,
                    "monthly": snap.active_full_monthly,
                    "display": _full_display(snap.active_full_enabled, snap.active_full_kind,
                                             snap.active_full_days, snap.active_full_monthly),
                },
                "compact": {
                    "enabled": _bool(snap.compact_full_enabled),
                    "kind": snap.compact_full_kind,
                    "days": snap.compact_full_days,
                    "monthly": snap.compact_full_monthly,
                    "display": _full_display(snap.compact_full_enabled, snap.compact_full_kind,
                                             snap.compact_full_days, snap.compact_full_monthly),
                },
            },
        },
        "objects": [
            {
                "name": o.object_name,
                "type": o.object_type,
                "path": o.object_path,
                "size": o.approx_size,
                "guest_processing": _bool(o.guest_processing_enabled),
                "aap": _bool(o.application_aware_enabled),
                "sql_mode": o.sql_mode,
                "sql_log_backup": _bool(o.sql_log_backup_enabled),
                "sql_freq_min": o.sql_log_freq_min,
                "sql_retain_days": o.sql_log_retain_days,
                "sql_use_db_retention": _bool(o.sql_use_db_retention),
                "sql_truncate": _bool(o.sql_truncate_enabled),
            }
            for o in objs
        ],
    })


@router.get("/api/job-config/analytics/history")
def api_history(
    request: Request,
    job_id: Optional[str]   = Query(None),
    job_name: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    """
    Histórico de configuração de UMA rotina ao longo das coletas + changelog
    (o que mudou entre coletas consecutivas nos campos auditáveis).
    Identifica a rotina por job_id (preferencial) ou job_name.
    """
    allowed = _scoped_client_ids(db, request)

    q = db.query(JobConfigSnapshot)
    if allowed is not None:
        if not allowed:
            return JSONResponse({"snapshots": [], "changes": []})
        q = q.filter(JobConfigSnapshot.client_id.in_(list(allowed)))
    if job_id:
        q = q.filter(JobConfigSnapshot.job_id == job_id)
    elif job_name:
        q = q.filter(JobConfigSnapshot.job_name == job_name)
    else:
        return JSONResponse({"error": "informe job_id ou job_name"}, status_code=400)

    # ordem cronológica (mais antigo -> mais novo) para calcular o diff
    snaps = q.order_by(JobConfigSnapshot.collected_at.asc().nullslast()).all()
    if not snaps:
        return JSONResponse({"snapshots": [], "changes": []})

    def _label_dt(s):
        return s.collected_at.strftime("%d/%m/%Y %H:%M") if s.collected_at else "—"

    # snapshots resumidos (mais novo primeiro p/ exibição)
    snap_list = [{
        "id": s.id,
        "collected_at": _label_dt(s),
        "values": dict(_audited_values(s)),
    } for s in reversed(snaps)]

    # changelog: compara coletas consecutivas (cronológico)
    changes = []
    for prev, curr in zip(snaps, snaps[1:]):
        prev_vals = dict(_audited_values(prev))
        curr_vals = dict(_audited_values(curr))
        diffs = []
        for label, new_val in _audited_values(curr):
            old_val = prev_vals.get(label)
            if str(old_val) != str(new_val):
                diffs.append({"field": label, "from": old_val, "to": new_val})
        if diffs:
            changes.append({
                "date":      _label_dt(curr),
                "from_date": _label_dt(prev),
                "diffs":     diffs,
            })
    changes.reverse()   # mais recente primeiro

    return JSONResponse({
        "routine": {"job_id": job_id, "job_name": job_name or (snaps[-1].job_name)},
        "collections": len(snaps),
        "snapshots": snap_list,
        "changes": changes,
    })


@router.get("/api/job-config/analytics/export-excel")
def api_export_excel(
    request: Request,
    client_id: Optional[int] = Query(None),
    search: Optional[str]    = Query(None),
    job_type: Optional[str]  = Query(None),
    full_day: Optional[str]  = Query(None),
    db: Session = Depends(get_db),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    allowed = _scoped_client_ids(db, request)
    if allowed is not None and not allowed:
        return JSONResponse({"error": "sem acesso"}, status_code=403)

    cid_filter = client_id if (allowed is None or client_id in allowed) else None
    snap_ids = _latest_snapshot_ids(db, allowed, client_id_filter=cid_filter,
                                    search=search, job_type=job_type)
    if not snap_ids:
        return JSONResponse({"error": "sem dados"}, status_code=404)

    snaps = (db.query(JobConfigSnapshot)
             .filter(JobConfigSnapshot.id.in_(snap_ids))
             .order_by(JobConfigSnapshot.job_name).all())
    if full_day:
        snaps = [s for s in snaps if _full_day_match(s, full_day)]
    if not snaps:
        return JSONResponse({"error": "sem dados"}, status_code=404)

    sql_counts: dict[int, int] = {}
    obj_rows = (db.query(JobConfigObject.job_config_snapshot_id)
                .filter(JobConfigObject.job_config_snapshot_id.in_(snap_ids),
                        JobConfigObject.sql_log_backup_enabled.is_(True))
                .all())
    for (sid,) in obj_rows:
        sql_counts[sid] = sql_counts.get(sid, 0) + 1

    wb = Workbook()
    ws = wb.active
    ws.title = "Configuração de Rotinas"

    header_fill = PatternFill("solid", fgColor="003366")
    header_font = Font(color="FFFFFF", bold=True)
    headers = [
        "Rotina", "Tipo", "Habilitada", "Horário", "Repositório", "SOBR",
        "Imutável", "GIP", "App-Aware", "VMs c/ SQL Log", "Retenção",
        "Política", "GFS", "Full (quando)", "Retry", "Proxy", "Coletado em",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    def yn(v):
        if v is None: return "—"
        return "Sim" if v else "Não"

    for s in snaps:
        ws.append([
            s.job_name,
            s.job_type or "—",
            yn(s.is_enabled),
            s.sched_daily_time or "—",
            s.repo_sobr_name or s.repo_name or "—",
            yn(s.repo_is_sobr),
            yn(s.repo_is_sobr and s.repo_immutability),
            s.gip_mode or "—",
            yn(s.aap_enabled),
            sql_counts.get(s.id, 0),
            _retention_display(s),
            s.ret_storage_type or "—",
            yn(s.ret_gfs_enabled),
            _full_when(s),
            s.sched_retry_count if s.sched_retry_count is not None else "—",
            s.proxy_mode or "—",
            s.collected_at.strftime("%d/%m/%Y %H:%M") if s.collected_at else "—",
        ])

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    fname = f"config_rotinas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
