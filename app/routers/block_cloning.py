"""
block_cloning.py — Tela "Block Cloning / Rotinas espalhadas".

Detecta rotinas de backup cuja cadeia ATIVA (arquivos em disco, performance tier)
está espalhada em 2+ extents do SOBR. Quando isso acontece, o Veeam grava parte
da cadeia em outro repositório (porque o primeiro encheu) e o block cloning
(fast clone ReFS/XFS), que só funciona DENTRO do mesmo extent, deixa de valer —
o que infla o espaço real ocupado.

Requer coleta com extent por ponto (collect_disk_backups.ps1 -WithExtent). Sem
isso o extent_name vem nulo (a coleta padrão só enxerga o nome do SOBR) e a tela
mostra as instruções da coleta.

Fonte: disk_backup_points (coletor PowerShell de disco, modo -WithExtent).
"""
from __future__ import annotations

import io
from datetime import datetime

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import func, or_
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Client, DiskBackupPoint, DiskUsageImport, UserClient
from app.services import branding_service

router = APIRouter()
templates = Jinja2Templates(directory="app/templates")
templates.env.globals["get_branding"] = branding_service.get_branding
templates.env.globals["now"] = datetime.now


def _fmt(n):
    n = float(n or 0)
    for unit in ("B", "KB", "MB", "GB", "TB", "PB"):
        if abs(n) < 1024 or unit == "PB":
            return (f"{n:.0f} {unit}" if unit in ("B", "KB", "MB") else f"{n:.1f} {unit}")
        n /= 1024
    return f"{n:.1f} PB"


templates.env.filters["fmt_bytes"] = _fmt


def _scoped(db: Session, request: Request):
    if request.session.get("is_admin"):
        return None
    uid = request.session.get("user_id")
    if not uid:
        return set()
    return {r[0] for r in db.query(UserClient.client_id).filter(UserClient.user_id == uid).all()}


def _latest_import(db: Session, client_id: int):
    return (db.query(DiskUsageImport)
            .filter(DiskUsageImport.client_id == client_id,
                    DiskUsageImport.status == "completed")
            .order_by(DiskUsageImport.imported_at.desc())
            .first())


def _analyze(db: Session, client_id: int):
    """Estado da análise de block cloning para um cliente.

    O coletor grava em ``extent_name`` o CONJUNTO de extents de performance onde a
    cadeia daquele VM tem arquivos (ex.: "SRV-EXEMPLO, SRV-EXEMPLO"). Uma cadeia com 2+
    extents = block cloning quebrado. A unidade correta é a cadeia (rotina + VM):
    uma rotina com VMs diferentes em extents diferentes NÃO quebra clone — só quando
    o MESMO VM está espalhado.
    """
    imp = _latest_import(db, client_id)
    if not imp:
        return {"has_import": False}
    iid = imp.id

    n_extent = (db.query(func.count(DiskBackupPoint.id))
                .filter(DiskBackupPoint.disk_usage_import_id == iid,
                        DiskBackupPoint.extent_name.isnot(None),
                        DiskBackupPoint.extent_name != "")
                .scalar()) or 0
    if n_extent == 0:
        return {"has_import": True, "has_extent": False,
                "import_at": imp.imported_at, "import_id": iid}

    # agregado por (rotina, vm, conjunto-de-extents) — só arquivos LOCAIS
    rows = (db.query(DiskBackupPoint.job_name,
                     DiskBackupPoint.vm_name,
                     DiskBackupPoint.extent_name,
                     func.count(DiskBackupPoint.id),
                     func.coalesce(func.sum(DiskBackupPoint.size_bytes), 0))
            .filter(DiskBackupPoint.disk_usage_import_id == iid,
                    DiskBackupPoint.extent_name.isnot(None),
                    DiskBackupPoint.extent_name != "",
                    or_(DiskBackupPoint.in_capacity_tier.is_(False),
                        DiskBackupPoint.in_capacity_tier.is_(None)))
            .group_by(DiskBackupPoint.job_name, DiskBackupPoint.vm_name,
                      DiskBackupPoint.extent_name)
            .all())

    # consolida por cadeia (rotina, vm) — união dos extents vistos
    chains: dict[tuple, dict] = {}
    all_extents: set[str] = set()
    for job, vm, extent_name, cnt, sz in rows:
        key = (job or "(sem rotina)", vm or "(sem VM)")
        c = chains.setdefault(key, {"exts": set(), "size": 0, "points": 0})
        for e in (extent_name or "").split(","):
            e = e.strip()
            if e:
                c["exts"].add(e); all_extents.add(e)
        c["size"] += int(sz or 0)
        c["points"] += int(cnt or 0)

    spread = []
    jobs_affected: set[str] = set()
    for (job, vm), c in chains.items():
        if len(c["exts"]) >= 2:
            jobs_affected.add(job)
            spread.append({
                "job": job, "vm": vm,
                "extents": sorted(c["exts"]),
                "extents_txt": ", ".join(sorted(c["exts"])),
                "n_extents": len(c["exts"]),
                "size": c["size"], "points": c["points"],
            })
    spread.sort(key=lambda x: -x["size"])

    return {
        "has_import": True, "has_extent": True,
        "import_at": imp.imported_at, "import_id": iid,
        "spread": spread,
        "n_spread": len(spread),
        "n_chains_total": len(chains),
        "n_jobs_affected": len(jobs_affected),
        "total_afetado": sum(x["size"] for x in spread),
        "all_extents": sorted(all_extents),
    }


@router.get("/analytics/block-cloning", response_class=HTMLResponse)
def block_cloning_page(request: Request, client_id: int | None = Query(None),
                       db: Session = Depends(get_db)):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=303)
    allowed = _scoped(db, request)
    q = db.query(Client).filter(Client.is_active.is_(True))
    if allowed is not None:
        q = q.filter(Client.id.in_(allowed)) if allowed else q.filter(False)
    clients = q.order_by(Client.name).all()

    data = {}
    if client_id and (allowed is None or client_id in allowed):
        data = _analyze(db, client_id)

    return templates.TemplateResponse("analytics_block_cloning.html", {
        "request": request, "clients": clients,
        "selected_client_id": client_id, "d": data,
    })


@router.get("/analytics/block-cloning/export-excel")
def block_cloning_export(request: Request, client_id: int = Query(...),
                         db: Session = Depends(get_db)):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=303)
    allowed = _scoped(db, request)
    if allowed is not None and client_id not in allowed:
        return RedirectResponse(url="/analytics/block-cloning", status_code=303)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    d = _analyze(db, client_id)
    cli = db.query(Client).get(client_id)
    wb = Workbook(); ws = wb.active; ws.title = "Cadeias espalhadas"
    hf = PatternFill("solid", fgColor="7C3AED"); hfont = Font(color="FFFFFF", bold=True)
    ws.append(["Rotina", "VM", "Nr. extents", "Extents", "Tamanho em disco (GB)", "Pontos"])
    for c in ws[1]:
        c.fill = hf; c.font = hfont; c.alignment = Alignment(horizontal="center")
    for r in d.get("spread", []):
        ws.append([r["job"], r["vm"], r["n_extents"], r["extents_txt"],
                   round(r["size"] / 1073741824, 1), r["points"]])
    for col in ws.columns:
        wd = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(wd + 3, 80)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    cname = (cli.name if cli else "cliente").replace(" ", "_")
    fname = f"block_cloning_rotinas_espalhadas_{cname}_{datetime.now():%Y%m%d}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})
