"""
vm_missing.py — Tela "VMs Não Localizadas".

Lista os backups que falharam porque a VM foi removida/não localizada no cluster
(motivo "is unavailable and will be skipped" / "not found" / "no longer exists").
Fonte: sessões do coletor PowerShell de performance (backup_vm_sessions). O nome
da VM é extraído da mensagem (nesses casos a falha é no nível do job).

Colunas: Data · Hora · Rotina · VM não localizada · Motivo.
"""
from __future__ import annotations

import io
import re
from datetime import datetime

from fastapi import APIRouter, Depends, Query, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import Client, BackupVmSession, UserClient
from app.services import branding_service

router = APIRouter()
templates = Jinja2Templates(directory="app/templates")
templates.env.globals["get_branding"] = branding_service.get_branding
templates.env.globals["now"] = datetime.now

# Padrões de "VM não localizada / removida" na mensagem do Veeam.
_PATTERNS = [
    "%unavailable and will be skipped%",
    "%no longer exists%",
    "%was not found%",
    "%cannot be found%",
    "%not found in%",
]
_VM_RE = [
    re.compile(r"Virtual [Mm]achine ['\"]?(.+?)['\"]? (?:is unavailable|was not found|no longer exists|cannot be found)"),
    re.compile(r"\bVM ['\"]?(.+?)['\"]? (?:is unavailable|was not found|no longer|not found)"),
]


def _extract_vm(msg: str) -> str:
    msg = msg or ""
    for rx in _VM_RE:
        m = rx.search(msg)
        if m:
            return m.group(1).strip()
    return "(ver motivo)"


def _scoped(db: Session, request: Request):
    if request.session.get("is_admin"):
        return None
    uid = request.session.get("user_id")
    if not uid:
        return set()
    return {r[0] for r in db.query(UserClient.client_id).filter(UserClient.user_id == uid).all()}


def _query(db: Session, request: Request, client_id, date_start, date_end):
    q = db.query(BackupVmSession.started_at, BackupVmSession.job_name_snapshot,
                 BackupVmSession.reason_raw).filter(
        or_(*[BackupVmSession.reason_raw.ilike(p) for p in _PATTERNS]))
    allowed = _scoped(db, request)
    if allowed is not None:
        if not allowed:
            return []
        q = q.filter(BackupVmSession.client_id.in_(allowed))
    if client_id:
        q = q.filter(BackupVmSession.client_id == client_id)
    if date_start:
        q = q.filter(BackupVmSession.started_at >= date_start)
    if date_end:
        q = q.filter(BackupVmSession.started_at < f"{date_end} 23:59:59")
    rows = q.order_by(BackupVmSession.started_at).all()
    recs = []
    for started, job, reason in rows:
        recs.append({
            "date": started.strftime("%d/%m/%Y"),
            "time": started.strftime("%H:%M:%S"),
            "iso": started,
            "job": job or "?",
            "vm": _extract_vm(reason),
            "reason": (reason or "").replace("\n", " ").strip(),
        })
    return recs


def _summary(recs):
    agg = {}
    for r in recs:
        a = agg.setdefault(r["vm"], {"n": 0, "jobs": set(), "first": r["iso"], "last": r["iso"]})
        a["n"] += 1
        a["jobs"].add(r["job"])
        a["first"] = min(a["first"], r["iso"])
        a["last"] = max(a["last"], r["iso"])
    return [{"vm": vm, "n": a["n"], "jobs": ", ".join(sorted(a["jobs"])),
             "first": a["first"].strftime("%d/%m/%Y"), "last": a["last"].strftime("%d/%m/%Y")}
            for vm, a in sorted(agg.items(), key=lambda x: -x[1]["n"])]


@router.get("/analytics/vm-missing", response_class=HTMLResponse)
def vm_missing_page(request: Request, client_id: int | None = Query(None),
                    date_start: str | None = Query(None), date_end: str | None = Query(None),
                    db: Session = Depends(get_db)):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=303)
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    recs = _query(db, request, client_id, (date_start or "").strip() or None,
                  (date_end or "").strip() or None) if client_id else []
    return templates.TemplateResponse("analytics_vm_missing.html", {
        "request": request, "clients": clients, "selected_client_id": client_id,
        "recs": recs, "summary": _summary(recs),
        "f": {"date_start": date_start, "date_end": date_end},
    })


@router.get("/analytics/vm-missing/export-excel")
def vm_missing_export(request: Request, client_id: int = Query(...),
                      date_start: str | None = Query(None), date_end: str | None = Query(None),
                      db: Session = Depends(get_db)):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=303)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    recs = _query(db, request, client_id, (date_start or "").strip() or None,
                  (date_end or "").strip() or None)
    cli = db.query(Client).get(client_id)
    wb = Workbook(); ws = wb.active; ws.title = "VMs nao localizadas"
    hf = PatternFill("solid", fgColor="C0392B"); hfont = Font(color="FFFFFF", bold=True)
    ws.append(["Data", "Hora", "Rotina", "VM nao localizada", "Motivo (Veeam)"])
    for c in ws[1]:
        c.fill = hf; c.font = hfont; c.alignment = Alignment(horizontal="center")
    for r in recs:
        ws.append([r["date"], r["time"], r["job"], r["vm"], r["reason"][:180]])
    ws2 = wb.create_sheet("Resumo por VM")
    ws2.append(["VM nao localizada", "Ocorrencias", "Primeiro", "Ultimo", "Rotina(s)"])
    for c in ws2[1]:
        c.fill = hf; c.font = hfont
    for s in _summary(recs):
        ws2.append([s["vm"], s["n"], s["first"], s["last"], s["jobs"]])
    for w in (ws, ws2):
        for col in w.columns:
            wd = max((len(str(c.value or "")) for c in col), default=10)
            w.column_dimensions[col[0].column_letter].width = min(wd + 3, 70)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    cname = (cli.name if cli else "cliente").replace(" ", "_")
    fname = f"vms_nao_localizadas_{cname}_{datetime.now():%Y%m%d}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})
