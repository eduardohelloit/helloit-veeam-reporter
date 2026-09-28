"""
env_report.py — Classificação de Ambiente + Consumo por Ambiente.

Automatiza o relatório mensal: separa o espaço em disco consumido por
HOMOLOGAÇÃO x PRODUÇÃO no Performance Tier do SOBR, mais o total e o livre —
gerando a pizza (Produção / Homologação / Livre).

Método (igual ao manual do usuário, "Backup size" da Veeam):
  • Total / Livre / Usado  -> repo_capacity_samples (tier Performance), físico;
  • Homologação            -> Σ size_bytes (BackupSize bruto) dos jobs marcados
                              como 'hml' na última carga de "Backups em Disco";
  • Produção               -> Usado(físico) − Homologação.

A marcação de quais jobs são de homologação vive em job_environments (flag por
cliente/rotina; sem tag = produção).
"""
from __future__ import annotations

import io
from datetime import datetime

from fastapi import APIRouter, Depends, Form, Query, Request
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.database import get_db
from app.models import (
    Client, DiskUsageImport, DiskBackupPoint, RepoCapacitySample, JobEnvironment,
)
from app.services import branding_service
from app.routers.disk_analytics import _fmt_bytes

router = APIRouter()
templates = Jinja2Templates(directory="app/templates")
templates.env.globals["get_branding"] = branding_service.get_branding
templates.env.globals["now"] = datetime.now
templates.env.filters["fmt_bytes"] = _fmt_bytes

GB = 1024 ** 3
TB = 1024 ** 4


def _latest_import(db: Session, client_id: int):
    return (db.query(DiskUsageImport)
            .filter(DiskUsageImport.client_id == client_id,
                    DiskUsageImport.status == "completed")
            .order_by(DiskUsageImport.imported_at.desc()).first())


def _job_sizes(db: Session, import_id: int):
    """Σ BackupSize (bruto) por job na carga; lista [(job, bytes, pts)]."""
    rows = (db.query(DiskBackupPoint.job_name,
                     func.coalesce(func.sum(DiskBackupPoint.size_bytes), 0),
                     func.count())
            .filter(DiskBackupPoint.disk_usage_import_id == import_id,
                    DiskBackupPoint.job_name.isnot(None))
            .group_by(DiskBackupPoint.job_name)
            .order_by(func.sum(DiskBackupPoint.size_bytes).desc()).all())
    return [(r[0], int(r[1] or 0), r[2]) for r in rows]


def _env_map(db: Session, client_id: int) -> dict:
    return {e.job_name: e.environment
            for e in db.query(JobEnvironment).filter(JobEnvironment.client_id == client_id).all()}


def _perf_capacity(db: Session, import_id: int):
    """Total / usado / livre físico do Performance tier."""
    r = (db.query(func.coalesce(func.sum(RepoCapacitySample.total_bytes), 0),
                  func.coalesce(func.sum(RepoCapacitySample.used_bytes), 0),
                  func.coalesce(func.sum(RepoCapacitySample.free_bytes), 0))
         .filter(RepoCapacitySample.disk_usage_import_id == import_id,
                 RepoCapacitySample.tier == "Performance").first())
    return int(r[0] or 0), int(r[1] or 0), int(r[2] or 0)


def _report(db: Session, client_id: int) -> dict:
    imp = _latest_import(db, client_id)
    if not imp:
        return {"has_data": False}
    sizes = _job_sizes(db, imp.id)
    envs = _env_map(db, client_id)
    total, used, free = _perf_capacity(db, imp.id)

    hml_bytes = sum(b for j, b, _p in sizes if envs.get(j) == "hml")
    hml_jobs = [{"job": j, "bytes": b, "pts": p}
                for j, b, p in sizes if envs.get(j) == "hml"]
    # Produção = usado físico − homologação (método do usuário)
    prod_bytes = max(used - hml_bytes, 0)

    return {
        "has_data": True,
        "imp": imp,
        "total": total, "used": used, "free": free,
        "hml": hml_bytes, "prod": prod_bytes,
        "hml_jobs": hml_jobs, "n_hml": len(hml_jobs), "n_jobs": len(sizes),
        "pie": {  # TB com 1 casa, ordem/cores iguais ao relatório atual
            "labels": ["Protegido Produção", "Protegido Homologação", "Espaço Livre"],
            "tb": [round(prod_bytes / TB, 1), round(hml_bytes / TB, 1), round(free / TB, 1)],
        },
    }


# ── TELA: relatório de consumo por ambiente ─────────────────────────────────
@router.get("/analytics/environment", response_class=HTMLResponse)
def environment_report(request: Request, client_id: int | None = Query(None),
                       db: Session = Depends(get_db)):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=303)
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    data = {"has_data": False}
    if client_id:
        data = _report(db, client_id)
    return templates.TemplateResponse("analytics_environment.html", {
        "request": request, "clients": clients,
        "selected_client_id": client_id, "data": data,
    })


# ── TELA: classificação (marcar jobs de homologação) ────────────────────────
@router.get("/environment/classify", response_class=HTMLResponse)
def classify_page(request: Request, client_id: int | None = Query(None),
                  db: Session = Depends(get_db)):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=303)
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    jobs = []
    if client_id:
        imp = _latest_import(db, client_id)
        if imp:
            envs = _env_map(db, client_id)
            for j, b, p in _job_sizes(db, imp.id):
                jobs.append({"job": j, "bytes": b, "pts": p,
                             "env": envs.get(j, "prod")})
    return templates.TemplateResponse("environment_classify.html", {
        "request": request, "clients": clients,
        "selected_client_id": client_id, "jobs": jobs,
    })


@router.post("/environment/classify")
async def classify_save(request: Request, db: Session = Depends(get_db)):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=303)
    form = await request.form()
    client_id = int(form.get("client_id"))
    # checkboxes marcados = homologação; o hidden 'all_jobs' traz todos os nomes
    hml = set(form.getlist("hml"))
    all_jobs = form.getlist("all_jobs")
    existing = {e.job_name: e for e in
                db.query(JobEnvironment).filter(JobEnvironment.client_id == client_id).all()}
    now = datetime.utcnow()
    for job in all_jobs:
        env = "hml" if job in hml else "prod"
        row = existing.get(job)
        if row:
            if row.environment != env:
                row.environment = env; row.updated_at = now
        else:
            db.add(JobEnvironment(client_id=client_id, job_name=job,
                                  environment=env, updated_at=now))
    db.commit()
    return RedirectResponse(url=f"/analytics/environment?client_id={client_id}", status_code=303)


# ── Export Excel ────────────────────────────────────────────────────────────
@router.get("/analytics/environment/export-excel")
def environment_export(request: Request, client_id: int = Query(...),
                       db: Session = Depends(get_db)):
    if not request.session.get("user_id"):
        return RedirectResponse(url="/login", status_code=303)
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    data = _report(db, client_id)
    if not data.get("has_data"):
        return RedirectResponse(url="/analytics/environment", status_code=303)

    wb = Workbook(); ws = wb.active; ws.title = "Consumo por Ambiente"
    hf = PatternFill("solid", fgColor="0B5ED7"); hfont = Font(color="FFFFFF", bold=True)
    ws.append(["Métrica", "TB"])
    for c in ws[1]:
        c.fill = hf; c.font = hfont
    ws.append(["Total do repositório (Performance)", round(data["total"] / TB, 1)])
    ws.append(["Usado", round(data["used"] / TB, 1)])
    ws.append(["Protegido Produção", round(data["prod"] / TB, 1)])
    ws.append(["Protegido Homologação", round(data["hml"] / TB, 1)])
    ws.append(["Espaço Livre", round(data["free"] / TB, 1)])

    ws2 = wb.create_sheet("Jobs Homologação")
    ws2.append(["Job", "Backup size (GB)", "Pontos"])
    for c in ws2[1]:
        c.fill = hf; c.font = hfont
    for j in data["hml_jobs"]:
        ws2.append([j["job"], round(j["bytes"] / GB, 1), j["pts"]])

    for w in (ws, ws2):
        for col in w.columns:
            wd = max((len(str(c.value or "")) for c in col), default=10)
            w.column_dimensions[col[0].column_letter].width = min(wd + 3, 48)

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    fname = f"consumo_ambiente_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})
