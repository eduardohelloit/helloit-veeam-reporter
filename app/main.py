"""
HelloIT Veeam Reporter

Princípio arquitetural:
  Upload  = carga de dados → alimenta ImportedEvent (base histórica)
  Relatório = consulta por cliente + período → gerado independente de upload
"""
import os
import uuid
import json
import time
import shutil
import secrets
import tempfile
import threading
from collections import defaultdict, deque
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional
from urllib.parse import urlparse

from fastapi import FastAPI, Request, UploadFile, File, Form, Depends, HTTPException, Query
from fastapi.responses import HTMLResponse, RedirectResponse, FileResponse, JSONResponse, Response
from fastapi import Body
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
from starlette.middleware.sessions import SessionMiddleware

from app.database import get_db, init_db, SessionLocal
from app.models import (
    Client, ImportedEvent, JobAction, ReportSession, JobSummary, UploadSession,
    OffenderCategory, OffenderRule,
    RpoPolicy, VmRpoAssignment, UserClient, User,
)
from app.services.week_service import format_week_label
from app.services.report_service import build_report
from app.services.classification_service import situation_badge_class
from app.services.excel_service import generate_excel
from app.routers.auth_admin import router as auth_router
from app.routers.offload_admin import router as offload_router
from app.routers.offload_analytics import router as offload_analytics_router
from app.routers.backup_admin import router as backup_router
from app.routers.backup_analytics import router as backup_analytics_router
from app.routers.job_config_admin import router as job_config_router
from app.routers.job_config_analytics import router as job_config_analytics_router
from app.routers.disk_admin import router as disk_router
from app.routers.disk_analytics import router as disk_analytics_router
from app.routers.disk_growth import router as disk_growth_router
from app.routers.env_report import router as env_report_router
from app.routers.ps_import import router as ps_import_router
from app.routers.vm_missing import router as vm_missing_router
from app.routers.block_cloning import router as block_cloning_router
from app.auth import require_admin
from app.services import branding_service
from app.services import offender_service

UPLOAD_DIR   = Path(os.environ.get("UPLOAD_DIR",   "./uploads"))
REPORTS_DIR  = Path(os.environ.get("REPORTS_DIR",  "./reports"))
BRANDING_DIR = Path(os.environ.get("BRANDING_DIR", "./branding"))
ALLOWED_EXTENSIONS = {".evtx", ".xml", ".csv"}
# Limite de tamanho para importação de Event Viewer (anti-DoS de disco/CPU)
MAX_IMPORT_BYTES = int(os.environ.get("MAX_IMPORT_BYTES", str(2 * 1024 * 1024 * 1024)))  # 2 GB

# ── Limites de upload de branding ────────────────────────────────────────────
ALLOWED_IMAGE_EXT = {".png", ".jpg", ".jpeg", ".svg", ".webp"}
ALLOWED_FAVICON_EXT = {".ico", ".png"}
ALLOWED_PPTX_EXT  = {".pptx"}
MAX_LOGO_BYTES    = 2 * 1024 * 1024   # 2 MB
MAX_FAV_BYTES     = 512 * 1024        # 512 KB
MAX_PPTX_BYTES    = 20 * 1024 * 1024  # 20 MB

APP_VERSION = "1.0.0"

# ── Security ────────────────────────────────────────────────────────────────────
_SESSION_SECRET = os.environ.get("SESSION_SECRET", "")
if not _SESSION_SECRET:
    import secrets
    _SESSION_SECRET = secrets.token_hex(32)

_HTTPS_ONLY = os.environ.get("HTTPS_ONLY", "false").lower() == "true"

# Content-Security-Policy — modo Report-Only para não quebrar os CDNs atuais
# (jsdelivr p/ Bootstrap/Chart.js/ícones, Google Fonts). Promover a bloqueio depois
# de hospedar assets localmente / validar violações.
_CSP_POLICY = (
    "default-src 'self'; "
    "script-src 'self' https://cdn.jsdelivr.net; "
    "style-src 'self' 'unsafe-inline' https://cdn.jsdelivr.net https://fonts.googleapis.com; "
    "font-src 'self' https://fonts.gstatic.com https://cdn.jsdelivr.net; "
    "img-src 'self' data:; "
    "connect-src 'self'; "
    "frame-ancestors 'self'; base-uri 'self'; form-action 'self'"
)
# Prefixos que podem ser cacheados (assets/estáticos); o resto recebe no-store.
_CACHEABLE_PREFIXES = ("/static/", "/branding/file/", "/branding.css")

# ── App setup ───────────────────────────────────────────────────────────────────
app = FastAPI(title="HelloIT Veeam Reporter", docs_url=None, redoc_url=None)

BASE_DIR = Path(__file__).parent
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))
templates.env.globals["situation_badge_class"] = situation_badge_class
templates.env.globals["now"] = datetime.now
templates.env.globals["APP_VERSION"] = APP_VERSION
templates.env.globals["get_branding"] = branding_service.get_branding
templates.env.filters["format_number"] = lambda n: f"{n:,}".replace(",", ".")

app.include_router(auth_router)
app.include_router(offload_router)
app.include_router(offload_analytics_router)
app.include_router(backup_router)
app.include_router(backup_analytics_router)
app.include_router(job_config_router)
app.include_router(job_config_analytics_router)
app.include_router(disk_router)
app.include_router(disk_analytics_router)
app.include_router(disk_growth_router)
app.include_router(env_report_router)
app.include_router(ps_import_router)
app.include_router(vm_missing_router)
app.include_router(block_cloning_router)

_PUBLIC_PREFIXES     = ("/static/", "/login", "/favicon.ico", "/branding.css", "/branding/file/")
_FORCE_CHANGE_EXEMPT = ("/auth/force-change-password", "/logout", "/static/", "/login")

# ── Rate limiting (in-memory; uvicorn single-worker) ─────────────────────────
_UNSAFE_METHODS = {"POST", "PUT", "PATCH", "DELETE"}
_RATE_EXEMPT_PREFIXES = ("/static/", "/branding/file/", "/branding.css", "/favicon.ico")
_rate_hits: "dict[str, deque]" = defaultdict(deque)


def _client_ip(request: Request) -> str:
    xff = request.headers.get("x-forwarded-for")
    if xff:
        return xff.split(",")[0].strip()
    return request.client.host if request.client else "unknown"


def _rate_bucket(path: str, method: str):
    """(nome_do_bucket, limite_por_minuto) conforme o tipo de rota."""
    if path == "/login" and method == "POST":
        return "login", 10
    if path.startswith("/api/export") or path == "/import" or path.startswith("/api/admin/reprocess"):
        return "heavy", 20
    return "general", 200


def _rate_limited(key: str, limit: int, window: float = 60.0) -> bool:
    now = time.monotonic()
    dq = _rate_hits[key]
    while dq and now - dq[0] > window:
        dq.popleft()
    if len(dq) >= limit:
        return True
    dq.append(now)
    return False


# ── CSRF: Origin/Referer same-origin + token de sessão (header X-CSRF-Token) ──
_CSRF_EXEMPT_PREFIXES = ("/login",)


def _same_origin(request: Request) -> bool:
    host = request.headers.get("host", "")
    for hdr in ("origin", "referer"):
        val = request.headers.get(hdr)
        if val:
            return urlparse(val).netloc == host
    return False  # método mutável sem Origin/Referer → tratado como suspeito


@app.middleware("http")
async def security_middleware(request: Request, call_next):
    path      = request.url.path
    method    = request.method
    is_public = any(path.startswith(p) for p in _PUBLIC_PREFIXES)

    # ── Rate limiting (ignora estáticos) ─────────────────────────────────────
    if not path.startswith(_RATE_EXEMPT_PREFIXES):
        bucket, limit = _rate_bucket(path, method)
        ident = _client_ip(request) if bucket == "login" else (
            request.session.get("username") or _client_ip(request)
        )
        if _rate_limited(f"{bucket}:{ident}", limit):
            return JSONResponse(
                {"detail": "Limite de requisições excedido. Tente novamente em instantes."},
                status_code=429,
            )

    if not is_public:
        username = request.session.get("username", "")
        if not username:
            return RedirectResponse(url="/login", status_code=302)
        must_change = request.session.get("must_change_password", False)
        if must_change and not any(path.startswith(p) for p in _FORCE_CHANGE_EXEMPT):
            return RedirectResponse(url="/auth/force-change-password", status_code=302)
        request.state.username = username
        request.state.is_admin = request.session.get("is_admin", False)
        # Garante token CSRF para a sessão autenticada
        if not request.session.get("csrf_token"):
            request.session["csrf_token"] = secrets.token_hex(32)
    else:
        request.state.username = request.session.get("username", "")
        request.state.is_admin = request.session.get("is_admin", False)

    # ── CSRF: métodos mutáveis em rotas autenticadas ─────────────────────────
    if (
        method in _UNSAFE_METHODS
        and not is_public
        and not any(path.startswith(p) for p in _CSRF_EXEMPT_PREFIXES)
    ):
        token = request.session.get("csrf_token")
        sent  = request.headers.get("x-csrf-token")
        header_ok = bool(token and sent and secrets.compare_digest(sent, token))
        if not (header_ok or _same_origin(request)):
            return JSONResponse(
                {"detail": "Requisição bloqueada por proteção CSRF (origem inválida ou token ausente)."},
                status_code=403,
            )

    response = await call_next(request)
    response.headers["X-Content-Type-Options"]  = "nosniff"
    response.headers["X-Frame-Options"]         = "SAMEORIGIN"
    response.headers["X-XSS-Protection"]        = "1; mode=block"
    response.headers["Referrer-Policy"]         = "strict-origin-when-cross-origin"
    response.headers["Permissions-Policy"]      = "geolocation=(), microphone=(), camera=()"
    response.headers["Content-Security-Policy-Report-Only"] = _CSP_POLICY
    response.headers["Cross-Origin-Opener-Policy"]   = "same-origin"
    response.headers["Cross-Origin-Resource-Policy"] = "same-origin"
    if not any(path.startswith(p) for p in _CACHEABLE_PREFIXES):
        response.headers["Cache-Control"] = "no-store"
    if _HTTPS_ONLY:
        response.headers["Strict-Transport-Security"] = "max-age=31536000; includeSubDomains"
    return response


@app.exception_handler(403)
async def forbidden_handler(request: Request, exc):
    return templates.TemplateResponse("auth/403.html", {"request": request}, status_code=403)


app.add_middleware(
    SessionMiddleware,
    secret_key=_SESSION_SECRET,
    session_cookie="veeam_session",
    max_age=28800,
    https_only=_HTTPS_ONLY,
    same_site="lax",
)


@app.on_event("startup")
def on_startup():
    UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
    REPORTS_DIR.mkdir(parents=True, exist_ok=True)
    (BRANDING_DIR / "logos").mkdir(parents=True, exist_ok=True)
    (BRANDING_DIR / "pptx").mkdir(parents=True, exist_ok=True)
    init_db()


# ── Login / Logout ──────────────────────────────────────────────────────────────

@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, error: str = Query(default="")):
    if request.session.get("username"):
        return RedirectResponse(url="/", status_code=302)
    return templates.TemplateResponse("login.html", {"request": request, "error": error})


@app.post("/login")
async def login_post(
    request: Request,
    username: str = Form(...),
    password: str = Form(...),
    db: Session = Depends(get_db),
):
    from app.auth import authenticate_user
    client_ip = request.client.host or "unknown"
    user, error = authenticate_user(db, username.strip(), password, ip=client_ip)

    if error:
        status = 429 if "bloqueada" in error.lower() else 401
        return templates.TemplateResponse(
            "login.html", {"request": request, "error": error}, status_code=status
        )

    request.session["username"]             = user.username
    request.session["user_id"]              = user.id
    request.session["is_admin"]             = user.is_admin
    request.session["must_change_password"] = user.must_change_password

    if user.must_change_password:
        return RedirectResponse(url="/auth/force-change-password", status_code=303)
    return RedirectResponse(url="/", status_code=303)


@app.get("/logout")
def logout(request: Request, db: Session = Depends(get_db)):
    from app.auth import log_security_event
    username = request.session.get("username")
    ip = request.client.host if request.client else "unknown"
    if username:
        log_security_event(db, "logout", username=username, actor=username, ip_address=ip)
    request.session.clear()
    return RedirectResponse(url="/login", status_code=302)


@app.get("/", response_class=HTMLResponse)
def index():
    return RedirectResponse(url="/import")


# ══════════════════════════════════════════════════════════════════════════════
# BRANDING / IDENTIDADE VISUAL
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/branding.css")
def branding_css():
    """CSS dinâmico com as variáveis de branding. Público (usado no login também)."""
    b = branding_service.get_branding()
    css = branding_service.build_css(b)
    return Response(content=css, media_type="text/css",
                    headers={"Cache-Control": "no-cache, must-revalidate"})


@app.get("/branding/file/{path:path}")
def branding_file(path: str):
    """Serve arquivos de branding (logos, templates). Público."""
    # Segurança: impede path traversal
    safe = path.replace("\\", "/").strip("/")
    if ".." in safe or safe.startswith("/"):
        raise HTTPException(status_code=400, detail="Path inválido")
    file_path = BRANDING_DIR / safe
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="Arquivo não encontrado")
    return FileResponse(str(file_path))


@app.get("/admin/branding", response_class=HTMLResponse)
def admin_branding_page(request: Request):
    """Página administrativa de Identidade Visual (apenas admin)."""
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    return templates.TemplateResponse("admin_branding.html", {"request": request})


@app.get("/api/admin/branding")
def api_list_branding(request: Request, db: Session = Depends(get_db)):
    """Lista todos os perfis de branding."""
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    return branding_service.get_all_profiles(db)


@app.get("/api/admin/branding/{profile_id}")
def api_get_branding(profile_id: int, request: Request, db: Session = Depends(get_db)):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    p = branding_service.get_profile_by_id(db, profile_id)
    if p is None:
        raise HTTPException(status_code=404, detail="Perfil não encontrado")
    return p


@app.post("/api/admin/branding")
def api_create_branding(
    request: Request,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    """Cria um novo perfil de branding (tipicamente por cliente)."""
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    from app.models import BrandingProfile

    client_id = data.get("client_id")
    name      = data.get("name", "Novo perfil")

    if client_id:
        existing = db.query(BrandingProfile).filter(
            BrandingProfile.client_id == client_id
        ).first()
        if existing:
            raise HTTPException(status_code=400, detail="Cliente já possui perfil de branding")

    # Cria baseado nos defaults HelloIT
    sd = branding_service.BASE_DEFAULTS
    from datetime import datetime as _dt
    profile = BrandingProfile(
        name          = name,
        is_default    = False,
        is_base   = False,
        client_id     = client_id,
        system_name   = sd["system_name"],
        short_name    = sd["short_name"],
        company_name  = sd["company_name"],
        header_text   = sd["header_text"],
        footer_text   = sd["footer_text"],
        website_url   = sd["website_url"],
        contact_email = "",
        primary_color        = sd["primary_color"],
        secondary_color      = sd["secondary_color"],
        accent_color         = sd["accent_color"],
        success_color        = sd["success_color"],
        warning_color        = sd["warning_color"],
        danger_color         = sd["danger_color"],
        background_color     = sd["background_color"],
        card_bg_color        = sd["card_bg_color"],
        table_header_color   = sd["table_header_color"],
        text_primary_color   = sd["text_primary_color"],
        text_secondary_color = sd["text_secondary_color"],
        created_at    = _dt.utcnow(),
        created_by    = request.state.username,
    )
    db.add(profile)
    db.commit()
    db.refresh(profile)
    branding_service.invalidate_cache()
    return branding_service._profile_to_dict(profile)


@app.put("/api/admin/branding/{profile_id}")
def api_update_branding(
    profile_id: int,
    request: Request,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    try:
        return branding_service.update_profile(db, profile_id, data, actor=request.state.username)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.delete("/api/admin/branding/{profile_id}")
def api_delete_branding(
    profile_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    from app.models import BrandingProfile
    profile = db.query(BrandingProfile).get(profile_id)
    if profile is None:
        raise HTTPException(status_code=404)
    if profile.is_default and profile.client_id is None:
        raise HTTPException(status_code=400, detail="Não é possível excluir o perfil global padrão")
    db.delete(profile)
    db.commit()
    branding_service.invalidate_cache()
    return {"ok": True}


@app.post("/api/admin/branding/{profile_id}/restore-base")
def api_restore_base(
    profile_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    try:
        return branding_service.reset_to_base(db, profile_id, actor=request.state.username)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))


@app.post("/api/admin/branding/{profile_id}/upload")
async def api_upload_branding_file(
    profile_id: int,
    request: Request,
    file: UploadFile = File(...),
    field: str = Form(...),
    db: Session = Depends(get_db),
):
    """Upload de logo ou template PPTX para um perfil de branding."""
    if not request.state.is_admin:
        raise HTTPException(status_code=403)

    ext = Path(file.filename or "").suffix.lower()
    is_pptx   = field in ("pptx_cover_path", "pptx_final_slide_path")
    is_favicon = field == "favicon_path"

    # Validação de extensão
    if is_pptx:
        if ext not in ALLOWED_PPTX_EXT:
            raise HTTPException(status_code=400, detail="Formato inválido. Use .pptx")
        max_bytes = MAX_PPTX_BYTES
        subdir = "pptx"
    elif is_favicon:
        if ext not in ALLOWED_FAVICON_EXT:
            raise HTTPException(status_code=400, detail="Formato inválido. Use .ico ou .png")
        max_bytes = MAX_FAV_BYTES
        subdir = "logos"
    else:
        if ext not in ALLOWED_IMAGE_EXT:
            raise HTTPException(status_code=400, detail="Formato inválido. Use PNG, JPG, SVG ou WEBP")
        max_bytes = MAX_LOGO_BYTES
        subdir = "logos"

    # Lê conteúdo com limite de tamanho
    content = await file.read(max_bytes + 1)
    if len(content) > max_bytes:
        raise HTTPException(status_code=413, detail=f"Arquivo muito grande. Máximo: {max_bytes//1024//1024} MB")

    # Salva com nome UUID para evitar colisões e path traversal
    filename = f"{uuid.uuid4().hex}{ext}"
    dest_dir = BRANDING_DIR / subdir
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / filename
    dest.write_bytes(content)

    # Remove arquivo anterior se existir
    from app.models import BrandingProfile
    profile = db.query(BrandingProfile).get(profile_id)
    if profile:
        old_path = getattr(profile, field, None)
        if old_path:
            old_file = BRANDING_DIR / old_path
            if old_file.exists():
                try:
                    old_file.unlink()
                except Exception:
                    pass

    relative_path = f"{subdir}/{filename}"
    try:
        branding_service.update_logo_path(db, profile_id, field, relative_path, actor=request.state.username)
    except ValueError as e:
        dest.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail=str(e))

    return {"ok": True, "path": relative_path}


@app.post("/api/admin/branding/{profile_id}/remove-file")
def api_remove_branding_file(
    profile_id: int,
    request: Request,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    """Remove um arquivo de logo/PPTX do perfil."""
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    field = data.get("field", "")
    from app.models import BrandingProfile
    profile = db.query(BrandingProfile).get(profile_id)
    if not profile:
        raise HTTPException(status_code=404)

    old_path = getattr(profile, field, None)
    if old_path:
        old_file = BRANDING_DIR / old_path
        if old_file.exists():
            try:
                old_file.unlink()
            except Exception:
                pass

    try:
        branding_service.update_logo_path(db, profile_id, field, None, actor=request.state.username)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))

    return {"ok": True}


@app.get("/api/clients")
def api_clients_list(request: Request, db: Session = Depends(get_db)):
    """Lista clientes ativos (escopados ao usuário; admin vê todos)."""
    return _all_clients_list(db, request)


# Coletores PowerShell servidos como TEXTO (para copiar na tela; evita download/AV).
_COLLECTOR_SCRIPTS = {
    "offload": ("SOBR.ps1",              "coleta_offload.ps1"),
    "backups": ("collect_backups.ps1",   "coleta_backups.ps1"),
    "config":  ("collect_job_config.ps1","coleta_config_rotinas.ps1"),
}


@app.get("/api/collector-script/{key}", dependencies=[Depends(require_admin)])
def api_collector_script(key: str):
    """Retorna o texto do coletor PowerShell (somente admin) para exibir/copiar na tela."""
    from fastapi.responses import PlainTextResponse
    entry = _COLLECTOR_SCRIPTS.get(key)
    if not entry:
        raise HTTPException(status_code=404, detail="Script não encontrado.")
    fname, _download = entry
    path = Path(__file__).parent.parent / "scripts" / "collectors" / fname
    if not path.exists():
        raise HTTPException(status_code=404, detail="Arquivo do script indisponível na instalação.")
    return PlainTextResponse(path.read_text(encoding="utf-8"),
                             media_type="text/plain; charset=utf-8")


# ══════════════════════════════════════════════════════════════════════════════
# IMPORTAÇÃO DE EVENTOS
# Upload alimenta a base histórica — não gera relatório automaticamente.
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/import", response_class=HTMLResponse, dependencies=[Depends(require_admin)])
def import_page(request: Request, db: Session = Depends(get_db)):
    sessions = (
        db.query(UploadSession)
        .order_by(UploadSession.uploaded_at.desc())
        .limit(15)
        .all()
    )
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse(
        "import.html",
        {"request": request, "sessions": sessions, "clients": clients},
    )




@app.post("/import", dependencies=[Depends(require_admin)])
async def import_file(
    request: Request,
    file: UploadFile = File(...),
    client_id: int = Form(...),
    db: Session = Depends(get_db),
):
    client = db.query(Client).get(client_id)
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()

    def _error(msg: str, status: int = 400):
        sessions = (
            db.query(UploadSession)
            .order_by(UploadSession.uploaded_at.desc())
            .limit(15)
            .all()
        )
        return templates.TemplateResponse(
            "import.html",
            {"request": request, "error": msg, "sessions": sessions, "clients": clients},
            status_code=status,
        )

    if not client:
        return _error("Cliente não encontrado.")

    suffix = Path(file.filename).suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        return _error(f"Formato não suportado: {suffix}. Use .evtx, .xml ou .csv.")

    # Rejeição precoce por Content-Length declarado (anti-DoS)
    declared = request.headers.get("content-length")
    if declared and declared.isdigit() and int(declared) > MAX_IMPORT_BYTES:
        return _error(
            f"Arquivo acima do limite permitido ({MAX_IMPORT_BYTES // (1024 * 1024)} MB).",
            status=413,
        )

    stored_name = f"{uuid.uuid4()}{suffix}"
    stored_path = UPLOAD_DIR / stored_name

    # Cópia em streaming com corte rígido: aborta e remove parcial se exceder o limite
    written = 0
    try:
        with stored_path.open("wb") as f_out:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                written += len(chunk)
                if written > MAX_IMPORT_BYTES:
                    f_out.close()
                    stored_path.unlink(missing_ok=True)
                    return _error(
                        f"Arquivo acima do limite permitido ({MAX_IMPORT_BYTES // (1024 * 1024)} MB).",
                        status=413,
                    )
                f_out.write(chunk)
    except Exception:
        stored_path.unlink(missing_ok=True)
        raise

    file_size_mb = stored_path.stat().st_size / (1024 * 1024)

    session = UploadSession(
        client_id=client_id,
        original_filename=file.filename,
        stored_filename=stored_name,
        file_format=suffix.lstrip("."),
        uploaded_by=request.session.get("username"),
        status="processing",
        status_message=f"Arquivo recebido ({file_size_mb:.1f} MB). Iniciando leitura...",
    )
    db.add(session)
    db.commit()
    db.refresh(session)
    session_id = session.id

    thread = threading.Thread(
        target=_process_import,
        args=(session_id, stored_path, suffix, client_id, file.filename),
        daemon=True,
    )
    thread.start()

    return RedirectResponse(url=f"/import/processing/{session_id}", status_code=303)


def _update_status(session_id: int, message: str, status: str = "processing"):
    db = SessionLocal()
    try:
        s = db.query(UploadSession).get(session_id)
        if s:
            s.status_message = message
            s.status = status
            db.commit()
    finally:
        db.close()


def _process_import(
    session_id: int,
    stored_path: Path,
    suffix: str,
    client_id: int,
    original_filename: str,
):
    """
    Worker em background: parse → dedup → insert ImportedEvents.
    NÃO cria ReportSession — apenas alimenta a base histórica.
    """
    from app.parser.csv_parser import CSVColumnError

    try:
        _update_status(session_id, "Lendo e filtrando eventos do arquivo...")

        events = _parse_all_events(stored_path, suffix, original_filename)
        total_read = len(events)

        # ── Detectar período ──────────────────────────────────────────────────
        dates = [ev.time_created for ev in events if ev.time_created]
        period_start = min(dates).replace(hour=0, minute=0, second=0, microsecond=0) if dates else None
        period_end   = max(dates).replace(hour=23, minute=59, second=59, microsecond=0) if dates else None

        # ── Contagens por categoria ────────────────────────────────────────────
        backup_count = sum(1 for e in events if e.event_category in ("backup", "log_backup"))
        audit_count  = sum(1 for e in events if e.event_category == "audit")

        _update_status(session_id, f"{total_read} eventos lidos. Verificando duplicatas...")

        # ── Deduplicação e insert via PostgreSQL ON CONFLICT DO NOTHING ───────
        db = SessionLocal()
        try:
            new_count, dup_count = _insert_events(db, session_id, client_id, events)

            s = db.query(UploadSession).get(session_id)
            s.period_start              = period_start
            s.period_end                = period_end
            s.total_events_read         = total_read
            s.new_events_inserted       = new_count
            s.duplicate_events_skipped  = dup_count
            s.backup_events_count       = backup_count
            s.audit_events_count        = audit_count
            s.status                    = "done"

            period_str = ""
            if period_start and period_end:
                period_str = (
                    f" | Período: {period_start.strftime('%d/%m/%Y')} → "
                    f"{period_end.strftime('%d/%m/%Y')}"
                )

            s.status_message = (
                f"Carga concluída. {new_count} eventos novos inseridos"
                + (f", {dup_count} duplicatas ignoradas" if dup_count else "")
                + period_str + "."
            )
            db.commit()

            # ── Classificação automática de ofensores ──────────────────────
            if new_count > 0:
                try:
                    from app.services.offender_service import classify_events_batch
                    n_classified = classify_events_batch(db, upload_session_id=session_id)
                    if n_classified:
                        print(f"[OFFENDER] {n_classified} eventos classificados (sessão {session_id}).")
                except Exception as _oe:
                    print(f"[OFFENDER] Erro na classificação automática: {_oe}")
        finally:
            db.close()

    except CSVColumnError as e:
        _update_status(session_id, f"Erro de formato CSV: {e}", status="error")
        db2 = SessionLocal()
        try:
            s = db2.query(UploadSession).get(session_id)
            if s:
                s.error_message = str(e)
                db2.commit()
        finally:
            db2.close()

    except Exception as exc:
        _update_status(session_id, f"Erro: {exc}", status="error")
        db2 = SessionLocal()
        try:
            s = db2.query(UploadSession).get(session_id)
            if s:
                s.error_message = str(exc)
                db2.commit()
        finally:
            db2.close()


def _parse_all_events(path: Path, suffix: str, original_filename: str):
    """Chama o parser correto e define source_filename em todos os eventos."""
    if suffix == ".evtx":
        from app.parser.evtx_parser import parse_evtx_file
        events = parse_evtx_file(str(path))
    elif suffix == ".xml":
        from app.parser.xml_parser import parse_xml_file
        events = parse_xml_file(str(path))
    elif suffix == ".csv":
        from app.parser.csv_parser import parse_csv_file
        events = parse_csv_file(str(path))
    else:
        raise ValueError(f"Formato desconhecido: {suffix}")

    for ev in events:
        ev.source_filename = ev.source_filename or original_filename
    return events


def _insert_events(db, session_id: int, client_id: int, events) -> tuple[int, int]:
    """
    Insere eventos usando INSERT ... ON CONFLICT DO NOTHING para deduplicação
    automática pela constraint UNIQUE(client_id, content_hash).
    Retorna (new_count, dup_count).
    """
    from sqlalchemy.dialects.postgresql import insert as pg_insert

    if not events:
        return 0, 0

    now = datetime.utcnow()
    rows = []
    for ev in events:
        rows.append({
            "upload_session_id":  session_id,
            "client_id":          client_id,
            "content_hash":       ev.content_hash,
            "raw_event_xml":      ev.raw_event_xml,
            "raw_event_json":     ev.raw_event_json,
            "original_message":   ev.original_message,
            "provider_name":      ev.provider_name,
            "channel":            ev.channel,
            "record_id":          ev.record_id,
            "source_filename":    ev.source_filename,
            "imported_at":        now,
            "event_id":           ev.event_id,
            "time_created":       ev.time_created,
            "vbr_hostname":       ev.vbr_hostname,
            "event_category":     ev.event_category,
            "job_name":           ev.job_name,
            "job_name_normalized":ev.job_name_normalized,
            "job_type":           ev.job_type,
            "job_result":         ev.job_result,
            "will_be_retried":    ev.will_be_retried,
            "operator":           ev.operator,
            "audit_event_type":   ev.audit_event_type,
            "audit_event_label":  ev.audit_event_label,
            "audit_details":      ev.audit_details,
            "parser_version":     ev.parser_version,
            "parsed_at":          now,
        })

    # Batch insert em chunks de 500
    total_new = 0
    chunk_size = 500
    for i in range(0, len(rows), chunk_size):
        chunk = rows[i:i + chunk_size]
        stmt = pg_insert(ImportedEvent.__table__).values(chunk)
        stmt = stmt.on_conflict_do_nothing(
            constraint="uq_events_client_hash"
        )
        result = db.execute(stmt)
        total_new += result.rowcount
        db.commit()

    dup_count = len(rows) - total_new
    return total_new, dup_count


# ── Status polling da importação ────────────────────────────────────────────────

@app.get("/import/processing/{session_id}", response_class=HTMLResponse)
def import_processing_page(session_id: int, request: Request, db: Session = Depends(get_db)):
    session = db.query(UploadSession).get(session_id)
    if not session:
        raise HTTPException(status_code=404, detail="Sessão não encontrada")
    if session.status == "done":
        return RedirectResponse(url=f"/import/{session_id}/summary")
    return templates.TemplateResponse(
        "import_processing.html",
        {"request": request, "session": session},
    )




@app.get("/api/import/status/{session_id}")
def api_import_status(session_id: int, db: Session = Depends(get_db)):
    session = db.query(UploadSession).get(session_id)
    if not session:
        raise HTTPException(status_code=404)
    return {
        "status":            session.status,
        "message":           session.status_message,
        "new_inserted":      session.new_events_inserted,
        "duplicates":        session.duplicate_events_skipped,
        "total_read":        session.total_events_read,
    }




@app.get("/import/{session_id}/summary", response_class=HTMLResponse)
def import_summary_page(session_id: int, request: Request, db: Session = Depends(get_db)):
    session = db.query(UploadSession).get(session_id)
    if not session:
        raise HTTPException(status_code=404)
    return templates.TemplateResponse(
        "import_summary.html",
        {"request": request, "session": session, "client": session.client},
    )


# ── Histórico de cargas ─────────────────────────────────────────────────────────

@app.get("/history", response_class=HTMLResponse)
def import_history_page(
    request: Request,
    client_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    q = db.query(UploadSession).order_by(UploadSession.uploaded_at.desc())
    if client_id:
        q = q.filter(UploadSession.client_id == client_id)
    sessions = q.limit(100).all()
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse(
        "import_history.html",
        {
            "request":           request,
            "sessions":          sessions,
            "clients":           clients,
            "selected_client_id": client_id,
        },
    )




@app.post("/history/{session_id}/delete")
def delete_import(session_id: int, db: Session = Depends(get_db)):
    session = db.query(UploadSession).get(session_id)
    if not session:
        raise HTTPException(status_code=404)
    stored = UPLOAD_DIR / session.stored_filename
    stored.unlink(missing_ok=True)
    db.delete(session)
    db.commit()
    return RedirectResponse(url="/history", status_code=303)


# ══════════════════════════════════════════════════════════════════════════════
# RELATÓRIOS POR PERÍODO
# Consulta a base histórica — sem vínculo com upload específico.
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/reports", response_class=HTMLResponse)
def reports_list_page(request: Request, db: Session = Depends(get_db)):
    reports = (
        db.query(ReportSession)
        .order_by(ReportSession.created_at.desc())
        .limit(50)
        .all()
    )
    return templates.TemplateResponse(
        "reports_list.html", {"request": request, "reports": reports}
    )


@app.get("/reports/new", response_class=HTMLResponse)
def new_report_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()

    # Sugerir período mais recente disponível
    latest = (
        db.query(ImportedEvent.time_created)
        .order_by(ImportedEvent.time_created.desc())
        .first()
    )
    default_end   = (latest[0] if latest else datetime.utcnow()).strftime("%Y-%m-%d")
    default_start = (
        (latest[0] - timedelta(days=6)) if latest else (datetime.utcnow() - timedelta(days=6))
    ).strftime("%Y-%m-%d")

    return templates.TemplateResponse(
        "report_query.html",
        {
            "request":       request,
            "clients":       clients,
            "default_start": default_start,
            "default_end":   default_end,
        },
    )


@app.post("/reports/new")
async def create_report(
    request: Request,
    client_id:          int  = Form(...),
    date_start:         str  = Form(...),
    date_end:           str  = Form(...),
    include_backup:     bool = Form(False),
    include_log_backup: bool = Form(False),
    show_all_jobs:      bool = Form(False),
    db: Session = Depends(get_db),
):
    client = db.query(Client).get(client_id)
    if not client:
        raise HTTPException(status_code=400, detail="Cliente não encontrado")

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    if not dt_start or not dt_end:
        raise HTTPException(status_code=400, detail="Datas inválidas")

    if dt_start > dt_end:
        raise HTTPException(status_code=400, detail="Data inicial deve ser anterior à data final")

    if not include_backup and not include_log_backup:
        include_backup = include_log_backup = True

    created_by = request.session.get("username")
    report = build_report(
        db=db,
        client_id=client_id,
        client_name=client.name,
        date_start=dt_start,
        date_end=dt_end,
        include_backup=include_backup,
        include_log_backup=include_log_backup,
        show_all_jobs=show_all_jobs,
        created_by=created_by,
    )

    return RedirectResponse(url=f"/review/{report.id}", status_code=303)


# ── Review e edição do relatório ────────────────────────────────────────────────

@app.get("/review/{report_id}", response_class=HTMLResponse)
def review_page(report_id: int, request: Request, db: Session = Depends(get_db)):
    report = _get_report_or_404(db, report_id)
    summaries = (
        db.query(JobSummary)
        .filter(JobSummary.report_id == report_id)
        .order_by(JobSummary.failed_count.desc(), JobSummary.job_name)
        .all()
    )
    week_label = format_week_label(report.date_start, report.date_end)

    # Totais do período
    period_events = (
        db.query(ImportedEvent)
        .filter(
            ImportedEvent.client_id == report.client_id,
            ImportedEvent.time_created >= report.date_start,
            ImportedEvent.time_created <= report.date_end,
            ImportedEvent.event_category.in_(["backup", "log_backup"]),
        )
        .all()
    )
    totals = {
        "success": sum(1 for e in period_events if e.job_result == 0),
        "warning": sum(1 for e in period_events if e.job_result == 1),
        "failed":  sum(1 for e in period_events if e.job_result == 2),
        "total":   len(period_events),
        "backup_total":       sum(1 for e in period_events if e.event_category == "backup"),
        "log_backup_total":   sum(1 for e in period_events if e.event_category == "log_backup"),
        "backup_success":     sum(1 for e in period_events if e.job_result == 0 and e.event_category == "backup"),
        "log_backup_success": sum(1 for e in period_events if e.job_result == 0 and e.event_category == "log_backup"),
    }

    return templates.TemplateResponse(
        "review.html",
        {
            "request":      request,
            "report":       report,
            "summaries":    summaries,
            "week_label":   week_label,
            "total_jobs":   len(summaries),
            "total_errors": sum(s.failed_count for s in summaries),
            "totals":       totals,
        },
    )


@app.post("/review/{report_id}/save")
async def save_review(report_id: int, request: Request, db: Session = Depends(get_db)):
    report = _get_report_or_404(db, report_id)
    form   = await request.form()
    actor  = request.session.get("username")
    now    = datetime.utcnow()

    summaries = db.query(JobSummary).filter(JobSummary.report_id == report_id).all()
    for js in summaries:
        prefix = f"job_{js.id}_"
        js.situation              = form.get(f"{prefix}situation",              js.situation)
        js.action_ongoing         = form.get(f"{prefix}action_ongoing",         "")
        js.technical_observation  = form.get(f"{prefix}technical_observation",  "")
        js.responsible            = form.get(f"{prefix}responsible",            "")
        js.ticket                 = form.get(f"{prefix}ticket",                 "")

        # Persistir em JobAction para pré-popular relatórios futuros
        job_norm = js.job_name_normalized or js.job_name
        ja = (
            db.query(JobAction)
            .filter(
                JobAction.client_id == report.client_id,
                JobAction.job_name_normalized == job_norm,
            )
            .first()
        )
        if ja is None:
            ja = JobAction(client_id=report.client_id, job_name_normalized=job_norm)
            db.add(ja)

        ja.situation              = js.situation
        ja.action_ongoing         = js.action_ongoing
        ja.technical_observation  = js.technical_observation
        ja.responsible            = js.responsible
        ja.ticket                 = js.ticket
        ja.updated_at             = now
        ja.updated_by             = actor

    report.client_name = form.get("client_name", report.client_name)
    db.commit()
    return RedirectResponse(url=f"/review/{report_id}?saved=1", status_code=303)


# ── Exportação Excel ────────────────────────────────────────────────────────────

@app.get("/review/{report_id}/export")
def export_excel(report_id: int, db: Session = Depends(get_db)):
    report = _get_report_or_404(db, report_id)
    summaries = (
        db.query(JobSummary)
        .filter(JobSummary.report_id == report_id)
        .order_by(JobSummary.failed_count.desc(), JobSummary.job_name)
        .all()
    )
    all_events = (
        db.query(ImportedEvent)
        .filter(
            ImportedEvent.client_id == report.client_id,
            ImportedEvent.time_created >= report.date_start,
            ImportedEvent.time_created <= report.date_end,
            ImportedEvent.event_category.in_(["backup", "log_backup"]),
        )
        .order_by(ImportedEvent.time_created)
        .all()
    )

    filename = (
        f"veeam_report_{report.date_start.strftime('%Y%m%d')}"
        f"_{report.date_end.strftime('%Y%m%d')}.xlsx"
    )
    output_path = REPORTS_DIR / filename
    _excel_branding = branding_service.get_branding(client_id=report.client_id)
    generate_excel(report, summaries, all_events, str(output_path), branding=_excel_branding)

    return FileResponse(
        path=str(output_path),
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ══════════════════════════════════════════════════════════════════════════════
# REPROCESSAMENTO DE EVENTOS
# Recalcula campos derivados dos eventos brutos sem necessidade de re-upload.
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/admin/reprocess", response_class=HTMLResponse)
async def reprocess_page(request: Request, db: Session = Depends(get_db)):
    from app.auth import require_admin
    require_admin(request)
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    total_events = db.query(ImportedEvent).count()
    return templates.TemplateResponse(
        "admin_reprocess.html",
        {"request": request, "clients": clients, "total_events": total_events},
    )


@app.post("/api/admin/reprocess")
async def reprocess_events(
    request: Request,
    db: Session = Depends(get_db),
):
    from app.auth import require_admin
    require_admin(request)

    body = await request.json()
    client_id = body.get("client_id")

    q = db.query(ImportedEvent)
    if client_id:
        q = q.filter(ImportedEvent.client_id == client_id)

    events = q.all()
    updated = 0

    for ev in events:
        changed = _reprocess_event(ev)
        if changed:
            updated += 1

    db.commit()
    return JSONResponse({"reprocessed": len(events), "updated": updated})


def _reprocess_event(ev: ImportedEvent) -> bool:
    """
    Recalcula campos derivados a partir dos dados brutos.
    Retorna True se algum campo foi alterado.
    """
    from app.parser.xml_parser import parse_xml_file
    from app.parser.common import (
        AUDIT_EVENT_IDS, AUDIT_EVENT_MAP, VEEAM_JOB_EVENT_IDS,
        extract_job_name_from_message, extract_result_from_message,
        is_log_backup, normalize_job_name,
    )
    import xml.etree.ElementTree as ET

    changed = False

    if ev.raw_event_xml:
        try:
            elem = ET.fromstring(ev.raw_event_xml)
            from app.parser.xml_parser import _get_event_id, _parse_backup_element, _parse_audit_element
            eid = _get_event_id(elem)
            if eid in AUDIT_EVENT_IDS:
                parsed = _parse_audit_element(elem, ev.raw_event_xml)
            elif eid in VEEAM_JOB_EVENT_IDS:
                parsed = _parse_backup_element(elem, ev.raw_event_xml)
            else:
                parsed = None

            if parsed:
                for field in (
                    "event_category", "job_name", "job_name_normalized", "job_type",
                    "job_result", "will_be_retried", "operator",
                    "audit_event_type", "audit_event_label", "audit_details",
                ):
                    new_val = getattr(parsed, field, None)
                    if getattr(ev, field) != new_val:
                        setattr(ev, field, new_val)
                        changed = True
        except Exception:
            pass
    elif ev.raw_event_json:
        # CSV: recalcular campos derivados a partir do JSON armazenado
        try:
            import json
            row = json.loads(ev.raw_event_json)
            message = next(
                (v for k, v in row.items() if k.lower() in ("message", "general", "description")),
                ev.original_message or "",
            )
            if message:
                new_job_name = extract_job_name_from_message(message)
                new_result   = extract_result_from_message(message)
                if ev.job_name != new_job_name:
                    ev.job_name = new_job_name
                    changed = True
                if ev.job_result != new_result:
                    ev.job_result = new_result
                    changed = True
            new_norm = normalize_job_name(ev.job_name)
            if ev.job_name_normalized != new_norm:
                ev.job_name_normalized = new_norm
                changed = True
        except Exception:
            pass

    if changed:
        from datetime import datetime
        from app.parser.common import PARSER_VERSION
        ev.parsed_at      = datetime.utcnow()
        ev.parser_version = PARSER_VERSION

    return changed


# ══════════════════════════════════════════════════════════════════════════════
# DASHBOARDS
# Todos os dashboards consultam ImportedEvent diretamente por cliente + período.
# ══════════════════════════════════════════════════════════════════════════════

def _parse_date(s: Optional[str], end_of_day: bool = False) -> Optional[datetime]:
    try:
        dt = datetime.strptime(s, "%Y-%m-%d")
        return dt.replace(hour=23, minute=59, second=59) if end_of_day else dt
    except (ValueError, TypeError):
        return None


# ── Isolamento por cliente (multi-tenant) ───────────────────────────────────
def _scoped_client_ids(db: Session, request: Optional[Request]):
    """None = admin (todos os clientes). Caso contrário, set de client_ids do escopo."""
    if request is not None and request.session.get("is_admin"):
        return None
    uid = request.session.get("user_id") if request is not None else None
    if not uid:
        return set()
    return {r[0] for r in db.query(UserClient.client_id).filter(UserClient.user_id == uid).all()}


def scope_guard(
    request: Request,
    client_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
) -> None:
    """Dependency: impede não-admin de consultar client_id fora do seu escopo (403)."""
    if request.session.get("is_admin"):
        return
    allowed = _scoped_client_ids(db, request)
    if client_id is None:
        raise HTTPException(status_code=403, detail="Selecione um cliente do seu escopo.")
    if client_id not in allowed:
        raise HTTPException(status_code=403, detail="Cliente fora do seu escopo.")


def _all_clients_list(db: Session, request: Optional[Request] = None) -> list:
    allowed = _scoped_client_ids(db, request) if request is not None else None
    q = db.query(Client).filter(Client.is_active.is_(True))
    if allowed is not None:
        if not allowed:
            return []
        q = q.filter(Client.id.in_(allowed))
    return [{"id": c.id, "name": c.name} for c in q.order_by(Client.name).all()]


def _all_reports_list(db: Session, client_id: Optional[int] = None) -> list:
    q = db.query(ReportSession)
    if client_id:
        q = q.filter(ReportSession.client_id == client_id)
    rows = q.order_by(ReportSession.date_start.desc()).limit(50).all()
    return [
        {
            "id":         r.id,
            "label":      (
                f"{r.client_name or 'Sem cliente'} — "
                f"{r.date_start.strftime('%d/%m/%Y')} → {r.date_end.strftime('%d/%m/%Y')}"
            ),
            "date_start": r.date_start.strftime("%Y-%m-%d"),
            "date_end":   r.date_end.strftime("%Y-%m-%d"),
            "client_id":  r.client_id,
        }
        for r in rows
    ]


def _build_event_query(db, dt_start, dt_end, job_type, client_id=None):
    """Query de ImportedEvent para dashboards de backup/log_backup."""
    q = (
        db.query(ImportedEvent)
        .filter(ImportedEvent.event_category.in_(["backup", "log_backup"]))
        # Excluir EV150 de sub-tarefas de VM
        .filter(
            ~(
                (ImportedEvent.event_id == 150)
                & (ImportedEvent.event_category == "backup")
            )
        )
    )
    if client_id:
        q = q.filter(ImportedEvent.client_id == client_id)
    if dt_start:
        q = q.filter(ImportedEvent.time_created >= dt_start)
    if dt_end:
        q = q.filter(ImportedEvent.time_created <= dt_end)
    if job_type and job_type != "all":
        q = q.filter(ImportedEvent.event_category == job_type)
    return q


@app.get("/dashboard", response_class=HTMLResponse)
def dashboard_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse("dashboard.html", {"request": request, "clients": clients})


@app.get("/api/dashboard", dependencies=[Depends(scope_guard)])
def api_dashboard(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    job_type:   str            = Query("all"),
    top_n:      int            = Query(10, ge=1, le=50),
    client_id:  Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final para consultar.", "code": "dates_required"},
            status_code=400,
        )

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    q      = _build_event_query(db, dt_start, dt_end, job_type, client_id)
    events = q.order_by(ImportedEvent.time_created).all()

    total       = len(events)
    success     = sum(1 for e in events if e.job_result == 0)
    warning     = sum(1 for e in events if e.job_result == 1)
    failed      = sum(1 for e in events if e.job_result == 2)
    unique_jobs = len({e.job_name for e in events if e.job_name and e.event_category == "backup"})
    success_rate = round(success / total * 100, 1) if total else 0

    by_day: dict = defaultdict(lambda: {"success": 0, "warning": 0, "failed": 0})
    for e in events:
        if e.time_created:
            key = e.time_created.strftime("%Y-%m-%d")
            if   e.job_result == 0: by_day[key]["success"] += 1
            elif e.job_result == 1: by_day[key]["warning"] += 1
            elif e.job_result == 2: by_day[key]["failed"]  += 1

    day_keys   = sorted(by_day.keys())
    day_labels = [datetime.strptime(d, "%Y-%m-%d").strftime("%d/%m") for d in day_keys]

    job_stats: dict = defaultdict(
        lambda: {"failed": 0, "warning": 0, "success": 0, "total": 0, "type": "backup"}
    )
    for e in events:
        if not e.job_name:
            continue
        s = job_stats[e.job_name]
        s["total"] += 1
        s["type"]   = e.event_category or "backup"
        if   e.job_result == 2: s["failed"]  += 1
        elif e.job_result == 1: s["warning"] += 1
        else:                   s["success"] += 1

    n        = top_n if top_n and top_n > 0 else len(job_stats)
    by_errors = sorted(job_stats.items(), key=lambda x: x[1]["failed"], reverse=True)[:n]
    by_execs  = sorted(job_stats.items(), key=lambda x: x[1]["total"],  reverse=True)[:n]

    def _trunc(name: str, max_len: int = 35) -> str:
        return name[:max_len] + "…" if len(name) > max_len else name

    return JSONResponse({
        "kpis": {
            "total": total, "success": success, "warning": warning,
            "failed": failed, "success_rate": success_rate, "unique_jobs": unique_jobs,
        },
        "result_distribution": {
            "labels": ["Sucesso", "Warning", "Falha"],
            "values": [success, warning, failed],
        },
        "errors_by_day": {
            "labels":  day_labels,
            "success": [by_day[d]["success"] for d in day_keys],
            "warning": [by_day[d]["warning"] for d in day_keys],
            "failed":  [by_day[d]["failed"]  for d in day_keys],
        },
        "top_by_errors": {
            "labels":  [_trunc(i[0]) for i in by_errors],
            "failed":  [i[1]["failed"]  for i in by_errors],
            "warning": [i[1]["warning"] for i in by_errors],
            "total":   [i[1]["total"]   for i in by_errors],
            "types":   [i[1]["type"]    for i in by_errors],
        },
        "top_by_executions": {
            "labels":  [_trunc(i[0]) for i in by_execs],
            "total":   [i[1]["total"]   for i in by_execs],
            "success": [i[1]["success"] for i in by_execs],
            "warning": [i[1]["warning"] for i in by_execs],
            "failed":  [i[1]["failed"]  for i in by_execs],
            "types":   [i[1]["type"]    for i in by_execs],
        },
        "reports": _all_reports_list(db, client_id=client_id),
        "clients": _all_clients_list(db),
    })


@app.get("/dashboard/admin", response_class=HTMLResponse)
def dashboard_admin_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse("dashboard_admin.html", {"request": request, "clients": clients})


@app.get("/api/admin-dashboard", dependencies=[Depends(require_admin)])
def api_admin_dashboard(
    date_start:      Optional[str] = Query(None),
    date_end:        Optional[str] = Query(None),
    job_type:        str            = Query("all"),
    client_id:       Optional[int] = Query(None),
    failed_page:     int            = Query(1, ge=1),
    failed_page_size:int            = Query(10, ge=1, le=100),
    db: Session = Depends(get_db),
):
    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final para consultar.", "code": "dates_required"},
            status_code=400,
        )

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    q      = _build_event_query(db, dt_start, dt_end, job_type, client_id)
    events = q.order_by(ImportedEvent.time_created.desc()).all()

    last_per_job: dict = {}
    for e in events:
        if e.job_name and e.job_name not in last_per_job:
            last_per_job[e.job_name] = e

    RESULT_LABEL = {0: "Sucesso", 1: "Warning", 2: "Falha"}
    failed_jobs_all  = []
    success_jobs = []

    for job_name, e in last_per_job.items():
        entry = {
            "job_name":       job_name,
            "job_type":       e.event_category or "backup",
            "last_execution": e.time_created.strftime("%d/%m/%Y %H:%M") if e.time_created else "—",
            "result":         e.job_result,
            "result_label":   RESULT_LABEL.get(e.job_result, "—"),
        }
        (failed_jobs_all if e.job_result == 2 else success_jobs).append(entry)

    failed_jobs_all.sort(key=lambda x: x["job_name"])
    success_jobs.sort(key=lambda x: x["job_name"])
    total = len(last_per_job)

    # Paginação server-side para jobs com falha
    total_failed = len(failed_jobs_all)
    page_size    = max(1, failed_page_size)
    total_pages  = max(1, (total_failed + page_size - 1) // page_size)
    page         = max(1, min(failed_page, total_pages))
    offset       = (page - 1) * page_size
    failed_jobs  = failed_jobs_all[offset : offset + page_size]

    return JSONResponse({
        "kpis": {
            "total":       total,
            "failed":      total_failed,
            "success":     len(success_jobs),
            "failed_pct":  round(total_failed      / total * 100, 1) if total else 0,
            "success_pct": round(len(success_jobs) / total * 100, 1) if total else 0,
        },
        "failed_jobs":        failed_jobs,
        "failed_pagination":  {"page": page, "page_size": page_size, "total": total_failed, "total_pages": total_pages},
        "success_jobs":       success_jobs,
        "reports":            _all_reports_list(db, client_id=client_id),
        "clients":            _all_clients_list(db),
    })


@app.get("/dashboard/protection", response_class=HTMLResponse)
def dashboard_protection_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse("dashboard_protection.html", {"request": request, "clients": clients})


@app.get("/api/protection-dashboard", dependencies=[Depends(scope_guard)])
def api_protection_dashboard(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    client_id:  Optional[int] = Query(None),
    period:     str            = Query("fortnight"),
    db: Session = Depends(get_db),
):
    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final para consultar.", "code": "dates_required"},
            status_code=400,
        )

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    def _q(eid):
        q = db.query(ImportedEvent).filter(ImportedEvent.event_id == eid)
        if client_id:
            q = q.filter(ImportedEvent.client_id == client_id)
        if dt_start:
            q = q.filter(ImportedEvent.time_created >= dt_start)
        if dt_end:
            q = q.filter(ImportedEvent.time_created <= dt_end)
        return q.all()

    ev150 = _q(150)
    ev190 = _q(190)

    protected_vms = len({e.job_name for e in ev150 if e.job_name})
    log_vms       = len({e.job_name for e in ev150 if e.job_name and e.event_category == "log_backup"})
    backup_jobs   = len({e.job_name for e in ev190 if e.job_name})

    dates = [e.time_created for e in ev150 if e.time_created]
    evolution = []
    if dates:
        range_start = (dt_start or min(dates)).replace(hour=0, minute=0, second=0, microsecond=0)
        range_end   = (dt_end   or max(dates)).replace(hour=23, minute=59, second=59, microsecond=0)
        delta_days  = {"week": 7, "month": 30}.get(period, 14)
        cursor = range_start
        while cursor <= range_end:
            p_end        = min(cursor + timedelta(days=delta_days - 1), range_end)
            p_end        = p_end.replace(hour=23, minute=59, second=59)
            period_evs   = [e for e in ev150 if e.time_created and cursor <= e.time_created <= p_end]
            unique_vms   = len({e.job_name for e in period_evs if e.job_name})
            unique_log   = len({e.job_name for e in period_evs if e.job_name and e.event_category == "log_backup"})
            evolution.append({
                "label":   f"{cursor.strftime('%d/%m')} – {p_end.strftime('%d/%m/%y')}",
                "vms":     unique_vms,
                "log_vms": unique_log,
                "bkp_vms": unique_vms - unique_log,
            })
            cursor += timedelta(days=delta_days)

    return JSONResponse({
        "kpis": {"protected_vms": protected_vms, "log_vms": log_vms, "backup_jobs": backup_jobs},
        "evolution": evolution,
        "clients":   _all_clients_list(db),
        "reports":   _all_reports_list(db, client_id=client_id),
    })


# ─────────────────────────────────────────────────────────────────────────────
# VM List: lista auditável de VMs protegidas no período  (EV150)
# ─────────────────────────────────────────────────────────────────────────────

def _vm_list_base_sql(dt_start, dt_end, client_id, search, protection_type, result_filter):
    """
    Monta SQL e parâmetros para a lista de VMs protegidas.
    Usa DISTINCT ON (PostgreSQL) para buscar o último resultado de cada VM
    sem subquery correlacionada.  Retorna (data_cte_sql, count_sql, params).
    """
    conds = [
        "e.event_id = 150",
        "e.job_name IS NOT NULL",
        "e.time_created BETWEEN :dt_start AND :dt_end",
    ]
    params: dict = {"dt_start": dt_start, "dt_end": dt_end}

    if client_id:
        conds.append("e.client_id = :client_id")
        params["client_id"] = client_id

    where = " AND ".join(conds)

    # CTE: agrega por (job_name, client_id)
    # CTE2: DISTINCT ON para pegar último resultado
    cte = f"""
    WITH vm_agg AS (
        SELECT
            e.job_name,
            e.client_id,
            COUNT(*)                                                   AS exec_count,
            MAX(e.time_created)                                        AS last_backup,
            MAX(CASE WHEN e.event_category = 'log_backup' THEN 1 ELSE 0 END) AS has_log,
            MAX(CASE WHEN e.event_category = 'backup'     THEN 1 ELSE 0 END) AS has_backup
        FROM imported_events e
        WHERE {where}
        GROUP BY e.job_name, e.client_id
    ),
    vm_last AS (
        SELECT DISTINCT ON (e.job_name, e.client_id)
            e.job_name,
            e.client_id,
            e.job_result,
            e.vbr_hostname
        FROM imported_events e
        WHERE {where}
        ORDER BY e.job_name, e.client_id, e.time_created DESC
    )
    SELECT
        a.job_name      AS vm_name,
        a.client_id,
        c.name          AS client_name,
        a.exec_count,
        a.last_backup,
        a.has_log,
        a.has_backup,
        l.job_result,
        l.vbr_hostname
    FROM vm_agg  a
    JOIN vm_last l ON l.job_name = a.job_name AND l.client_id = a.client_id
    JOIN clients c ON c.id = a.client_id
    WHERE 1=1
    """

    # Filtro de busca
    if search:
        cte += " AND LOWER(a.job_name) LIKE :search"
        params["search"] = f"%{search.lower()}%"

    # Filtro por tipo de proteção
    if protection_type == "with_log":
        cte += " AND a.has_log = 1"
    elif protection_type == "without_log":
        cte += " AND a.has_log = 0"

    # Filtro por resultado
    if result_filter == "success":
        cte += " AND l.job_result = 0"
    elif result_filter == "warning":
        cte += " AND l.job_result = 1"
    elif result_filter == "failed":
        cte += " AND l.job_result = 2"

    count_sql = f"SELECT COUNT(*) FROM ({cte}) _cnt"
    return cte, count_sql, params


@app.get("/api/protection-vm-list", dependencies=[Depends(scope_guard)])
def api_protection_vm_list(
    date_start:      Optional[str] = Query(None),
    date_end:        Optional[str] = Query(None),
    client_id:       Optional[int] = Query(None),
    search:          Optional[str] = Query(None),
    protection_type: str            = Query("all"),   # all | with_log | without_log
    result_filter:   str            = Query("all"),   # all | success | warning | failed
    page:            int            = Query(1, ge=1),
    page_size:       int            = Query(20, ge=1, le=200),
    db: Session = Depends(get_db),
):
    from sqlalchemy import text as sa_text

    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final.", "code": "dates_required"},
            status_code=400,
        )

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    data_sql, count_sql, params = _vm_list_base_sql(
        dt_start, dt_end, client_id, search, protection_type, result_filter
    )

    total = db.execute(sa_text(count_sql), params).scalar() or 0

    page_size   = max(1, min(page_size, 500))
    total_pages = max(1, (total + page_size - 1) // page_size)
    page        = max(1, min(page, total_pages))
    offset      = (page - 1) * page_size

    paged_sql = data_sql + " ORDER BY a.job_name, a.client_id LIMIT :limit OFFSET :offset"
    params["limit"]  = page_size
    params["offset"] = offset

    rows = db.execute(sa_text(paged_sql), params).fetchall()

    RESULT_LABEL = {0: "Sucesso", 1: "Warning", 2: "Falha"}

    vms = []
    for r in rows:
        has_log    = bool(r.has_log)
        has_backup = bool(r.has_backup)
        if has_log and has_backup:
            prot_type = "Backup + Trans. Log"
        elif has_log:
            prot_type = "Transaction Log"
        else:
            prot_type = "Backup"

        vms.append({
            "vm_name":           r.vm_name,
            "client_id":         r.client_id,
            "client_name":       r.client_name,
            "exec_count":        r.exec_count,
            "last_backup":       r.last_backup.strftime("%d/%m/%Y %H:%M") if r.last_backup else "—",
            "has_log":           has_log,
            "protection_type":   prot_type,
            "last_result":       r.job_result,
            "last_result_label": RESULT_LABEL.get(r.job_result, "—"),
            "vbr_hostname":      r.vbr_hostname or "—",
        })

    return JSONResponse({
        "total":       total,
        "page":        page,
        "page_size":   page_size,
        "total_pages": total_pages,
        "vms":         vms,
    })


@app.get("/api/protection-vm-detail", dependencies=[Depends(scope_guard)])
def api_protection_vm_detail(
    vm_name:    str            = Query(...),
    client_id:  Optional[int] = Query(None),
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    """Histórico de execuções de uma VM específica (drill-down)."""
    q = (
        db.query(ImportedEvent)
        .filter(
            ImportedEvent.event_id == 150,
            ImportedEvent.job_name == vm_name,
        )
        .order_by(ImportedEvent.time_created.desc())
    )
    if client_id:
        q = q.filter(ImportedEvent.client_id == client_id)
    if date_start:
        q = q.filter(ImportedEvent.time_created >= _parse_date(date_start))
    if date_end:
        q = q.filter(ImportedEvent.time_created <= _parse_date(date_end, end_of_day=True))

    rows = q.limit(500).all()

    RESULT_LABEL = {0: "Sucesso", 1: "Warning", 2: "Falha"}
    TYPE_LABEL   = {"backup": "Backup", "log_backup": "Transaction Log"}

    return JSONResponse({
        "executions": [
            {
                "id":               e.id,
                "time":             e.time_created.strftime("%d/%m/%Y %H:%M") if e.time_created else "—",
                "event_id":         e.event_id,
                "event_category":   e.event_category,
                "type_label":       TYPE_LABEL.get(e.event_category or "", e.event_category or "—"),
                "result":           e.job_result,
                "result_label":     RESULT_LABEL.get(e.job_result, "—"),
                "vbr_hostname":     e.vbr_hostname or "—",
                "message":          (e.original_message or "")[:400],
                "upload_session_id": e.upload_session_id,
                "source_filename":  e.source_filename or "—",
            }
            for e in rows
        ]
    })


@app.get("/api/export/protection-vm-excel", dependencies=[Depends(scope_guard)])
def export_protection_vm_excel(
    date_start:      Optional[str] = Query(None),
    date_end:        Optional[str] = Query(None),
    client_id:       Optional[int] = Query(None),
    search:          Optional[str] = Query(None),
    protection_type: str            = Query("all"),
    result_filter:   str            = Query("all"),
    db: Session = Depends(get_db),
):
    from sqlalchemy import text as sa_text
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if not date_start or not date_end:
        return JSONResponse({"error": "Informe data inicial e data final."}, status_code=400)

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    data_sql, _, params = _vm_list_base_sql(
        dt_start, dt_end, client_id, search, protection_type, result_filter
    )
    full_sql = data_sql + " ORDER BY a.job_name, a.client_id LIMIT 10000"
    rows = db.execute(sa_text(full_sql), params).fetchall()

    RESULT_LABEL = {0: "Sucesso", 1: "Warning", 2: "Falha"}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VMs Protegidas"

    # Header style
    hdr_fill  = PatternFill("solid", fgColor="7C3AED")
    hdr_font  = Font(bold=True, color="FFFFFF", size=10)
    hdr_align = Alignment(horizontal="center", vertical="center")
    thin      = Side(style="thin", color="CCCCCC")
    border    = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [
        "Máquina Virtual", "Ambiente / Cliente", "Tipo de Proteção",
        "Último Backup", "Último Resultado", "Servidor VBR",
        "Com Transaction Log", "Qtd. Execuções no Período",
    ]
    widths = [45, 30, 25, 20, 15, 25, 20, 22]

    ws.row_dimensions[1].height = 20
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font      = hdr_font
        cell.fill      = hdr_fill
        cell.alignment = hdr_align
        cell.border    = border
        ws.column_dimensions[get_column_letter(col)].width = w

    RESULT_COLOR = {0: "D4EDDA", 1: "FFF3CD", 2: "F8D7DA"}

    for row_idx, r in enumerate(rows, 2):
        has_log    = bool(r.has_log)
        has_backup = bool(r.has_backup)
        if has_log and has_backup:
            prot_type = "Backup + Trans. Log"
        elif has_log:
            prot_type = "Transaction Log"
        else:
            prot_type = "Backup"

        result_label = RESULT_LABEL.get(r.job_result, "—")
        last_backup  = r.last_backup.strftime("%d/%m/%Y %H:%M") if r.last_backup else "—"
        result_color = RESULT_COLOR.get(r.job_result, "FFFFFF")

        row_data = [
            r.vm_name,
            r.client_name,
            prot_type,
            last_backup,
            result_label,
            r.vbr_hostname or "—",
            "Sim" if has_log else "Não",
            r.exec_count,
        ]

        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col, value=val)
            cell.border    = border
            cell.alignment = Alignment(vertical="center")
            cell.font      = Font(size=9)
            # Color result column
            if col == 5:
                cell.fill = PatternFill("solid", fgColor=result_color)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            if col == 8:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Auto-filter
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.freeze_panes    = "A2"

    # Metadata
    ws2 = wb.create_sheet("Informações")
    ws2["A1"] = "Gerado em"
    ws2["B1"] = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws2["A2"] = "Período"
    ws2["B2"] = f"{dt_start.strftime('%d/%m/%Y')} → {dt_end.strftime('%d/%m/%Y')}"
    ws2["A3"] = "Total de VMs"
    ws2["B3"] = len(rows)

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(tmp.name)
    tmp.close()

    period_label = f"{date_start}_a_{date_end}"
    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"vms_protegidas_{period_label}.xlsx",
    )


@app.get("/api/export/protection-pptx", dependencies=[Depends(scope_guard)])
def export_protection_pptx(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    client_id:  Optional[int] = Query(None),
    period:     str            = Query("fortnight"),
    db: Session = Depends(get_db),
):
    """Gera o PPTX de Proteção de VMs com capa dinâmica e identidade visual ativa."""
    from fastapi.responses import Response as FastAPIResponse
    from app.services.pptx_service import generate_protection_pptx
    from sqlalchemy import text as sa_text

    if not date_start or not date_end:
        raise HTTPException(status_code=400, detail="Informe data inicial e data final.")

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    # ── KPIs de proteção ──────────────────────────────────────────────────────
    def _q(eid):
        q = db.query(ImportedEvent).filter(ImportedEvent.event_id == eid)
        if client_id:
            q = q.filter(ImportedEvent.client_id == client_id)
        if dt_start:
            q = q.filter(ImportedEvent.time_created >= dt_start)
        if dt_end:
            q = q.filter(ImportedEvent.time_created <= dt_end)
        return q.all()

    ev150 = _q(150)
    ev190 = _q(190)

    protected_vms = len({e.job_name for e in ev150 if e.job_name})
    log_vms       = len({e.job_name for e in ev150 if e.job_name and e.event_category == "log_backup"})
    backup_jobs   = len({e.job_name for e in ev190 if e.job_name})
    kpis = {"protected_vms": protected_vms, "log_vms": log_vms, "backup_jobs": backup_jobs}

    # ── Evolução por período (mesma lógica do dashboard) ─────────────────────
    dates = [e.time_created for e in ev150 if e.time_created]
    evolution: list = []
    if dates:
        range_start = (dt_start or min(dates)).replace(hour=0, minute=0, second=0, microsecond=0)
        range_end   = (dt_end   or max(dates)).replace(hour=23, minute=59, second=59, microsecond=0)
        delta_days  = {"week": 7, "month": 30}.get(period, 14)
        cursor = range_start
        while cursor <= range_end:
            p_end      = min(cursor + timedelta(days=delta_days - 1), range_end)
            p_end      = p_end.replace(hour=23, minute=59, second=59)
            period_evs = [e for e in ev150 if e.time_created and cursor <= e.time_created <= p_end]
            unique_vms = len({e.job_name for e in period_evs if e.job_name})
            unique_log = len({e.job_name for e in period_evs if e.job_name and e.event_category == "log_backup"})
            evolution.append({
                "label":   f"{cursor.strftime('%d/%m')} – {p_end.strftime('%d/%m/%y')}",
                "vms":     unique_vms,
                "log_vms": unique_log,
                "bkp_vms": unique_vms - unique_log,
            })
            cursor += timedelta(days=delta_days)

    # ── Lista de VMs (até 50 para caber no slide; detalhes ficam no Excel) ────
    data_sql, _, params = _vm_list_base_sql(
        dt_start, dt_end, client_id,
        search=None, protection_type="all", result_filter="all",
    )
    vm_rows = db.execute(
        sa_text(data_sql + " ORDER BY a.job_name LIMIT 50"), params
    ).fetchall()
    vm_list = [dict(r._mapping) for r in vm_rows]

    # ── Período legível ───────────────────────────────────────────────────────
    if dt_start and dt_end:
        period_str = f"{dt_start.strftime('%d/%m/%Y')} → {dt_end.strftime('%d/%m/%Y')}"
    elif dates:
        period_str = f"{min(dates).strftime('%d/%m/%Y')} → {max(dates).strftime('%d/%m/%Y')}"
    else:
        period_str = "—"

    # ── Nome do cliente ───────────────────────────────────────────────────────
    client_name = ""
    if client_id:
        c = db.query(Client).get(client_id)
        if c:
            client_name = c.name

    # ── Gerar PPTX ────────────────────────────────────────────────────────────
    _branding  = branding_service.get_branding(client_id=client_id)
    pptx_bytes = generate_protection_pptx(
        client_name=client_name,
        period=period_str,
        kpis=kpis,
        evolution=evolution,
        vm_list=vm_list,
        branding=_branding,
        branding_dir=str(BRANDING_DIR),
    )

    safe_name = (client_name or "todos").replace(" ", "_")
    fname     = f"Protecao_VMs_{safe_name}_{date_start}_{date_end}.pptx"
    return FastAPIResponse(
        content=pptx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


# ══════════════════════════════════════════════════════════════════════════════
# RPO OBSERVADO POR VM — com suporte a RPO esperado por VM
# ══════════════════════════════════════════════════════════════════════════════

# Global default: VMs sem política configurada usam este valor
DEFAULT_RPO_MINUTES = 1440  # 24h


def _rpo_status(backup_count: int, avg_rpo_h, expected_h: float = None) -> str:
    """
    Retorna o status do RPO comparando observado x esperado.
    Se expected_h não for fornecido, usa DEFAULT_RPO_MINUTES como referência.
    """
    if backup_count == 0:
        return "no_backup"
    if backup_count == 1:
        return "no_history"
    if avg_rpo_h is None:
        return "no_history"
    eh = float(expected_h) if expected_h is not None else DEFAULT_RPO_MINUTES / 60.0
    h  = float(avg_rpo_h)
    if h <= eh:
        return "ok"
    if h <= eh * 1.5:
        return "attention"
    return "critical"


def _calc_deviation_pct(avg_rpo_h, expected_h) -> Optional[float]:
    """Desvio percentual: (observado - esperado) / esperado * 100."""
    if avg_rpo_h is None or not expected_h:
        return None
    return round((float(avg_rpo_h) - float(expected_h)) / float(expected_h) * 100, 1)


def _rpo_base_sql(dt_start, dt_end, client_id, search, rpo_status,
                   policy_id: Optional[int] = None):
    """
    Constrói SQL com CTEs para cálculo de RPO observado + esperado por VM.

    Inclui LEFT JOIN em vm_rpo_assignments e rpo_policies para obter
    o RPO esperado configurado por VM. Fallback global: DEFAULT_RPO_MINUTES.

    Colunas do SELECT final:
      vm_name, client_id, client_name, backup_count, last_backup,
      avg_rpo_h, max_rpo_h, min_rpo_h,
      assignment_id, expected_minutes, criticality,
      policy_id, policy_name, expected_h, rpo_source
    """
    params: dict = {"dt_start": dt_start, "dt_end": dt_end}
    c_cond = ""
    if client_id:
        c_cond = "AND e.client_id = :client_id"
        params["client_id"] = client_id

    cte = f"""
    WITH
    all_vms AS (
        SELECT DISTINCT
            COALESCE(e.job_name_normalized, e.job_name) AS vm_name,
            e.client_id
        FROM imported_events e
        WHERE e.event_id = 150
          AND COALESCE(e.event_category, 'backup') != 'log_backup'
          AND e.job_name IS NOT NULL
          AND e.time_created BETWEEN :dt_start AND :dt_end
          {c_cond}
    ),
    success_evts AS (
        SELECT
            COALESCE(e.job_name_normalized, e.job_name) AS vm_name,
            e.client_id, e.time_created
        FROM imported_events e
        WHERE e.event_id = 150
          AND COALESCE(e.event_category, 'backup') != 'log_backup'
          AND e.job_result = 0
          AND e.job_name IS NOT NULL
          AND e.time_created BETWEEN :dt_start AND :dt_end
          {c_cond}
    ),
    ranked_evts AS (
        SELECT vm_name, client_id, time_created,
               LAG(time_created) OVER (
                   PARTITION BY vm_name, client_id ORDER BY time_created
               ) AS prev_time
        FROM success_evts
    ),
    intervals AS (
        SELECT vm_name, client_id,
               EXTRACT(EPOCH FROM (time_created - prev_time)) / 3600.0 AS interval_h
        FROM ranked_evts
        WHERE prev_time IS NOT NULL
    ),
    success_agg AS (
        SELECT vm_name, client_id,
               COUNT(*)          AS backup_count,
               MAX(time_created) AS last_backup
        FROM success_evts
        GROUP BY vm_name, client_id
    ),
    interval_agg AS (
        SELECT vm_name, client_id,
               AVG(interval_h) AS avg_rpo_h,
               MAX(interval_h) AS max_rpo_h,
               MIN(interval_h) AS min_rpo_h
        FROM intervals
        GROUP BY vm_name, client_id
    ),
    vm_stats AS (
        SELECT
            av.vm_name, av.client_id,
            COALESCE(sa.backup_count, 0) AS backup_count,
            sa.last_backup,
            ia.avg_rpo_h, ia.max_rpo_h, ia.min_rpo_h
        FROM all_vms av
        LEFT JOIN success_agg  sa ON sa.vm_name = av.vm_name AND sa.client_id = av.client_id
        LEFT JOIN interval_agg ia ON ia.vm_name = av.vm_name AND ia.client_id = av.client_id
    ),
    vm_rpo AS (
        SELECT
            s.vm_name, s.client_id,
            c.name              AS client_name,
            s.backup_count, s.last_backup,
            s.avg_rpo_h, s.max_rpo_h, s.min_rpo_h,
            vra.id              AS assignment_id,
            vra.expected_minutes,
            vra.criticality,
            rp.id               AS policy_id,
            rp.name             AS policy_name,
            COALESCE(vra.expected_minutes, {DEFAULT_RPO_MINUTES}) / 60.0  AS expected_h,
            CASE WHEN vra.id IS NOT NULL THEN 'configured' ELSE 'default' END AS rpo_source
        FROM vm_stats s
        JOIN clients c ON c.id = s.client_id
        LEFT JOIN vm_rpo_assignments vra
            ON LOWER(vra.vm_name) = LOWER(s.vm_name)
            AND vra.client_id = s.client_id
            AND vra.is_active = true
        LEFT JOIN rpo_policies rp ON rp.id = vra.rpo_policy_id
    )
    SELECT * FROM vm_rpo
    WHERE 1=1
    """

    if search:
        cte += " AND LOWER(vm_name) LIKE :search"
        params["search"] = f"%{search.lower()}%"

    # Status filter usa expected_h (coluna da CTE vm_rpo)
    if rpo_status == "ok":
        cte += " AND backup_count >= 2 AND avg_rpo_h <= expected_h"
    elif rpo_status == "attention":
        cte += " AND backup_count >= 2 AND avg_rpo_h > expected_h AND avg_rpo_h <= expected_h * 1.5"
    elif rpo_status == "critical":
        cte += " AND backup_count >= 2 AND avg_rpo_h > expected_h * 1.5"
    elif rpo_status == "no_history":
        cte += " AND backup_count = 1"
    elif rpo_status == "no_backup":
        cte += " AND backup_count = 0"
    elif rpo_status == "no_configured":
        cte += " AND rpo_source = 'default'"
    elif rpo_status == "critical_all":
        cte += " AND (backup_count = 0 OR (backup_count >= 2 AND avg_rpo_h > expected_h * 1.5))"

    if policy_id:
        cte += " AND policy_id = :policy_id"
        params["policy_id"] = policy_id

    count_sql = f"SELECT COUNT(*) FROM ({cte}) _cnt"
    return cte, count_sql, params


# ══════════════════════════════════════════════════════════════════════════════
# Servidores duplicados (mesma VM em mais de uma rotina) — base Event Viewer
#   EV150 (categoria 'backup') = tarefa por VM;  EV190 = conclusão do job.
#   Ambos compartilham o GUID de sessão (1º <Data> do EventData) → VM↔rotina.
# ══════════════════════════════════════════════════════════════════════════════

@app.get("/analytics/duplicates", response_class=HTMLResponse)
def analytics_duplicates_page(request: Request, db: Session = Depends(get_db)):
    clients = _all_clients_list(db, request)
    return templates.TemplateResponse(
        "analytics_duplicates.html",
        {"request": request, "clients": clients},
    )


@app.get("/api/analytics/duplicates", dependencies=[Depends(scope_guard)])
def api_analytics_duplicates(
    request: Request,
    client_id: Optional[int] = Query(None),
    days: int = Query(30, ge=1, le=365),
    db: Session = Depends(get_db),
):
    """
    Lista VMs/servidores cuja imagem é protegida por MAIS DE UMA rotina no período.
    Detecção 100% via Event Viewer: junta EV150('backup') × EV190 pelo GUID de sessão.
    """
    from sqlalchemy import text as sa_text

    allowed = _scoped_client_ids(db, request)
    params: dict = {"days": days}
    client_clause = ""
    if client_id is not None:
        client_clause = "AND e.client_id = :cid"
        params["cid"] = client_id
    elif allowed is not None:
        # admin sempre tem allowed=None; aqui é não-admin sem client_id (scope_guard já barra),
        # mas por segurança limita ao escopo.
        if not allowed:
            return JSONResponse({"items": [], "total": 0})
        client_clause = "AND e.client_id = ANY(:cids)"
        params["cids"] = list(allowed)

    sql = sa_text(f"""
        WITH j AS (
            SELECT DISTINCT
                e.client_id,
                substring(e.raw_event_xml from '<Data>([0-9a-fA-F-]{{36}})</Data>') AS sess,
                e.job_name AS rotina
            FROM imported_events e
            WHERE e.event_id = 190 AND e.event_category = 'backup'
              AND e.raw_event_xml IS NOT NULL
              AND e.time_created > NOW() - make_interval(days => :days)
              {client_clause}
        ),
        v AS (
            SELECT DISTINCT
                e.client_id,
                substring(e.raw_event_xml from '<Data>([0-9a-fA-F-]{{36}})</Data>') AS sess,
                e.job_name AS vm
            FROM imported_events e
            WHERE e.event_id = 150 AND e.event_category = 'backup'
              AND e.raw_event_xml IS NOT NULL
              AND e.time_created > NOW() - make_interval(days => :days)
              {client_clause}
        )
        SELECT v.client_id, c.name AS client_name, v.vm,
               COUNT(DISTINCT j.rotina) AS qtd,
               STRING_AGG(DISTINCT j.rotina, '||') AS rotinas
        FROM v
        JOIN j ON j.sess = v.sess AND j.client_id = v.client_id
        JOIN clients c ON c.id = v.client_id
        WHERE v.sess IS NOT NULL
        GROUP BY v.client_id, c.name, v.vm
        HAVING COUNT(DISTINCT j.rotina) > 1
        ORDER BY qtd DESC, v.vm
    """)

    rows = db.execute(sql, params).fetchall()
    items = [{
        "server":   r.vm,
        "client":   r.client_name,
        "count":    r.qtd,
        "routines": sorted((r.rotinas or "").split("||")),
    } for r in rows]
    return JSONResponse({"items": items, "total": len(items)})


@app.get("/dashboard/rpo", response_class=HTMLResponse)
def dashboard_rpo_page(request: Request, db: Session = Depends(get_db)):
    clients  = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    policies = (
        db.query(RpoPolicy)
        .filter(RpoPolicy.is_active.is_(True))
        .order_by(RpoPolicy.expected_minutes)
        .all()
    )
    return templates.TemplateResponse(
        "dashboard_rpo.html",
        {"request": request, "clients": clients, "policies": policies,
         "default_rpo_h": DEFAULT_RPO_MINUTES // 60},
    )


@app.get("/api/rpo-dashboard", dependencies=[Depends(scope_guard)])
def api_rpo_dashboard(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    client_id:  Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    """KPIs resumo do RPO observado vs esperado."""
    from sqlalchemy import text as sa_text

    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final.", "code": "dates_required"},
            status_code=400,
        )

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    data_sql, _, params = _rpo_base_sql(dt_start, dt_end, client_id, None, "all")
    rows = db.execute(sa_text(data_sql), params).fetchall()

    vms_ok = vms_attention = vms_critical = vms_no_history = vms_no_backup = vms_no_configured = 0
    rpo_vals: list[float] = []

    for r in rows:
        st = _rpo_status(r.backup_count, r.avg_rpo_h, r.expected_h)
        if r.rpo_source == "default":
            vms_no_configured += 1
        if st == "ok":
            vms_ok += 1
            rpo_vals.append(float(r.avg_rpo_h))
        elif st == "attention":
            vms_attention += 1
            rpo_vals.append(float(r.avg_rpo_h))
        elif st == "critical":
            vms_critical += 1
            rpo_vals.append(float(r.avg_rpo_h))
        elif st == "no_history":
            vms_no_history += 1
        elif st == "no_backup":
            vms_no_backup += 1

    global_avg = round(sum(rpo_vals) / len(rpo_vals), 1) if rpo_vals else None

    return JSONResponse({
        "kpis": {
            "global_avg_rpo_h":   global_avg,
            "vms_ok":             vms_ok,
            "vms_attention":      vms_attention,
            "vms_critical":       vms_critical,
            "vms_no_history":     vms_no_history,
            "vms_no_backup":      vms_no_backup,
            "vms_no_configured":  vms_no_configured,
            "total_vms":          len(rows),
            "default_rpo_h":      DEFAULT_RPO_MINUTES // 60,
        }
    })


@app.get("/api/rpo-vm-list", dependencies=[Depends(scope_guard)])
def api_rpo_vm_list(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    client_id:  Optional[int] = Query(None),
    search:     Optional[str] = Query(None),
    rpo_status: str            = Query("all"),
    policy_id:  Optional[int] = Query(None),
    page:       int            = Query(1, ge=1),
    page_size:  int            = Query(25, ge=1, le=200),
    db: Session = Depends(get_db),
):
    """Lista paginada de VMs com RPO observado + esperado, desvio e status."""
    from sqlalchemy import text as sa_text

    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final.", "code": "dates_required"},
            status_code=400,
        )

    dt_start  = _parse_date(date_start)
    dt_end    = _parse_date(date_end, end_of_day=True)
    page      = max(1, page)
    page_size = min(max(10, page_size), 200)

    data_sql, count_sql, params = _rpo_base_sql(
        dt_start, dt_end, client_id, search, rpo_status, policy_id=policy_id
    )

    total  = db.execute(sa_text(count_sql), params).scalar() or 0
    offset = (page - 1) * page_size
    paged  = (data_sql +
              " ORDER BY avg_rpo_h DESC NULLS LAST, backup_count ASC, vm_name"
              f" LIMIT {page_size} OFFSET {offset}")
    rows   = db.execute(sa_text(paged), params).fetchall()

    ref_dt = dt_end

    vms = []
    for r in rows:
        bc    = r.backup_count
        avg_h = float(r.avg_rpo_h) if r.avg_rpo_h is not None else None
        max_h = float(r.max_rpo_h) if r.max_rpo_h is not None else None
        min_h = float(r.min_rpo_h) if r.min_rpo_h is not None else None
        exp_h = float(r.expected_h) if r.expected_h is not None else DEFAULT_RPO_MINUTES / 60.0
        st    = _rpo_status(bc, avg_h, exp_h)

        since_h = None
        if r.last_backup and ref_dt:
            since_h = round((ref_dt - r.last_backup).total_seconds() / 3600, 1)

        vms.append({
            "vm_name":       r.vm_name,
            "client_id":     r.client_id,
            "client_name":   r.client_name,
            "backup_count":  bc,
            "last_backup":   r.last_backup.strftime("%d/%m/%Y %H:%M") if r.last_backup else None,
            "avg_rpo_h":     round(avg_h, 2) if avg_h is not None else None,
            "max_rpo_h":     round(max_h, 2) if max_h is not None else None,
            "min_rpo_h":     round(min_h, 2) if min_h is not None else None,
            "expected_h":    round(exp_h, 2),
            "expected_minutes": r.expected_minutes,
            "rpo_source":    r.rpo_source,
            "policy_id":     r.policy_id,
            "policy_name":   r.policy_name,
            "assignment_id": r.assignment_id,
            "criticality":   r.criticality,
            "deviation_pct": _calc_deviation_pct(avg_h, exp_h),
            "since_h":       since_h,
            "status":        st,
        })

    return JSONResponse({
        "vms":         vms,
        "total":       total,
        "page":        page,
        "page_size":   page_size,
        "total_pages": max(1, (total + page_size - 1) // page_size),
    })


@app.get("/api/export/rpo-excel", dependencies=[Depends(scope_guard)])
def export_rpo_excel(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    client_id:  Optional[int] = Query(None),
    search:     Optional[str] = Query(None),
    rpo_status: str            = Query("all"),
    policy_id:  Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    """Exporta RPO observado + esperado por VM para Excel."""
    from sqlalchemy import text as sa_text
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if not date_start or not date_end:
        return JSONResponse({"error": "Informe data inicial e data final."}, status_code=400)

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    data_sql, _, params = _rpo_base_sql(dt_start, dt_end, client_id, search, rpo_status, policy_id=policy_id)
    rows = db.execute(
        sa_text(data_sql + " ORDER BY avg_rpo_h DESC NULLS LAST, vm_name LIMIT 10000"),
        params,
    ).fetchall()

    def _fmt_h(h) -> str:
        if h is None: return "—"
        h = float(h)
        d  = int(h // 24); h_rem = h % 24
        hr = int(h_rem);   m     = int(h_rem * 60 % 60)
        if d  > 0: return f"{d}d {hr}h {m}m"
        if hr > 0: return f"{hr}h {m}m"
        return f"{m}m"

    STATUS_LABELS = {
        "ok":         "Dentro do RPO",
        "attention":  "Atenção (até 150% do RPO)",
        "critical":   "Fora do RPO (>150%)",
        "no_history": "Histórico insuficiente",
        "no_backup":  "Sem backup com sucesso",
    }
    STATUS_COLORS = {
        "ok":         "D4EDDA", "attention":  "FFF3CD",
        "critical":   "F8D7DA", "no_history": "E2E3E5", "no_backup":  "F8D7DA",
    }

    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = "RPO por VM"

    hdr_fill  = PatternFill("solid", fgColor="7C3AED")
    hdr_font  = Font(bold=True, color="FFFFFF", size=10)
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin      = Side(style="thin", color="CCCCCC")
    border    = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = [
        "Ambiente / Cliente", "Máquina Virtual",
        "Último Backup com Sucesso",
        "RPO Esperado", "Origem do RPO", "Política",
        "RPO Médio Observado", "Maior RPO", "Menor RPO",
        "Desvio (%)", "Qtd. Backups",
        "Tempo Desde Último Backup", "Status do RPO",
    ]
    widths = [28, 45, 22, 16, 16, 20, 22, 18, 18, 12, 12, 22, 25]

    ws.row_dimensions[1].height = 30
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = hdr_align; cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = w

    for ri, r in enumerate(rows, 2):
        bc    = r.backup_count
        avg_h = float(r.avg_rpo_h) if r.avg_rpo_h is not None else None
        exp_h = float(r.expected_h) if r.expected_h is not None else DEFAULT_RPO_MINUTES / 60.0
        st    = _rpo_status(bc, avg_h, exp_h)

        since_h = None
        if r.last_backup and dt_end:
            since_h = (dt_end - r.last_backup).total_seconds() / 3600

        source_lbl = "Configurado" if r.rpo_source == "configured" else f"Padrão global ({DEFAULT_RPO_MINUTES//60}h)"
        dev = _calc_deviation_pct(avg_h, exp_h)

        row_data = [
            r.client_name, r.vm_name,
            r.last_backup.strftime("%d/%m/%Y %H:%M") if r.last_backup else "—",
            _fmt_h(exp_h), source_lbl, r.policy_name or "—",
            _fmt_h(avg_h), _fmt_h(r.max_rpo_h), _fmt_h(r.min_rpo_h),
            f"+{dev}%" if dev and dev > 0 else (f"{dev}%" if dev is not None else "—"),
            bc, _fmt_h(since_h), STATUS_LABELS.get(st, st),
        ]
        sc = STATUS_COLORS.get(st, "FFFFFF")
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=ri, column=col, value=val)
            cell.border = border; cell.alignment = Alignment(vertical="center")
            cell.font = Font(size=9)
            if col == 13:
                cell.fill = PatternFill("solid", fgColor=sc)
                cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.freeze_panes    = "A2"

    ws2 = wb.create_sheet("Informações")
    ws2["A1"] = "Gerado em"
    ws2["B1"] = datetime.now().strftime("%d/%m/%Y %H:%M")
    ws2["A2"] = "Período analisado"
    ws2["B2"] = f"{dt_start.strftime('%d/%m/%Y')} → {dt_end.strftime('%d/%m/%Y')}"
    ws2["A3"] = "Total de VMs"
    ws2["B3"] = len(rows)
    ws2["A4"] = "Nota"
    ws2["B4"] = (
        "RPO calculado com base em backups bem-sucedidos (EV150, backup). "
        "Transaction Logs excluídos. Status: compara RPO observado x RPO esperado configurado."
    )

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(tmp.name); tmp.close()
    return FileResponse(tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"rpo_por_vm_{date_start}_a_{date_end}.xlsx")


# ── Admin: Políticas de RPO ───────────────────────────────────────────────────

@app.get("/admin/rpo-policies", response_class=HTMLResponse)
def admin_rpo_policies_page(request: Request, db: Session = Depends(get_db)):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    clients  = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    policies = db.query(RpoPolicy).order_by(RpoPolicy.expected_minutes).all()
    return templates.TemplateResponse(
        "admin_rpo_policies.html",
        {"request": request, "clients": clients, "policies": policies,
         "default_rpo_h": DEFAULT_RPO_MINUTES // 60},
    )


@app.get("/api/admin/rpo-policies", dependencies=[Depends(require_admin)])
def api_list_rpo_policies(db: Session = Depends(get_db)):
    """Lista todas as políticas de RPO."""
    policies = db.query(RpoPolicy).order_by(RpoPolicy.expected_minutes).all()
    return [
        {
            "id":               p.id,
            "name":             p.name,
            "description":      p.description,
            "expected_minutes": p.expected_minutes,
            "expected_h":       round(p.expected_minutes / 60, 2),
            "severity_default": p.severity_default,
            "is_active":        p.is_active,
        }
        for p in policies
    ]


@app.post("/api/admin/rpo-policies")
def api_create_rpo_policy(
    request: Request,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    name = (data.get("name") or "").strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nome é obrigatório.")
    minutes = int(data.get("expected_minutes") or 0)
    if minutes <= 0:
        raise HTTPException(status_code=400, detail="expected_minutes deve ser > 0.")
    p = RpoPolicy(
        name             = name,
        description      = data.get("description"),
        expected_minutes = minutes,
        severity_default = data.get("severity_default", "medium"),
        is_active        = bool(data.get("is_active", True)),
        created_at       = datetime.utcnow(),
    )
    db.add(p); db.commit(); db.refresh(p)
    return {"id": p.id, "name": p.name, "expected_minutes": p.expected_minutes}


@app.put("/api/admin/rpo-policies/{pol_id}")
def api_update_rpo_policy(
    pol_id: int, request: Request,
    data: dict = Body(...), db: Session = Depends(get_db),
):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    p = db.query(RpoPolicy).get(pol_id)
    if not p:
        raise HTTPException(status_code=404)
    for field in ("name", "description", "severity_default"):
        if field in data:
            setattr(p, field, data[field])
    if "expected_minutes" in data:
        mins = int(data["expected_minutes"] or 0)
        if mins > 0:
            p.expected_minutes = mins
    if "is_active" in data:
        p.is_active = bool(data["is_active"])
    p.updated_at = datetime.utcnow()
    db.commit(); db.refresh(p)
    return {"id": p.id, "name": p.name, "expected_minutes": p.expected_minutes}


@app.delete("/api/admin/rpo-policies/{pol_id}")
def api_delete_rpo_policy(
    pol_id: int, request: Request, db: Session = Depends(get_db),
):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    p = db.query(RpoPolicy).get(pol_id)
    if not p:
        raise HTTPException(status_code=404)
    # Desvincula atribuições (SET NULL via FK) antes de excluir
    db.query(VmRpoAssignment).filter(VmRpoAssignment.rpo_policy_id == pol_id).update(
        {"rpo_policy_id": None}, synchronize_session=False
    )
    db.delete(p); db.commit()
    return {"ok": True}


# ── VM RPO Assignments ────────────────────────────────────────────────────────

@app.get("/api/rpo-assignment", dependencies=[Depends(scope_guard)])
def api_get_rpo_assignment(
    client_id: int            = Query(...),
    vm_name:   str            = Query(...),
    db: Session = Depends(get_db),
):
    """Retorna a atribuição de RPO atual de uma VM (ou null)."""
    a = (
        db.query(VmRpoAssignment)
        .filter(
            VmRpoAssignment.client_id == client_id,
            VmRpoAssignment.vm_name   == vm_name,
            VmRpoAssignment.is_active.is_(True),
        )
        .first()
    )
    if not a:
        return JSONResponse({"assignment": None})
    return JSONResponse({
        "assignment": {
            "id":               a.id,
            "rpo_policy_id":    a.rpo_policy_id,
            "expected_minutes": a.expected_minutes,
            "criticality":      a.criticality,
            "notes":            a.notes,
        }
    })


@app.post("/api/rpo-assignment", dependencies=[Depends(scope_guard)])
def api_save_rpo_assignment(
    request: Request,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    """Cria ou atualiza o vínculo de uma VM com uma política/valor de RPO."""
    client_id = data.get("client_id")
    vm_name   = (data.get("vm_name") or "").strip()
    if not client_id or not vm_name:
        raise HTTPException(status_code=400, detail="client_id e vm_name são obrigatórios.")
    minutes   = int(data.get("expected_minutes") or 0)
    if minutes <= 0:
        raise HTTPException(status_code=400, detail="expected_minutes deve ser > 0.")

    actor = getattr(request.state, "username", "sistema")

    # Upsert via UniqueConstraint
    existing = (
        db.query(VmRpoAssignment)
        .filter(
            VmRpoAssignment.client_id == client_id,
            VmRpoAssignment.vm_name   == vm_name,
        )
        .first()
    )
    if existing:
        existing.rpo_policy_id    = data.get("rpo_policy_id")
        existing.expected_minutes = minutes
        existing.criticality      = data.get("criticality")
        existing.notes            = data.get("notes")
        existing.is_active        = True
        existing.updated_by       = actor
        existing.updated_at       = datetime.utcnow()
        db.commit()
        return {"id": existing.id, "action": "updated"}
    else:
        a = VmRpoAssignment(
            client_id         = client_id,
            vm_name           = vm_name,
            rpo_policy_id     = data.get("rpo_policy_id"),
            expected_minutes  = minutes,
            criticality       = data.get("criticality"),
            notes             = data.get("notes"),
            is_active         = True,
            created_by        = actor,
            created_at        = datetime.utcnow(),
        )
        db.add(a); db.commit(); db.refresh(a)
        return {"id": a.id, "action": "created"}


@app.delete("/api/rpo-assignment/{assignment_id}")
def api_delete_rpo_assignment(
    assignment_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    """Remove a atribuição de RPO de uma VM."""
    a = db.query(VmRpoAssignment).get(assignment_id)
    if not a:
        raise HTTPException(status_code=404)
    db.delete(a); db.commit()
    return {"ok": True}


@app.get("/api/admin/vm-rpo-assignments")
def api_list_vm_rpo_assignments(
    client_id: Optional[int] = Query(None),
    search:    Optional[str] = Query(None),
    page:      int            = Query(1, ge=1),
    page_size: int            = Query(30, ge=1, le=200),
    db: Session = Depends(get_db),
):
    """Lista atribuições de RPO com filtro e paginação."""
    q = (
        db.query(VmRpoAssignment, Client.name.label("client_name"),
                 RpoPolicy.name.label("policy_name"))
        .join(Client, Client.id == VmRpoAssignment.client_id)
        .outerjoin(RpoPolicy, RpoPolicy.id == VmRpoAssignment.rpo_policy_id)
        .filter(VmRpoAssignment.is_active.is_(True))
    )
    if client_id:
        q = q.filter(VmRpoAssignment.client_id == client_id)
    if search:
        q = q.filter(VmRpoAssignment.vm_name.ilike(f"%{search}%"))

    total = q.count()
    rows  = (q.order_by(VmRpoAssignment.vm_name)
              .limit(page_size).offset((page-1)*page_size).all())

    return JSONResponse({
        "assignments": [
            {
                "id":               a.id,
                "client_id":        a.client_id,
                "client_name":      client_name,
                "vm_name":          a.vm_name,
                "rpo_policy_id":    a.rpo_policy_id,
                "policy_name":      policy_name,
                "expected_minutes": a.expected_minutes,
                "expected_h":       round(a.expected_minutes / 60, 2),
                "criticality":      a.criticality,
                "notes":            a.notes,
                "created_by":       a.created_by,
                "updated_by":       a.updated_by,
            }
            for a, client_name, policy_name in rows
        ],
        "total":       total,
        "page":        page,
        "page_size":   page_size,
        "total_pages": max(1, (total + page_size - 1) // page_size),
    })


# ══════════════════════════════════════════════════════════════════════════════
# CLASSIFICAÇÃO DE OFENSORES — Admin CRUD + Incidents Dashboard
# ══════════════════════════════════════════════════════════════════════════════

# ── Helpers ───────────────────────────────────────────────────────────────────

def _category_to_dict(cat: OffenderCategory, include_rules: bool = True) -> dict:
    d: dict = {
        "id":          cat.id,
        "name":        cat.name,
        "description": cat.description,
        "severity":    cat.severity,
        "color":       cat.color,
        "is_active":   cat.is_active,
        "priority":    cat.priority,
        "notes":       cat.notes,
        "created_at":  cat.created_at.isoformat() if cat.created_at else None,
        "updated_at":  cat.updated_at.isoformat() if cat.updated_at else None,
    }
    if include_rules:
        d["rules"] = [_rule_to_dict(r) for r in sorted(cat.rules, key=lambda x: (-x.priority, x.id))]
    return d


def _rule_to_dict(rule: OffenderRule) -> dict:
    return {
        "id":             rule.id,
        "category_id":   rule.category_id,
        "pattern":        rule.pattern,
        "match_type":     rule.match_type,
        "case_sensitive": rule.case_sensitive,
        "is_active":      rule.is_active,
        "priority":       rule.priority,
        "created_at":     rule.created_at.isoformat() if rule.created_at else None,
    }


# ── Página admin ──────────────────────────────────────────────────────────────

@app.get("/admin/offenders", response_class=HTMLResponse)
def admin_offenders_page(request: Request, db: Session = Depends(get_db)):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    return templates.TemplateResponse("admin_offenders.html", {"request": request})


# ── API: Categorias ───────────────────────────────────────────────────────────

@app.get("/api/admin/offenders", dependencies=[Depends(require_admin)])
def api_list_offenders(db: Session = Depends(get_db)):
    cats = (
        db.query(OffenderCategory)
        .order_by(OffenderCategory.priority.desc(), OffenderCategory.name)
        .all()
    )
    return [_category_to_dict(c, include_rules=True) for c in cats]


@app.post("/api/admin/offenders")
def api_create_offender(
    request: Request,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    cat = OffenderCategory(
        name        = (data.get("name") or "").strip(),
        description = data.get("description"),
        severity    = data.get("severity", "medium"),
        color       = data.get("color"),
        priority    = int(data.get("priority") or 100),
        is_active   = bool(data.get("is_active", True)),
        notes       = data.get("notes"),
        created_at  = datetime.utcnow(),
    )
    if not cat.name:
        raise HTTPException(status_code=400, detail="Nome é obrigatório.")
    db.add(cat)
    db.commit()
    db.refresh(cat)
    offender_service.invalidate_cache()
    return _category_to_dict(cat)


@app.put("/api/admin/offenders/{cat_id}")
def api_update_offender(
    cat_id: int,
    request: Request,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    cat = db.query(OffenderCategory).get(cat_id)
    if not cat:
        raise HTTPException(status_code=404)
    for field in ("name", "description", "severity", "color", "notes"):
        if field in data:
            setattr(cat, field, data[field])
    if "priority" in data:
        cat.priority  = int(data["priority"] or 100)
    if "is_active" in data:
        cat.is_active = bool(data["is_active"])
    cat.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(cat)
    offender_service.invalidate_cache()
    return _category_to_dict(cat)


@app.delete("/api/admin/offenders/{cat_id}")
def api_delete_offender(
    cat_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    cat = db.query(OffenderCategory).get(cat_id)
    if not cat:
        raise HTTPException(status_code=404)
    db.delete(cat)
    db.commit()
    offender_service.invalidate_cache()
    return {"ok": True}


# ── API: Regras ───────────────────────────────────────────────────────────────

@app.post("/api/admin/offenders/{cat_id}/rules")
def api_add_rule(
    cat_id: int,
    request: Request,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    cat = db.query(OffenderCategory).get(cat_id)
    if not cat:
        raise HTTPException(status_code=404)
    pattern = (data.get("pattern") or "").strip()
    if not pattern:
        raise HTTPException(status_code=400, detail="Pattern é obrigatório.")
    rule = OffenderRule(
        category_id    = cat_id,
        pattern        = pattern,
        match_type     = data.get("match_type", "contains"),
        case_sensitive = bool(data.get("case_sensitive", False)),
        is_active      = bool(data.get("is_active", True)),
        priority       = int(data.get("priority") or 100),
        created_at     = datetime.utcnow(),
    )
    db.add(rule)
    db.commit()
    db.refresh(rule)
    offender_service.invalidate_cache()
    return _rule_to_dict(rule)


@app.put("/api/admin/offenders/rules/{rule_id}")
def api_update_rule(
    rule_id: int,
    request: Request,
    data: dict = Body(...),
    db: Session = Depends(get_db),
):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    rule = db.query(OffenderRule).get(rule_id)
    if not rule:
        raise HTTPException(status_code=404)
    for field in ("pattern", "match_type"):
        if field in data:
            setattr(rule, field, data[field])
    if "case_sensitive" in data:
        rule.case_sensitive = bool(data["case_sensitive"])
    if "is_active" in data:
        rule.is_active = bool(data["is_active"])
    if "priority" in data:
        rule.priority = int(data["priority"] or 100)
    rule.updated_at = datetime.utcnow()
    db.commit()
    offender_service.invalidate_cache()
    return _rule_to_dict(rule)


@app.delete("/api/admin/offenders/rules/{rule_id}")
def api_delete_rule(
    rule_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    rule = db.query(OffenderRule).get(rule_id)
    if not rule:
        raise HTTPException(status_code=404)
    db.delete(rule)
    db.commit()
    offender_service.invalidate_cache()
    return {"ok": True}


# ── API: Reprocessamento ──────────────────────────────────────────────────────

@app.post("/api/admin/offenders/reprocess")
def api_reprocess_offenders(
    request: Request,
    data: dict = Body(default={}),
    db: Session = Depends(get_db),
):
    """Reclassifica todos (ou filtrados por cliente) os eventos de falha."""
    if not request.state.is_admin:
        raise HTTPException(status_code=403)
    client_id = data.get("client_id")
    n = offender_service.classify_events_batch(
        db,
        client_id=client_id if client_id else None,
        reprocess_all=True,
    )
    return {"classified": n}


# ── Dashboard: Incidentes ─────────────────────────────────────────────────────

@app.get("/dashboard/incidents", response_class=HTMLResponse)
def dashboard_incidents_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    categories = (
        db.query(OffenderCategory)
        .filter(OffenderCategory.is_active.is_(True))
        .order_by(OffenderCategory.priority.desc())
        .all()
    )
    return templates.TemplateResponse(
        "dashboard_incidents.html",
        {"request": request, "clients": clients, "categories": categories},
    )


@app.get("/api/incidents-dashboard", dependencies=[Depends(scope_guard)])
def api_incidents_dashboard(
    date_start:  Optional[str] = Query(None),
    date_end:    Optional[str] = Query(None),
    client_id:   Optional[int] = Query(None),
    offender_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    from sqlalchemy import text as sa_text

    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final.", "code": "dates_required"},
            status_code=400,
        )

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    # ── Base filter ───────────────────────────────────────────────────────────
    def _base_q():
        q = db.query(ImportedEvent).filter(
            ImportedEvent.job_result == 2,
            ImportedEvent.event_category.in_(["backup", "log_backup"]),
            ImportedEvent.time_created >= dt_start,
            ImportedEvent.time_created <= dt_end,
        )
        if client_id:
            q = q.filter(ImportedEvent.client_id == client_id)
        return q

    total_errors  = _base_q().count()
    classified    = _base_q().filter(ImportedEvent.offender_category_id.isnot(None)).count()
    unclassified  = total_errors - classified

    # ── Por ambiente (cliente) ─────────────────────────────────────────────────
    from sqlalchemy import func as sqla_func

    env_rows = (
        db.query(Client.name, sqla_func.count(ImportedEvent.id).label("cnt"))
        .join(ImportedEvent, ImportedEvent.client_id == Client.id)
        .filter(
            ImportedEvent.job_result == 2,
            ImportedEvent.event_category.in_(["backup", "log_backup"]),
            ImportedEvent.time_created >= dt_start,
            ImportedEvent.time_created <= dt_end,
        )
    )
    if client_id:
        env_rows = env_rows.filter(ImportedEvent.client_id == client_id)
    env_rows = env_rows.group_by(Client.name).order_by(sqla_func.count(ImportedEvent.id).desc()).all()

    by_environment = [{"name": r.name, "count": r.cnt} for r in env_rows]

    # ── Por ofensor ───────────────────────────────────────────────────────────
    off_sql = """
    SELECT
        oc.id,
        oc.name,
        oc.color,
        oc.severity,
        COUNT(ie.id)    AS event_count,
        ROUND(COUNT(ie.id) * 100.0 / NULLIF(:total, 0), 1) AS pct,
        MODE() WITHIN GROUP (ORDER BY ie.offender_match_pattern) AS top_pattern
    FROM offender_categories oc
    JOIN imported_events ie ON ie.offender_category_id = oc.id
    WHERE ie.job_result = 2
      AND ie.event_category IN ('backup', 'log_backup')
      AND ie.time_created BETWEEN :dt_start AND :dt_end
      {client_cond}
      {offender_cond}
    GROUP BY oc.id, oc.name, oc.color, oc.severity
    ORDER BY event_count DESC
    """
    off_params: dict = {"total": total_errors, "dt_start": dt_start, "dt_end": dt_end}
    client_cond  = "AND ie.client_id = :client_id" if client_id else ""
    offender_cond = "AND oc.id = :offender_id" if offender_id else ""
    if client_id:
        off_params["client_id"] = client_id
    if offender_id:
        off_params["offender_id"] = offender_id

    off_rows = db.execute(
        sa_text(off_sql.format(client_cond=client_cond, offender_cond=offender_cond)),
        off_params,
    ).fetchall()

    offenders = [
        {
            "id":          r.id,
            "name":        r.name,
            "color":       r.color or "#6c757d",
            "severity":    r.severity,
            "count":       r.event_count,
            "pct":         float(r.pct or 0),
            "top_pattern": r.top_pattern,
        }
        for r in off_rows
    ]

    # ── Evolução mensal ───────────────────────────────────────────────────────
    monthly_sql = """
    SELECT
        TO_CHAR(DATE_TRUNC('month', ie.time_created), 'YYYY-MM') AS month,
        c.name      AS client_name,
        COUNT(ie.id) AS error_count
    FROM imported_events ie
    JOIN clients c ON c.id = ie.client_id
    WHERE ie.job_result = 2
      AND ie.event_category IN ('backup', 'log_backup')
      AND ie.time_created BETWEEN :dt_start AND :dt_end
      {client_cond}
    GROUP BY DATE_TRUNC('month', ie.time_created), c.name
    ORDER BY month, client_name
    """
    monthly_rows = db.execute(
        sa_text(monthly_sql.format(client_cond=client_cond)),
        {"dt_start": dt_start, "dt_end": dt_end, **({} if not client_id else {"client_id": client_id})},
    ).fetchall()

    # Pivot: {month_label → {client → count}}
    monthly_map: dict = {}
    clients_seen: list = []
    for r in monthly_rows:
        lbl = r.month  # "AAAA-MM"
        if lbl not in monthly_map:
            monthly_map[lbl] = {}
        monthly_map[lbl][r.client_name] = r.error_count
        if r.client_name not in clients_seen:
            clients_seen.append(r.client_name)

    monthly = [
        {"month": m, "label": _month_label(m), "by_client": monthly_map[m]}
        for m in sorted(monthly_map)
    ]

    return JSONResponse({
        "kpis": {
            "total_errors":  total_errors,
            "classified":    classified,
            "unclassified":  unclassified,
            "classification_rate": round(classified * 100 / total_errors, 1) if total_errors else 0,
            "top_offender":  offenders[0]["name"] if offenders else None,
            "active_offenders": len(offenders),
        },
        "by_environment": by_environment,
        "offenders":      offenders,
        "monthly":        monthly,
        "monthly_clients": clients_seen,
    })


def _month_label(month_str: str) -> str:
    """'2026-01' → 'Jan/26'"""
    MONTHS = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun",
              "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
    try:
        y, m = month_str.split("-")
        return f"{MONTHS[int(m)-1]}/{y[2:]}"
    except Exception:
        return month_str


@app.get("/api/incidents-unclassified", dependencies=[Depends(scope_guard)])
def api_incidents_unclassified(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    client_id:  Optional[int] = Query(None),
    page:       int            = Query(1, ge=1),
    page_size:  int            = Query(20, ge=1, le=200),
    db: Session = Depends(get_db),
):
    """Eventos de falha sem classificação de ofensor (para revisão e criação de regras)."""
    from sqlalchemy import text as sa_text

    if not date_start or not date_end:
        return JSONResponse({"error": "Informe data inicial e data final."}, status_code=400)

    dt_start  = _parse_date(date_start)
    dt_end    = _parse_date(date_end, end_of_day=True)
    page      = max(1, page)
    page_size = min(max(10, page_size), 100)

    cond = ""
    params: dict = {"dt_start": dt_start, "dt_end": dt_end}
    if client_id:
        cond = "AND ie.client_id = :client_id"
        params["client_id"] = client_id

    # Agrupar por mensagem truncada para exibir padrões recorrentes
    sql = f"""
    SELECT
        LEFT(ie.original_message, 300)  AS msg_preview,
        COUNT(*)                        AS occurrence_count,
        MAX(ie.time_created)            AS last_seen,
        MAX(c.name)                     AS client_name,
        MAX(ie.job_name)                AS job_name
    FROM imported_events ie
    JOIN clients c ON c.id = ie.client_id
    WHERE ie.job_result = 2
      AND ie.event_category IN ('backup', 'log_backup')
      AND ie.offender_category_id IS NULL
      AND ie.original_message IS NOT NULL
      AND ie.time_created BETWEEN :dt_start AND :dt_end
      {cond}
    GROUP BY LEFT(ie.original_message, 300)
    ORDER BY occurrence_count DESC
    LIMIT {page_size} OFFSET {(page-1)*page_size}
    """
    cnt_sql = f"""
    SELECT COUNT(DISTINCT LEFT(ie.original_message, 300))
    FROM imported_events ie
    WHERE ie.job_result = 2
      AND ie.event_category IN ('backup', 'log_backup')
      AND ie.offender_category_id IS NULL
      AND ie.original_message IS NOT NULL
      AND ie.time_created BETWEEN :dt_start AND :dt_end
      {cond}
    """

    rows  = db.execute(sa_text(sql),     params).fetchall()
    total = db.execute(sa_text(cnt_sql), params).scalar() or 0

    events = [
        {
            "msg_preview":       r.msg_preview,
            "occurrence_count":  r.occurrence_count,
            "last_seen":         r.last_seen.strftime("%d/%m/%Y %H:%M") if r.last_seen else "—",
            "client_name":       r.client_name,
            "job_name":          r.job_name,
        }
        for r in rows
    ]

    return JSONResponse({
        "events":      events,
        "total":       total,
        "page":        page,
        "page_size":   page_size,
        "total_pages": max(1, (total + page_size - 1) // page_size),
    })


@app.get("/api/incidents-offender-detail", dependencies=[Depends(scope_guard)])
def api_incidents_offender_detail(
    offender_id: int            = Query(...),
    date_start:  Optional[str] = Query(None),
    date_end:    Optional[str] = Query(None),
    client_id:   Optional[int] = Query(None),
    page:        int            = Query(1, ge=1),
    page_size:   int            = Query(20, ge=1, le=200),
    db: Session = Depends(get_db),
):
    """Drill-down: eventos de falha classificados em um ofensor."""
    if not date_start or not date_end:
        return JSONResponse({"error": "Informe datas."}, status_code=400)

    dt_start  = _parse_date(date_start)
    dt_end    = _parse_date(date_end, end_of_day=True)
    page      = max(1, page)
    page_size = min(max(10, page_size), 100)
    offset    = (page - 1) * page_size

    q = (
        db.query(ImportedEvent, Client.name.label("client_name"))
        .join(Client, Client.id == ImportedEvent.client_id)
        .filter(
            ImportedEvent.offender_category_id == offender_id,
            ImportedEvent.job_result == 2,
            ImportedEvent.time_created >= dt_start,
            ImportedEvent.time_created <= dt_end,
        )
    )
    if client_id:
        q = q.filter(ImportedEvent.client_id == client_id)

    total = q.count()
    rows  = q.order_by(ImportedEvent.time_created.desc()).limit(page_size).offset(offset).all()

    events = [
        {
            "id":            ev.id,
            "time_created":  ev.time_created.strftime("%d/%m/%Y %H:%M") if ev.time_created else "—",
            "client_name":   client_name,
            "job_name":      ev.job_name or "—",
            "msg_preview":   (ev.original_message or "")[:250],
            "match_pattern": ev.offender_match_pattern,
        }
        for ev, client_name in rows
    ]

    return JSONResponse({
        "events":      events,
        "total":       total,
        "page":        page,
        "page_size":   page_size,
        "total_pages": max(1, (total + page_size - 1) // page_size),
    })


@app.get("/api/export/incidents-excel", dependencies=[Depends(scope_guard)])
def export_incidents_excel(
    date_start:  Optional[str] = Query(None),
    date_end:    Optional[str] = Query(None),
    client_id:   Optional[int] = Query(None),
    offender_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    """Exporta incidentes classificados para Excel."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    if not date_start or not date_end:
        return JSONResponse({"error": "Informe datas."}, status_code=400)

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    q = (
        db.query(ImportedEvent, Client.name.label("client_name"),
                 OffenderCategory.name.label("offender_name"))
        .join(Client, Client.id == ImportedEvent.client_id)
        .outerjoin(OffenderCategory, OffenderCategory.id == ImportedEvent.offender_category_id)
        .filter(
            ImportedEvent.job_result == 2,
            ImportedEvent.event_category.in_(["backup", "log_backup"]),
            ImportedEvent.time_created >= dt_start,
            ImportedEvent.time_created <= dt_end,
        )
    )
    if client_id:
        q = q.filter(ImportedEvent.client_id == client_id)
    if offender_id:
        q = q.filter(ImportedEvent.offender_category_id == offender_id)
    rows = q.order_by(ImportedEvent.time_created.desc()).limit(20000).all()

    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = "Incidentes"

    thin      = Side(style="thin", color="CCCCCC")
    border    = Border(left=thin, right=thin, top=thin, bottom=thin)
    hdr_fill  = PatternFill("solid", fgColor="7C3AED")
    hdr_font  = Font(bold=True, color="FFFFFF", size=10)
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = ["Data/Hora", "Ambiente / Cliente", "Job", "Ofensor", "Pattern encontrado",
               "Mensagem (resumo)"]
    widths  = [20, 30, 45, 30, 35, 60]

    ws.row_dimensions[1].height = 24
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = hdr_align; cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = w

    for ri, (ev, client_name, offender_name) in enumerate(rows, 2):
        row_data = [
            ev.time_created.strftime("%d/%m/%Y %H:%M") if ev.time_created else "—",
            client_name,
            ev.job_name or "—",
            offender_name or "Não classificado",
            ev.offender_match_pattern or "—",
            (ev.original_message or "")[:300],
        ]
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=ri, column=col, value=val)
            cell.border = border
            cell.font   = Font(size=9)
            cell.alignment = Alignment(vertical="center", wrap_text=(col == 6))

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.freeze_panes    = "A2"

    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
    wb.save(tmp.name)
    tmp.close()
    return FileResponse(
        tmp.name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"incidentes_{date_start}_a_{date_end}.xlsx",
    )


@app.get("/dashboard/weekly", response_class=HTMLResponse)
def dashboard_weekly_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse("dashboard_weekly.html", {"request": request, "clients": clients})


@app.get("/api/weekly-report-data", dependencies=[Depends(scope_guard)])
def api_weekly_report_data(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    client_id:  Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final para consultar.", "code": "dates_required"},
            status_code=400,
        )

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    q      = _build_event_query(db, dt_start, dt_end, "all", client_id)
    events = q.order_by(ImportedEvent.time_created.desc()).all()

    last_per_job: dict = {}
    for e in events:
        if e.job_name and e.job_name not in last_per_job:
            last_per_job[e.job_name] = e

    RESULT_LABEL = {0: "Sucesso", 1: "Warning", 2: "Falha"}
    failed_jobs:  list = []
    success_jobs: list = []
    for job_name, e in last_per_job.items():
        entry = {
            "job_name":       job_name,
            "job_type":       e.event_category or "backup",
            "last_execution": e.time_created.strftime("%d/%m/%Y %H:%M") if e.time_created else "—",
        }
        (failed_jobs if e.job_result == 2 else success_jobs).append(entry)
    failed_jobs.sort(key=lambda x: x["job_name"])

    total = len(last_per_job)
    admin_kpis = {
        "failed":      len(failed_jobs),
        "success":     len(success_jobs),
        "failed_pct":  round(len(failed_jobs)  / total * 100, 1) if total else 0.0,
        "success_pct": round(len(success_jobs) / total * 100, 1) if total else 0.0,
    }

    def _q150_190(eid):
        q = db.query(ImportedEvent).filter(ImportedEvent.event_id == eid)
        if client_id:
            q = q.filter(ImportedEvent.client_id == client_id)
        if dt_start:
            q = q.filter(ImportedEvent.time_created >= dt_start)
        if dt_end:
            q = q.filter(ImportedEvent.time_created <= dt_end)
        return q.all()

    ev150 = _q150_190(150)
    ev190 = _q150_190(190)
    protection_kpis = {
        "protected_vms": len({e.job_name for e in ev150 if e.job_name}),
        "log_vms":       len({e.job_name for e in ev150 if e.job_name and e.event_category == "log_backup"}),
        "backup_jobs":   len({e.job_name for e in ev190 if e.job_name}),
    }

    all_dates = [e.time_created for e in events if e.time_created]
    all_dates += [e.time_created for e in ev150 if e.time_created]
    if dt_start and dt_end:
        period_str = f"{dt_start.strftime('%d/%m/%Y')} → {dt_end.strftime('%d/%m/%Y')}"
    elif all_dates:
        period_str = f"{min(all_dates).strftime('%d/%m/%Y')} → {max(all_dates).strftime('%d/%m/%Y')}"
    else:
        period_str = "—"

    client_name = ""
    if client_id:
        c = db.query(Client).get(client_id)
        if c:
            client_name = c.name

    return JSONResponse({
        "admin":       {"kpis": admin_kpis, "failed_jobs": failed_jobs, "success_jobs": success_jobs},
        "protection":  {"kpis": protection_kpis},
        "period":      period_str,
        "client_name": client_name,
        "clients":     _all_clients_list(db),
    })


@app.post("/api/export/weekly-pptx", dependencies=[Depends(scope_guard)])
async def export_weekly_pptx(request: Request, db: Session = Depends(get_db)):
    from fastapi.responses import Response as FastAPIResponse
    from app.services.pptx_service import generate_weekly_pptx

    body      = await request.json()
    dt_start  = _parse_date(body.get("date_start")) if body.get("date_start") else None
    dt_end    = _parse_date(body.get("date_end"), end_of_day=True) if body.get("date_end") else None
    client_id = body.get("client_id")

    q      = _build_event_query(db, dt_start, dt_end, "all", client_id)
    events = q.order_by(ImportedEvent.time_created.desc()).all()

    last_per_job: dict = {}
    for e in events:
        if e.job_name and e.job_name not in last_per_job:
            last_per_job[e.job_name] = e

    failed_jobs:  list = []
    success_jobs: list = []
    for job_name, e in last_per_job.items():
        entry = {
            "job_name":       job_name,
            "job_type":       e.event_category or "backup",
            "last_execution": e.time_created.strftime("%d/%m/%Y %H:%M") if e.time_created else "—",
        }
        (failed_jobs if e.job_result == 2 else success_jobs).append(entry)
    failed_jobs.sort(key=lambda x: x["job_name"])

    total = len(last_per_job)
    admin_kpis = {
        "failed":      len(failed_jobs),
        "success":     len(success_jobs),
        "failed_pct":  round(len(failed_jobs)  / total * 100, 1) if total else 0.0,
        "success_pct": round(len(success_jobs) / total * 100, 1) if total else 0.0,
    }

    def _q2(eid):
        q2 = db.query(ImportedEvent).filter(ImportedEvent.event_id == eid)
        if client_id:
            q2 = q2.filter(ImportedEvent.client_id == client_id)
        if dt_start:
            q2 = q2.filter(ImportedEvent.time_created >= dt_start)
        if dt_end:
            q2 = q2.filter(ImportedEvent.time_created <= dt_end)
        return q2.all()

    ev150 = _q2(150)
    ev190 = _q2(190)
    protection_kpis = {
        "protected_vms": len({e.job_name for e in ev150 if e.job_name}),
        "log_vms":       len({e.job_name for e in ev150 if e.job_name and e.event_category == "log_backup"}),
        "backup_jobs":   len({e.job_name for e in ev190 if e.job_name}),
    }

    all_dates = [e.time_created for e in events if e.time_created]
    if dt_start and dt_end:
        period_str = f"{dt_start.strftime('%d/%m/%Y')} → {dt_end.strftime('%d/%m/%Y')}"
    elif all_dates:
        period_str = f"{min(all_dates).strftime('%d/%m/%Y')} → {max(all_dates).strftime('%d/%m/%Y')}"
    else:
        period_str = "—"

    client_name = ""
    if client_id:
        c = db.query(Client).get(client_id)
        if c:
            client_name = c.name

    _branding = branding_service.get_branding(client_id=client_id)
    pptx_bytes = generate_weekly_pptx(
        client_name=client_name, period=period_str,
        kpis=admin_kpis, protection_kpis=protection_kpis, failed_jobs=failed_jobs,
        branding=_branding, branding_dir=str(BRANDING_DIR),
    )
    safe_name = (client_name or "todos").replace(" ", "_")
    return FastAPIResponse(
        content=pptx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="Relatorio_Semanal_{safe_name}.pptx"'},
    )


# ════════════════════════════════════════════════════════════════════════════════
# RELATÓRIO MENSAL — visão mensal de backup regular (independente do semanal)
# ════════════════════════════════════════════════════════════════════════════════
#
# Diferenças vs. semanal:
#   • Sucesso × Falha contabilizados por EXECUÇÃO (eventos de sessão EV190 regulares),
#     não pelo último status de cada job.
#   • Transaction Logs (event_category='log_backup') são EXCLUÍDOS do gráfico e do top-5.

def _monthly_report_payload(db: Session, dt_start, dt_end, client_id):
    """Calcula KPIs, distribuição sucesso×falha (por execução) e top-5 falhas do período."""
    from collections import defaultdict

    def _q_eid(eid):
        q = db.query(ImportedEvent).filter(ImportedEvent.event_id == eid)
        if client_id: q = q.filter(ImportedEvent.client_id == client_id)
        if dt_start:  q = q.filter(ImportedEvent.time_created >= dt_start)
        if dt_end:    q = q.filter(ImportedEvent.time_created <= dt_end)
        return q.all()

    # ── Cards (EV150 = proteção) ──────────────────────────────────────────────
    ev150 = _q_eid(150)
    protected_vms       = len({e.job_name for e in ev150 if e.job_name})
    transaction_log_vms = len({e.job_name for e in ev150
                               if e.job_name and e.event_category == "log_backup"})

    # ── Execuções de backup REGULAR (sessão de job, EV190–200, exclui log) ────
    qj = (
        db.query(ImportedEvent)
        .filter(ImportedEvent.event_id >= 190, ImportedEvent.event_id <= 200)
        .filter(ImportedEvent.event_category != "log_backup")
        .filter(ImportedEvent.job_name.isnot(None))
    )
    if client_id: qj = qj.filter(ImportedEvent.client_id == client_id)
    if dt_start:  qj = qj.filter(ImportedEvent.time_created >= dt_start)
    if dt_end:    qj = qj.filter(ImportedEvent.time_created <= dt_end)
    reg = qj.order_by(ImportedEvent.time_created.asc()).all()

    backup_jobs      = len({e.job_name for e in reg})
    success_count    = sum(1 for e in reg if e.job_result == 0)
    failed_count     = sum(1 for e in reg if e.job_result == 2)
    warning_count    = sum(1 for e in reg if e.job_result == 1)
    total_executions = len(reg)
    base = success_count + failed_count            # base do donut (Sucesso × Falha)
    success_percent  = round(success_count / base * 100, 1) if base else 0.0
    failed_percent   = round(failed_count  / base * 100, 1) if base else 0.0

    # ── Top 5 rotinas com erro ────────────────────────────────────────────────
    agg: dict = defaultdict(lambda: {
        "failed": 0, "last_failure": None, "last_execution": None, "last_result": None,
    })
    for e in reg:   # ordem asc → o último visto é a execução mais recente
        s = agg[e.job_name]
        if e.job_result == 2:
            s["failed"] += 1
            if e.time_created and (s["last_failure"] is None or e.time_created > s["last_failure"]):
                s["last_failure"] = e.time_created
        if e.time_created and (s["last_execution"] is None or e.time_created >= s["last_execution"]):
            s["last_execution"] = e.time_created
            s["last_result"]    = e.job_result

    STATUS = {0: "Sucesso", 1: "Warning", 2: "Falha"}
    ranked = sorted(
        ((jn, s) for jn, s in agg.items() if s["failed"] > 0),
        key=lambda kv: kv[1]["failed"], reverse=True,
    )[:5]
    top_failed_jobs = [{
        "job_name":       jn,
        "job_type":       "Backup",
        "failed_count":   s["failed"],
        "last_failure":   s["last_failure"].strftime("%d/%m/%Y %H:%M") if s["last_failure"] else "—",
        "last_execution": s["last_execution"].strftime("%d/%m/%Y %H:%M") if s["last_execution"] else "—",
        "current_status": STATUS.get(s["last_result"], "Sem histórico suficiente"),
    } for jn, s in ranked]

    period_str = (f"{dt_start.strftime('%d/%m/%Y')} → {dt_end.strftime('%d/%m/%Y')}"
                  if dt_start and dt_end else "—")
    client_name = ""
    if client_id:
        c = db.query(Client).get(client_id)
        if c:
            client_name = c.name

    return {
        "client_name": client_name,
        "date_start":  dt_start.strftime("%Y-%m-%d") if dt_start else None,
        "date_end":    dt_end.strftime("%Y-%m-%d") if dt_end else None,
        "period":      period_str,
        "kpis": {
            "protected_vms":       protected_vms,
            "transaction_log_vms": transaction_log_vms,
            "backup_jobs":         backup_jobs,
        },
        "backup_distribution": {
            "success_count":    success_count,
            "failed_count":     failed_count,
            "warning_count":    warning_count,
            "success_percent":  success_percent,
            "failed_percent":   failed_percent,
            "total_executions": total_executions,
        },
        "top_failed_jobs": top_failed_jobs,
    }


def _monthly_report_payload_ps(db: Session, dt_start, dt_end, client_id):
    """
    Igual ao _monthly_report_payload, mas 100% baseado no PowerShell (backup_vm_sessions).
    A unidade é a EXECUÇÃO DE JOB (job_session_id) — Top 5 e distribuição são POR ROTINA,
    não por VM. Resultado da execução: Falha se qualquer VM falhou, senão Warning se
    houve aviso, senão Sucesso.
    """
    from sqlalchemy import text as sa_text

    conds = ["client_id = :cid"] if client_id else []
    params: dict = {}
    if client_id:
        params["cid"] = client_id
    if dt_start:
        conds.append("started_at >= :d0"); params["d0"] = dt_start
    if dt_end:
        conds.append("started_at <= :d1"); params["d1"] = dt_end
    where = (" WHERE " + " AND ".join(conds)) if conds else ""

    sess_cte = f"""
        WITH sess AS (
            SELECT COALESCE(job_session_id, 'row-' || id::text) AS sk,
                   max(job_name_snapshot) AS job_name,
                   min(started_at)        AS started_at,
                   bool_or(result = 'Failed')  AS f,
                   bool_or(result = 'Warning') AS w
            FROM backup_vm_sessions{where}
            GROUP BY 1
        )
    """

    # Distribuição por execução de job
    agg = db.execute(sa_text(sess_cte + """
        SELECT count(*) AS total,
               count(*) FILTER (WHERE f) AS failed,
               count(*) FILTER (WHERE (NOT f) AND w) AS warning,
               count(*) FILTER (WHERE (NOT f) AND (NOT w)) AS success,
               count(DISTINCT job_name) AS jobs
        FROM sess
    """), params).first()

    total_executions = agg.total or 0
    failed_count  = agg.failed or 0
    warning_count = agg.warning or 0
    success_count = agg.success or 0
    backup_jobs   = agg.jobs or 0
    base = success_count + failed_count
    success_percent = round(success_count / base * 100, 1) if base else 0.0
    failed_percent  = round(failed_count  / base * 100, 1) if base else 0.0

    # VMs protegidas (distintas)
    vms = db.execute(sa_text(f"SELECT count(DISTINCT vm_name) AS n FROM backup_vm_sessions{where}"),
                     params).first()
    protected_vms = (vms.n if vms else 0) or 0

    # Top 5 rotinas com mais falhas (por execução de job)
    top_rows = db.execute(sa_text(sess_cte + """
        , sj AS (
            SELECT job_name, started_at,
                   (CASE WHEN f THEN 2 WHEN w THEN 1 ELSE 0 END) AS res
            FROM sess
        ),
        last AS (
            SELECT DISTINCT ON (job_name) job_name, res AS last_res
            FROM sj ORDER BY job_name, started_at DESC
        )
        SELECT sj.job_name AS job_name,
               count(*) FILTER (WHERE res = 2)               AS failed,
               max(started_at) FILTER (WHERE res = 2)        AS last_failure,
               max(started_at)                               AS last_execution,
               max(l.last_res)                               AS last_result
        FROM sj JOIN last l ON l.job_name = sj.job_name
        GROUP BY sj.job_name
        HAVING count(*) FILTER (WHERE res = 2) > 0
        ORDER BY failed DESC, sj.job_name ASC
        LIMIT 5
    """), params).all()

    STATUS = {0: "Sucesso", 1: "Warning", 2: "Falha"}
    top_failed_jobs = [{
        "job_name":       r.job_name,
        "job_type":       "Backup",
        "failed_count":   r.failed,
        "last_failure":   r.last_failure.strftime("%d/%m/%Y %H:%M") if r.last_failure else "—",
        "last_execution": r.last_execution.strftime("%d/%m/%Y %H:%M") if r.last_execution else "—",
        "current_status": STATUS.get(r.last_result, "Sem histórico suficiente"),
    } for r in top_rows]

    period_str = (f"{dt_start.strftime('%d/%m/%Y')} → {dt_end.strftime('%d/%m/%Y')}"
                  if dt_start and dt_end else "—")
    client_name = ""
    if client_id:
        c = db.query(Client).get(client_id)
        if c:
            client_name = c.name

    return {
        "client_name": client_name,
        "date_start":  dt_start.strftime("%Y-%m-%d") if dt_start else None,
        "date_end":    dt_end.strftime("%Y-%m-%d") if dt_end else None,
        "period":      period_str,
        "kpis": {
            "protected_vms":       protected_vms,
            "transaction_log_vms": None,               # não se aplica ao PowerShell
            "backup_jobs":         backup_jobs,
            "job_executions":      total_executions,
        },
        "backup_distribution": {
            "success_count":    success_count,
            "failed_count":     failed_count,
            "warning_count":    warning_count,
            "success_percent":  success_percent,
            "failed_percent":   failed_percent,
            "total_executions": total_executions,
        },
        "top_failed_jobs": top_failed_jobs,
    }


@app.get("/dashboard/monthly-ps", response_class=HTMLResponse)
def dashboard_monthly_ps_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse("dashboard_monthly_ps.html", {"request": request, "clients": clients})


@app.get("/api/monthly-ps-report-data", dependencies=[Depends(scope_guard)])
def api_monthly_ps_report_data(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    client_id:  Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    if not date_start or not date_end:
        return JSONResponse({"error": "Informe data inicial e data final para consultar.",
                             "code": "dates_required"}, status_code=400)
    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)
    if not dt_start or not dt_end:
        return JSONResponse({"error": "Datas inválidas.", "code": "dates_invalid"}, status_code=400)
    payload = _monthly_report_payload_ps(db, dt_start, dt_end, client_id)
    payload["clients"] = _all_clients_list(db)
    return JSONResponse(payload)


@app.post("/api/export/monthly-ps-pptx", dependencies=[Depends(scope_guard)])
async def export_monthly_ps_pptx(request: Request, db: Session = Depends(get_db)):
    from fastapi.responses import Response as FastAPIResponse
    from app.services.pptx_service import generate_monthly_backup_pptx

    body      = await request.json()
    dt_start  = _parse_date(body.get("date_start")) if body.get("date_start") else None
    dt_end    = _parse_date(body.get("date_end"), end_of_day=True) if body.get("date_end") else None
    client_id = body.get("client_id")

    p   = _monthly_report_payload_ps(db, dt_start, dt_end, client_id)
    dist = p["backup_distribution"]
    kpis = {
        "failed":      dist["failed_count"],
        "success":     dist["success_count"],
        "failed_pct":  dist["failed_percent"],
        "success_pct": dist["success_percent"],
        "total":       dist["total_executions"],
    }
    # card do meio: no PowerShell não há Transaction Log -> "Execuções (job)"
    protection = {
        "protected_vms":       p["kpis"]["protected_vms"],
        "transaction_log_vms": p["kpis"]["job_executions"],
        "backup_jobs":         p["kpis"]["backup_jobs"],
    }

    _branding = branding_service.get_branding(client_id=client_id)
    pptx_bytes = generate_monthly_backup_pptx(
        client_name=p["client_name"], period=p["period"],
        kpis=kpis, protection_kpis=protection, top_failed_jobs=p["top_failed_jobs"],
        branding=_branding, branding_dir=str(BRANDING_DIR),
        middle_kpi_label="Execuções (job)", middle_kpi_desc="Execuções de rotina no período",
        subtitle="Execuções de rotina (job) no período  •  base PowerShell",
    )
    safe_name = (p["client_name"] or "todos").replace(" ", "_")
    return FastAPIResponse(
        content=pptx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="Relatorio_Mensal_PowerShell_{safe_name}.pptx"'},
    )


# ══════════════════════════════════════════════════════════════════════════════
# RELATÓRIO DE CLIENTE (rotina) — dossiê Word por rotina de backup (provedor cloud)
# ══════════════════════════════════════════════════════════════════════════════

def _fmt_dur_hm(secs) -> str:
    if secs is None:
        return "—"
    secs = int(secs)
    h, m = secs // 3600, (secs % 3600) // 60
    if h:
        return f"{h}h{m:02d}m"
    return f"{m}m{secs % 60:02d}s"


def _fmt_hours(hours) -> str:
    if hours is None:
        return "—"
    if hours >= 48:
        return f"{hours/24:.1f} dias"
    return f"{hours:.1f} h"


def _job_names_from_body(body: dict) -> list:
    """Extrai a lista de rotinas do corpo, aceitando job_names[] (novo) ou
    job_name (compat). Remove vazios e duplicatas preservando a ordem."""
    raw = body.get("job_names")
    if not isinstance(raw, list):
        raw = [body.get("job_name")]
    seen, out = set(), []
    for j in raw:
        j = (j or "").strip()
        if j and j not in seen:
            seen.add(j)
            out.append(j)
    return out


def _config_items(snap, *, include_encryption: bool) -> list:
    """Lista rotulada de parâmetros vigentes de UM snapshot (privacidade:
    sem nomes de proxies/extents — só a contagem de extents)."""
    from app.routers.job_config_analytics import (
        _schedule_display, _retention_display, _gfs_display, _parse_json,
    )
    yn = lambda v: "Sim" if v else ("Não" if v is not None else "—")
    extents = _parse_json(snap.repo_extent_names) or []
    cap = snap.repo_capacity_name or "—"
    if snap.repo_capacity_name:
        cap += " (imutável)" if snap.repo_capacity_immutable else " (sem imutabilidade)"
    items = [
        ("Agendamento",              _schedule_display(snap)),
        ("Retenção",                 _retention_display(snap)),
        ("GFS (retenção estendida)", _gfs_display(snap)),
        ("Repositório",              snap.repo_sobr_name or snap.repo_name or "—"),
        ("Imutabilidade (performance tier)", yn(snap.repo_immutability)),
        ("Linux Hardened",           yn(snap.repo_is_linux_hardened)),
        ("Capacity tier",            cap),
    ]
    if include_encryption:
        items.append(("Criptografia no armazenamento", yn(snap.stg_encryption)))
    items.append(("Application-Aware", yn(snap.aap_enabled)))
    items.append(("Rotina habilitada", yn(snap.is_enabled)))
    if extents:
        items.insert(4, ("Redundância (performance tier)",
                         f"{len(extents)} repositórios (extents)"))
    return items


def _rpo_from_snapshot(snap, editorial_rpo: str = "") -> str:
    """RPO objetivo: valor editorial (se houver) ou derivado do agendamento."""
    rpo = (editorial_rpo or "").strip()
    if rpo:
        return rpo
    if snap.sched_periodically_enabled and snap.sched_periodically_every:
        return (f"{snap.sched_periodically_every} "
                f"{snap.sched_periodically_unit or 'hora(s)'} (execução periódica)")
    if snap.sched_daily_enabled:
        import re as _re2
        m = _re2.search(r"(\d{1,2}:\d{2})", snap.sched_daily_time or "")
        return "24 horas (execução diária" + (f" às {m.group(1)})" if m else ")")
    return "Conforme agendamento da rotina"


def _compliance_rows(snap, restore_test_freq: str) -> list:
    """Matriz de conformidade (ISO 27001/27040, NIST CSF, CIS v8) de UM snapshot."""
    from app.routers.job_config_analytics import (
        _schedule_display, _retention_display, _gfs_display,
    )
    immut = bool(snap.repo_immutability or snap.repo_is_linux_hardened
                 or snap.repo_capacity_immutable)
    immut_parts = []
    if snap.repo_immutability:
        immut_parts.append("performance tier imutável")
    if snap.repo_is_linux_hardened:
        immut_parts.append("repositório Linux Hardened")
    if snap.repo_capacity_immutable:
        d = f" ({snap.repo_capacity_immut_days} dias)" if snap.repo_capacity_immut_days else ""
        immut_parts.append(f"capacity tier imutável{d}")
    sched_ok = bool(snap.is_enabled) and bool(
        snap.is_schedule_enabled or snap.sched_daily_enabled
        or snap.sched_periodically_enabled)
    has_ret = any(v is not None for v in
                  (snap.ret_restore_points, snap.ret_cycles, snap.ret_days))
    tri = lambda v: ("CONFORME" if v else "ATENÇÃO") if v is not None else "N/D"
    return [
        {"ref": "ISO/IEC 27001 A.8.13",
         "req": "Retenção de backup definida",
         "current": _retention_display(snap),
         "status": "CONFORME" if has_ret else "N/D"},
        {"ref": "ISO/IEC 27001 A.8.13 • ISO 22301",
         "req": "Execução automática agendada",
         "current": _schedule_display(snap)
                    + ("" if snap.is_enabled else " — ROTINA DESABILITADA"),
         "status": "CONFORME" if sched_ok else "NÃO CONFORME"},
        {"ref": "ISO/IEC 27001 A.8.13",
         "req": "Retenção de longo prazo (GFS)",
         "current": _gfs_display(snap),
         "status": tri(snap.ret_gfs_enabled)},
        {"ref": "ISO/IEC 27040 • NIST CSF PR.DS-11",
         "req": "Imutabilidade dos backups (anti-ransomware)",
         "current": ", ".join(immut_parts).capitalize() if immut_parts
                    else "Sem camada imutável identificada",
         "status": "CONFORME" if immut else "ATENÇÃO"},
        {"ref": "CIS v8 — 11.4",
         "req": "Cópia secundária de dados (regra 3-2-1)",
         "current": (f"Capacity tier: {snap.repo_capacity_name}"
                     if snap.repo_capacity_name
                     else "Sem camada secundária identificada"),
         "status": "CONFORME" if snap.repo_capacity_name else "ATENÇÃO"},
        {"ref": "ISO/IEC 27001 A.8.24 • ISO/IEC 27040",
         "req": "Criptografia dos dados de backup",
         "current": ("Habilitada no armazenamento" if snap.stg_encryption
                     else "Não habilitada na rotina" if snap.stg_encryption is not None
                     else "Não informado pela coleta"),
         "status": tri(snap.stg_encryption)},
        {"ref": "ISO/IEC 27001 A.8.13",
         "req": "Consistência de aplicação (application-aware)",
         "current": ("Habilitada" if snap.aap_enabled
                     else "Não habilitada" if snap.aap_enabled is not None
                     else "Não informado pela coleta"),
         "status": tri(snap.aap_enabled)},
        {"ref": "NIST CSF PR.DS-11 • CIS v8 — 11.1",
         "req": "Monitoramento e relatório de execução",
         "current": "Monitoramento contínuo + relatório mensal de backup",
         "status": "CONFORME"},
        {"ref": "CIS v8 — 11.5",
         "req": "Testes de restauração periódicos",
         "current": "Periodicidade definida nesta política: "
                    + (restore_test_freq or "Trimestral"),
         "status": "CONFORME"},
    ]


def _latest_snapshot(db, client_id: int, job_name: str):
    from app.models import JobConfigSnapshot
    return (db.query(JobConfigSnapshot)
            .filter(JobConfigSnapshot.client_id == client_id,
                    JobConfigSnapshot.job_name == job_name)
            .order_by(JobConfigSnapshot.collected_at.desc().nullslast())
            .first())


def _client_report_payload(db: Session, client_id: int, job_names, dt_start, dt_end,
                           report_name: str, tlp: str,
                           author: Optional[str] = None) -> dict:
    """Agrega dados de UMA OU MAIS rotinas (backup_vm_sessions + config audit).

    KPIs de execução são COMBINADOS entre as rotinas; a configuração e as
    alterações são apresentadas POR ROTINA (cada job tem sua config)."""
    from collections import Counter, defaultdict
    from app.models import BackupVmSession, JobConfigSnapshot
    from app.routers.job_config_analytics import _parse_json, _audited_values

    if isinstance(job_names, str):
        job_names = [job_names]
    job_names = [j for j in (jn.strip() for jn in job_names) if j]
    multi_job = len(job_names) > 1

    rows = (db.query(BackupVmSession)
            .filter(BackupVmSession.client_id == client_id,
                    BackupVmSession.job_name_snapshot.in_(job_names),
                    BackupVmSession.started_at >= dt_start,
                    BackupVmSession.started_at <= dt_end)
            .order_by(BackupVmSession.started_at.asc())
            .all())

    # ── Sessões (execuções de job) ────────────────────────────────────────────
    sess: dict = defaultdict(list)
    for s in rows:
        sess[s.job_session_id or f"row-{s.id}"].append(s)

    executions = len(sess)
    success = warning = failed = 0
    durations, processed = [], []
    failures = []
    day_counts: dict = defaultdict(lambda: {"execs": 0, "fails": 0})
    for grp in sess.values():
        starts = [g.started_at for g in grp if g.started_at]
        ends   = [g.ended_at for g in grp if g.ended_at]
        dur = int((max(ends) - min(starts)).total_seconds()) if starts and ends else None
        if dur is not None and dur >= 0:
            durations.append(dur)
        gbs = [g.processed_gb for g in grp if g.processed_gb is not None]
        if gbs:
            processed.append(sum(gbs))
        res = ("Failed" if any(g.result == "Failed" for g in grp)
               else "Warning" if any(g.result == "Warning" for g in grp) else "Success")
        if starts:
            dkey = min(starts).date()
            day_counts[dkey]["execs"] += 1
            if res == "Failed":
                day_counts[dkey]["fails"] += 1
        if res == "Failed":
            failed += 1
            fvms = sorted({g.vm_name for g in grp if g.result == "Failed" and g.vm_name})
            reasons = Counter((g.reason_raw or "").strip()
                              for g in grp if g.result == "Failed" and (g.reason_raw or "").strip())
            top = reasons.most_common(1)[0][0] if reasons else "—"
            if len(reasons) > 1:
                top += f"  (+{len(reasons)-1} outro(s))"
            failures.append({
                "date":   min(starts).strftime("%d/%m/%Y %H:%M") if starts else "—",
                "vms":    f"{len(fvms)}: " + ", ".join(fvms) if fvms else "—",
                "reason": top,
            })
        elif res == "Warning":
            warning += 1
        else:
            success += 1

    base = success + failed
    success_rate = f"{(success / base * 100):.1f}%" if base else "—"

    # ── Por VM (RPO observado) ────────────────────────────────────────────────
    # "(sem VM)" é fallback de sessão sem detalhe por VM — não é máquina real,
    # não entra no detalhe nem na contagem de VMs protegidas do report.
    byvm: dict = defaultdict(list)
    for s in rows:
        if s.vm_name and s.vm_name != "(sem VM)":
            byvm[s.vm_name].append(s)
    vms_out = []
    for vm in sorted(byvm):
        lst = sorted(byvm[vm], key=lambda x: x.started_at or dt_start)
        okd = [x.started_at for x in lst if x.result == "Success" and x.started_at]
        gaps = [(okd[i+1] - okd[i]).total_seconds() / 3600 for i in range(len(okd)-1)]
        durs = [x.duration_seconds for x in lst if x.duration_seconds is not None]
        gbs  = [x.processed_gb for x in lst if x.processed_gb is not None]
        vms_out.append({
            "vm":           vm,
            "executions":   len(lst),
            "successes":    len(okd),
            "last_success": okd[-1].strftime("%d/%m/%Y %H:%M") if okd else "NUNCA no período",
            "rpo_avg":      _fmt_hours(sum(gaps)/len(gaps)) if gaps else "—",
            "rpo_max":      _fmt_hours(max(gaps)) if gaps else "—",
            "avg_dur":      _fmt_dur_hm(sum(durs)/len(durs)) if durs else "—",
            "avg_gb":       f"{sum(gbs)/len(gbs):.1f} GB" if gbs else "—",
        })

    # ── Configuração POR ROTINA (snapshot mais recente de cada job) ──────────
    # REGRA: infraestrutura interna do provedor NÃO aparece (sem PROXIES/EXTENTS).
    configs = []
    for jn in job_names:
        snap = _latest_snapshot(db, client_id, jn)
        if not snap:
            continue
        configs.append({
            "job_name":     jn,
            "items":        _config_items(snap, include_encryption=False),
            "collected_at": (snap.collected_at.strftime("%d/%m/%Y %H:%M")
                             if snap.collected_at else None),
        })
    # compat com caminho single-job do docx
    config = configs[0]["items"] if configs else None
    config_collected_at = configs[0]["collected_at"] if configs else None

    # ── Alterações de configuração NO PERÍODO (config drift por rotina) ──────
    # Compara coletas consecutivas (mesma lógica do histórico da Config Audit),
    # considerando só mudanças detectadas em coletas dentro do período.
    # PRIVACIDADE: campos de infra do provedor (proxies/GIP) ficam FORA.
    _HIDDEN_FIELDS = {"Proxies", "Proxy (modo)", "GIP (modo)"}
    config_changes = []
    for jn in job_names:
        snaps_all = (db.query(JobConfigSnapshot)
                     .filter(JobConfigSnapshot.client_id == client_id,
                             JobConfigSnapshot.job_name == jn)
                     .order_by(JobConfigSnapshot.collected_at.asc().nullslast())
                     .all())
        for prev_s, curr_s in zip(snaps_all, snaps_all[1:]):
            if (not curr_s.collected_at
                    or curr_s.collected_at < dt_start or curr_s.collected_at > dt_end):
                continue
            prev_vals = dict(_audited_values(prev_s))
            for label, new_val in _audited_values(curr_s):
                if label in _HIDDEN_FIELDS:
                    continue
                old_val = prev_vals.get(label)
                if str(old_val) != str(new_val):
                    config_changes.append({
                        "date":  curr_s.collected_at.strftime("%d/%m/%Y"),
                        "job":   jn,
                        "field": (f"[{jn}] {label}" if multi_job else label),
                        "from":  old_val if old_val is not None else "—",
                        "to":    new_val if new_val is not None else "—",
                    })

    # série diária p/ o gráfico (todos os dias do período, zeros incluídos)
    from datetime import timedelta as _td
    daily = []
    d = dt_start.date()
    while d <= dt_end.date() and len(daily) <= 62:
        c = day_counts.get(d, {"execs": 0, "fails": 0})
        daily.append({"label": d.strftime("%d"), "execs": c["execs"], "fails": c["fails"]})
        d += _td(days=1)

    # ── Trilha de AUDITORIA do Veeam (Event Viewer) — quem/quando/o quê ───────
    # Fonte mais detalhada que o config-drift: operador, hora exata e de→para
    # por propriedade. PRIVACIDADE: propriedades de proxy/GIP ficam fora e o
    # domínio/host do operador é removido (BMHV04\user → user).
    import re as _re

    def _short(v, n=70):
        # valores de agendamento vêm como XML cru do Veeam — remove as tags
        v = _re.sub(r"<[^>]+>", " ", str(v or ""))
        v = " ".join(v.split())
        return (v[:n] + "…") if len(v) > n else (v or "—")

    audit_changes = []
    audit_events = (db.query(ImportedEvent)
                    .filter(ImportedEvent.client_id == client_id,
                            ImportedEvent.event_category == "audit",
                            ImportedEvent.job_name.in_(job_names),
                            ImportedEvent.time_created >= dt_start,
                            ImportedEvent.time_created <= dt_end)
                    .order_by(ImportedEvent.time_created.asc())
                    .all())
    for ev in audit_events:
        det = ev.audit_details or ""
        lines = []
        if ev.audit_event_type == "job_updated":
            for name, old, new in _re.findall(
                    r'<property name="([^"]+)"[^>]*>(?:<old>(.*?)</old>)?(?:<new>(.*?)</new>)?</property>',
                    det, _re.S):
                if "proxy" in name.lower():          # infra do provedor
                    continue
                o, nv = _short(old), _short(new)
                if o == "—" and nv == "—":           # valor era só XML técnico
                    lines.append(f"{name} (alterado)")
                else:
                    lines.append(f"{name}:  {o}  →  {nv}")
        elif ev.audit_event_type in ("objects_added", "objects_deleted", "objects_changed"):
            objs = _re.findall(r'<Property internal="Name"[^>]*>([^<]+)</Property>', det)
            verbo = {"objects_added": "VM(s) adicionada(s)",
                     "objects_deleted": "VM(s) removida(s)",
                     "objects_changed": "Objeto(s) alterado(s)"}[ev.audit_event_type]
            lines.append(f"{verbo}: {', '.join(objs) if objs else '—'}")
        if not lines:
            lines = ["—"]
        op = (ev.operator or "—").split("\\")[-1]
        audit_changes.append({
            "dt":       ev.time_created.strftime("%d/%m/%Y %H:%M") if ev.time_created else "—",
            "job":      ev.job_name or "—",
            "action":   ev.audit_event_label or ev.audit_event_type or "—",
            "operator": op,
            "lines":    lines,
        })
    audit_changes_more = max(0, len(audit_changes) - 60)
    audit_changes = audit_changes[:60]

    # ── PRIVACIDADE: mascara nomes de infra interna em TEXTO CRU do Veeam ─────
    # Motivos de falha (e eventualmente linhas de auditoria) citam extents e
    # proxies pelo nome (ex.: "vmwarerep01 extent is offline") — infra do
    # provedor não aparece no report do cliente final.
    # Nomes vêm de TODAS as coletas do ambiente (o extent citado numa falha
    # pode pertencer a outro SOBR que não o desta rotina).
    internal_names = set()
    for ext_json, prox_json in (db.query(JobConfigSnapshot.repo_extent_names,
                                         JobConfigSnapshot.proxy_selected_names)
                                .filter(JobConfigSnapshot.client_id == client_id)
                                .distinct().all()):
        internal_names.update(_parse_json(ext_json) or [])
        internal_names.update(_parse_json(prox_json) or [])
    internal_names = {n for n in internal_names if n and len(str(n)) >= 3}
    if internal_names:
        mask_pat = _re.compile(
            "|".join(_re.escape(str(n)) for n in
                     sorted(internal_names, key=lambda x: len(str(x)),
                            reverse=True)),
            _re.IGNORECASE)
        for f in failures:
            f["reason"] = mask_pat.sub("[infraestrutura do provedor]", f["reason"])
        for ch in audit_changes:
            ch["lines"] = [mask_pat.sub("[infraestrutura do provedor]", ln)
                           for ln in ch["lines"]]

    # dia mais crítico (mais falhas) p/ o painel do gráfico
    worst_day = None
    if day_counts:
        wd, wc = max(day_counts.items(), key=lambda kv: kv[1]["fails"])
        if wc["fails"] > 0:
            worst_day = {"label": wd.strftime("%d/%m"), "fails": wc["fails"]}

    _MESES = ["", "JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO",
              "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]
    if dt_start.month == dt_end.month and dt_start.year == dt_end.year:
        period_label = f"{_MESES[dt_start.month]}/{dt_start.year}"
    else:
        period_label = f"{dt_start.strftime('%d/%m/%Y')} a {dt_end.strftime('%d/%m/%Y')}"

    client = db.query(Client).get(client_id)
    slug = "".join(ch for ch in report_name.upper() if ch.isalnum())[:24] or "CLIENTE"
    return {
        "report_name":  report_name,
        "tlp":          tlp,
        "period":       f"{dt_start.strftime('%d/%m/%Y')} a {dt_end.strftime('%d/%m/%Y')}",
        "period_label": period_label,
        "env_name":     client.name if client else "",
        "job_name":     "; ".join(job_names),
        "job_names":    job_names,
        "multi_job":    multi_job,
        "doc_ref":      f"REL-{dt_start.strftime('%Y%m')}-{slug}",
        "generated_at": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "generated_date": datetime.now().strftime("%d/%m/%Y"),
        "author":       author or "—",
        "daily":        daily,
        "worst_day":    worst_day,
        "config_changes": config_changes,
        "audit_changes": audit_changes,
        "audit_changes_more": audit_changes_more,
        "kpis": {
            "protected_vms": len(byvm),
            "executions":    executions,
            "success":       success,
            "warning":       warning,
            "failed":        failed,
            "success_rate":  success_rate,
            "avg_duration":  _fmt_dur_hm(sum(durations)/len(durations)) if durations else "—",
            "max_duration":  _fmt_dur_hm(max(durations)) if durations else "—",
            "avg_processed": f"{sum(processed)/len(processed):.1f} GB" if processed else "—",
        },
        "config":              config,
        "config_collected_at": config_collected_at,
        "configs":             configs,
        "vms":                 vms_out,
        "failures":            failures,
    }


@app.get("/dashboard/client-report", response_class=HTMLResponse)
def dashboard_client_report_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse("dashboard_client_report.html",
                                      {"request": request, "clients": clients})


@app.get("/api/client-report/routines", dependencies=[Depends(scope_guard)])
def api_client_report_routines(
    client_id: int = Query(...),
    search: Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    """Rotinas (job_name distintos) do cliente, para o autocomplete da tela."""
    from sqlalchemy import text as sa_text
    params = {"cid": client_id}
    cond = ""
    if search:
        cond = "AND LOWER(job_name_snapshot) LIKE :s"
        params["s"] = f"%{search.lower()}%"
    rows = db.execute(sa_text(f"""
        SELECT job_name_snapshot AS job_name,
               count(DISTINCT COALESCE(job_session_id, 'r'||id::text)) AS executions,
               max(started_at) AS last_exec
        FROM backup_vm_sessions
        WHERE client_id = :cid AND job_name_snapshot IS NOT NULL {cond}
        GROUP BY job_name_snapshot
        ORDER BY job_name_snapshot
        LIMIT 500
    """), params).all()
    return JSONResponse({"items": [{
        "job_name":   r.job_name,
        "executions": r.executions,
        "last_exec":  r.last_exec.strftime("%d/%m/%Y %H:%M") if r.last_exec else None,
    } for r in rows]})


@app.post("/api/export/client-report-docx", dependencies=[Depends(scope_guard)])
async def export_client_report_docx(request: Request, db: Session = Depends(get_db)):
    from fastapi.responses import Response as FastAPIResponse
    from app.services.docx_service import generate_client_report_docx

    body        = await request.json()
    client_id   = body.get("client_id")
    job_names   = _job_names_from_body(body)
    report_name = (body.get("report_name") or "").strip()
    tlp         = (body.get("tlp") or "AMBER").upper()
    dt_start    = _parse_date(body.get("date_start")) if body.get("date_start") else None
    dt_end      = _parse_date(body.get("date_end"), end_of_day=True) if body.get("date_end") else None

    if not client_id or not job_names or not report_name or not dt_start or not dt_end:
        raise HTTPException(status_code=400,
                            detail="Informe cliente, ao menos uma rotina, nome do report e período.")
    if tlp not in ("CLEAR", "GREEN", "AMBER", "RED"):
        tlp = "AMBER"

    author = (request.session.get("full_name")
              or request.session.get("username") or None)
    payload = _client_report_payload(db, client_id, job_names, dt_start, dt_end,
                                     report_name, tlp, author=author)
    if payload["kpis"]["executions"] == 0:
        raise HTTPException(status_code=404,
                            detail="Nenhuma execução das rotinas selecionadas no período.")

    _branding = branding_service.get_branding(client_id=client_id)
    data = generate_client_report_docx(payload, branding=_branding,
                                       branding_dir=str(BRANDING_DIR))

    slug = "".join(ch if ch.isalnum() else "_" for ch in report_name)[:40]
    fname = f"Relatorio_Backup_{slug}_{dt_start.strftime('%Y%m')}.docx"
    return FastAPIResponse(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


def _backup_policy_payload(db: Session, client_id: int, job_names,
                           editorial: dict, author: Optional[str] = None) -> Optional[dict]:
    """
    Dados da Política de Backup (PSI) de UMA OU MAIS rotinas: parâmetros
    vigentes da Config Audit + matriz de conformidade automática (ISO
    27001/27040, NIST CSF, CIS v8) + campos editoriais do formulário.

    Cada rotina tem SUA configuração e SUA matriz — apresentadas por rotina.
    O escopo de VMs é a união das rotinas.

    Retorna:
      • None                          → nenhuma rotina informada
      • {"missing_config": [jobs]}     → alguma rotina sem coleta de config
      • dict completo                  → ok

    PRIVACIDADE: mesma regra do Relatório de Cliente — nomes de proxies e
    extents (infra do provedor) NÃO aparecem.
    """
    from app.models import JobConfigObject

    if isinstance(job_names, str):
        job_names = [job_names]
    job_names = [j for j in (jn.strip() for jn in job_names) if j]
    if not job_names:
        return None
    multi_job = len(job_names) > 1

    # snapshot mais recente de cada rotina; toda rotina precisa de coleta
    snaps = [(jn, _latest_snapshot(db, client_id, jn)) for jn in job_names]
    missing = [jn for jn, s in snaps if s is None]
    if missing:
        return {"missing_config": missing}

    ed_rpo = (editorial.get("rpo") or "").strip()
    test_freq = editorial.get("restore_test_freq") or "Trimestral"

    configs, compliance_by_job, rpo_by_job = [], [], []
    all_vms, collected_dates = [], []
    platform = "Veeam Backup & Replication"
    for jn, snap in snaps:
        configs.append({
            "job_name":     jn,
            "items":        _config_items(snap, include_encryption=True),
            "collected_at": (snap.collected_at.strftime("%d/%m/%Y %H:%M")
                             if snap.collected_at else None),
        })
        compliance_by_job.append({"job_name": jn,
                                  "rows": _compliance_rows(snap, test_freq)})
        rpo_by_job.append({"job_name": jn, "value": _rpo_from_snapshot(snap, ed_rpo)})
        if snap.collected_at:
            collected_dates.append(snap.collected_at)
        if snap.platform:
            platform = f"Veeam Backup & Replication ({snap.platform})"
        objs = (db.query(JobConfigObject)
                .filter(JobConfigObject.job_config_snapshot_id == snap.id)
                .order_by(JobConfigObject.object_name.asc())
                .all())
        all_vms.extend(o.object_name for o in objs if o.object_name)

    # escopo = união das VMs das rotinas (sem duplicar)
    seen, vm_names = set(), []
    for v in all_vms:
        if v not in seen:
            seen.add(v); vm_names.append(v)
    scope = {
        "vm_count": len(vm_names) if vm_names else "—",
        "vms":      vm_names[:40] + ([f"… (+{len(vm_names)-40})"] if len(vm_names) > 40 else []),
        "platform": platform,
    }

    # RPO: se todas as rotinas têm o mesmo → um valor; senão, por rotina
    distinct_rpo = {r["value"] for r in rpo_by_job}
    rpo_target = rpo_by_job[0]["value"] if len(distinct_rpo) == 1 else None

    client = db.query(Client).get(client_id)
    client_name = editorial.get("client_name") or "; ".join(job_names)
    slug = "".join(ch for ch in client_name.upper() if ch.isalnum())[:24] or "CLIENTE"
    collected_at = (max(collected_dates).strftime("%d/%m/%Y %H:%M")
                    if collected_dates else None)
    return {
        "client_name":    client_name,
        "tlp":            editorial.get("tlp") or "AMBER",
        "env_name":       client.name if client else "",
        "job_name":       "; ".join(job_names),
        "job_names":      job_names,
        "multi_job":      multi_job,
        "doc_ref":        f"PSI-BKP-{slug}",
        "doc_version":    (editorial.get("doc_version") or "1.0").strip() or "1.0",
        "generated_date": datetime.now().strftime("%d/%m/%Y"),
        "author":         author or "—",
        "scope":          scope,
        "config":         configs[0]["items"],          # compat single-job
        "configs":        configs,
        "config_collected_at": collected_at,
        "rpo_target":        rpo_target,
        "rpo_by_job":        rpo_by_job,
        "rto":               (editorial.get("rto") or "").strip() or "Conforme contrato",
        "restore_test_freq": test_freq,
        "review_cycle":      editorial.get("review_cycle") or "Anual",
        "responsible_client": (editorial.get("responsible_client") or "").strip()
                              or "A designar pelo cliente",
        "legal_retention":   (editorial.get("legal_retention") or "").strip(),
        "compliance":        compliance_by_job[0]["rows"],   # compat single-job
        "compliance_by_job": compliance_by_job,
    }


@app.get("/dashboard/backup-policy", response_class=HTMLResponse)
def dashboard_backup_policy_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse("dashboard_backup_policy.html",
                                      {"request": request, "clients": clients})


@app.post("/api/export/backup-policy-docx", dependencies=[Depends(scope_guard)])
async def export_backup_policy_docx(request: Request, db: Session = Depends(get_db)):
    from fastapi.responses import Response as FastAPIResponse
    from app.services.docx_service import generate_backup_policy_docx

    body        = await request.json()
    client_id   = body.get("client_id")
    job_names   = _job_names_from_body(body)
    client_name = (body.get("client_name") or "").strip()
    tlp         = (body.get("tlp") or "AMBER").upper()
    if tlp not in ("CLEAR", "GREEN", "AMBER", "RED"):
        tlp = "AMBER"

    if not client_id or not job_names or not client_name:
        raise HTTPException(status_code=400,
                            detail="Informe ambiente, ao menos uma rotina e o nome do cliente.")

    editorial = {
        "client_name":        client_name,
        "tlp":                tlp,
        "rto":                body.get("rto"),
        "rpo":                body.get("rpo"),
        "restore_test_freq":  body.get("restore_test_freq"),
        "review_cycle":       body.get("review_cycle"),
        "responsible_client": body.get("responsible_client"),
        "legal_retention":    body.get("legal_retention"),
        "doc_version":        body.get("doc_version"),
    }
    author = (request.session.get("full_name")
              or request.session.get("username") or None)
    payload = _backup_policy_payload(db, client_id, job_names, editorial, author=author)
    if payload is None:
        raise HTTPException(status_code=400, detail="Informe ao menos uma rotina.")
    if payload.get("missing_config"):
        faltando = ", ".join(payload["missing_config"])
        raise HTTPException(
            status_code=404,
            detail=f"Sem coleta de configuração para: {faltando}. Importe a "
                   "Auditoria de Configuração dessas rotinas antes de gerar a política.")

    _branding = branding_service.get_branding(client_id=client_id)
    data = generate_backup_policy_docx(payload, branding=_branding,
                                       branding_dir=str(BRANDING_DIR))

    slug = "".join(ch if ch.isalnum() else "_" for ch in client_name)[:40]
    fname = f"Politica_Backup_{slug}_v{payload['doc_version'].replace('.', '_')}.docx"
    return FastAPIResponse(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@app.get("/dashboard/monthly", response_class=HTMLResponse)
def dashboard_monthly_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse("dashboard_monthly.html", {"request": request, "clients": clients})


@app.get("/api/monthly-report-data", dependencies=[Depends(scope_guard)])
def api_monthly_report_data(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    client_id:  Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final para consultar.", "code": "dates_required"},
            status_code=400,
        )
    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)
    if not dt_start or not dt_end:
        return JSONResponse({"error": "Datas inválidas.", "code": "dates_invalid"}, status_code=400)

    payload = _monthly_report_payload(db, dt_start, dt_end, client_id)
    payload["clients"] = _all_clients_list(db)
    return JSONResponse(payload)


@app.post("/api/export/monthly-pptx", dependencies=[Depends(scope_guard)])
async def export_monthly_pptx(request: Request, db: Session = Depends(get_db)):
    from fastapi.responses import Response as FastAPIResponse
    from app.services.pptx_service import generate_monthly_backup_pptx

    body      = await request.json()
    dt_start  = _parse_date(body.get("date_start")) if body.get("date_start") else None
    dt_end    = _parse_date(body.get("date_end"), end_of_day=True) if body.get("date_end") else None
    client_id = body.get("client_id")

    p   = _monthly_report_payload(db, dt_start, dt_end, client_id)
    dist = p["backup_distribution"]
    kpis = {
        "failed":      dist["failed_count"],
        "success":     dist["success_count"],
        "failed_pct":  dist["failed_percent"],
        "success_pct": dist["success_percent"],
        "total":       dist["total_executions"],
    }

    _branding = branding_service.get_branding(client_id=client_id)
    pptx_bytes = generate_monthly_backup_pptx(
        client_name=p["client_name"], period=p["period"],
        kpis=kpis, protection_kpis=p["kpis"], top_failed_jobs=p["top_failed_jobs"],
        branding=_branding, branding_dir=str(BRANDING_DIR),
    )
    safe_name = (p["client_name"] or "todos").replace(" ", "_")
    return FastAPIResponse(
        content=pptx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="Relatorio_Mensal_{safe_name}.pptx"'},
    )


# ══════════════════════════════════════════════════════════════════════════════
# DASHBOARD OFFLOAD / CAPACITY TIER
# ══════════════════════════════════════════════════════════════════════════════

def _offload_report_payload(db: Session, dt_start, dt_end, client_id) -> dict:
    """Calcula KPIs, top rotinas com falha e distribuição por categoria de offload."""
    from app.models import OffloadSession, OffloadJob, OffloadReasonCategory

    q = db.query(OffloadSession)
    if client_id:
        q = q.filter(OffloadSession.client_id == client_id)
    if dt_start:
        q = q.filter(OffloadSession.started_at >= dt_start)
    if dt_end:
        q = q.filter(OffloadSession.started_at <= dt_end)

    sessions = q.all()

    total       = len(sessions)
    failed      = sum(1 for s in sessions if s.result == "Failed")
    warning     = sum(1 for s in sessions if s.result == "Warning")
    success     = sum(1 for s in sessions if s.result == "Success")
    unique_jobs = len({s.offload_job_id for s in sessions})
    failed_pct  = round(failed / total * 100, 1) if total else 0.0

    kpis = dict(total=total, unique_jobs=unique_jobs,
                failed=failed, warning=warning, success=success, failed_pct=failed_pct)

    # Top rotinas com mais falhas
    job_fail_cnt: dict  = {}   # job_id → count
    job_last_fail: dict = {}   # job_id → last started_at
    job_cat: dict       = {}   # job_id → reason_category_id (first seen)
    for s in sessions:
        if s.result != "Failed":
            continue
        jid = s.offload_job_id
        job_fail_cnt[jid]  = job_fail_cnt.get(jid, 0) + 1
        if s.started_at and (jid not in job_last_fail or s.started_at > job_last_fail[jid]):
            job_last_fail[jid] = s.started_at
        if s.reason_category_id and jid not in job_cat:
            job_cat[jid] = s.reason_category_id

    job_ids  = list(job_fail_cnt.keys())
    jobs_map = {j.id: j for j in db.query(OffloadJob).filter(OffloadJob.id.in_(job_ids)).all()} if job_ids else {}

    all_cat_ids = list({v for v in job_cat.values()})
    cats_map: dict = {}
    if all_cat_ids:
        for c in db.query(OffloadReasonCategory).filter(OffloadReasonCategory.id.in_(all_cat_ids)).all():
            cats_map[c.id] = c.name

    top_failed_jobs = []
    for jid, cnt in sorted(job_fail_cnt.items(), key=lambda x: -x[1])[:10]:
        job = jobs_map.get(jid)
        if not job:
            continue
        cat_id = job_cat.get(jid)
        lf     = job_last_fail.get(jid)
        top_failed_jobs.append({
            "job_name":    job.job_name,
            "failed_count": cnt,
            "last_failed": lf.strftime("%d/%m/%Y %H:%M") if lf else "—",
            "category":    cats_map.get(cat_id, "Não classificado") if cat_id else "—",
        })

    # Distribuição por categoria (falhas)
    cat_cnt: dict = {}
    for s in sessions:
        if s.result == "Failed" and s.reason_category_id:
            cat_cnt[s.reason_category_id] = cat_cnt.get(s.reason_category_id, 0) + 1

    by_category = []
    all_cat_ids2 = list(cat_cnt.keys())
    if all_cat_ids2:
        cats_map2: dict = {}
        for c in db.query(OffloadReasonCategory).filter(OffloadReasonCategory.id.in_(all_cat_ids2)).all():
            cats_map2[c.id] = c.name
        for cid, cnt in sorted(cat_cnt.items(), key=lambda x: -x[1])[:7]:
            by_category.append({
                "category": cats_map2.get(cid, "Desconhecida"),
                "count":    cnt,
                "pct":      round(cnt / failed * 100, 1) if failed else 0.0,
            })

    # ── Série diária de falhas/avisos (gráfico de onda) ───────────────────────
    daily_failed: dict  = {}
    daily_warning: dict = {}
    for s in sessions:
        if not s.started_at:
            continue
        d = s.started_at.date()
        if s.result == "Failed":
            daily_failed[d]  = daily_failed.get(d, 0) + 1
        elif s.result == "Warning":
            daily_warning[d] = daily_warning.get(d, 0) + 1

    # Limites do eixo: período informado tem prioridade; senão usa o range dos dados
    d0 = dt_start.date() if dt_start else None
    d1 = dt_end.date()   if dt_end   else None
    if (d0 is None or d1 is None):
        present = sorted(set(daily_failed) | set(daily_warning))
        if present:
            d0 = d0 or present[0]
            d1 = d1 or present[-1]

    daily_series = []
    if d0 and d1 and d0 <= d1:
        # série contínua dia-a-dia (zeros incluídos) — limitada a 400 dias
        span = (d1 - d0).days
        if span <= 400:
            cur = d0
            while cur <= d1:
                daily_series.append({
                    "label":   cur.strftime("%d/%m"),
                    "date":    cur.strftime("%Y-%m-%d"),
                    "failed":  daily_failed.get(cur, 0),
                    "warning": daily_warning.get(cur, 0),
                })
                cur += timedelta(days=1)
        else:
            # range muito grande: só os dias com ocorrência (evita milhares de pontos)
            for cur in sorted(set(daily_failed) | set(daily_warning)):
                daily_series.append({
                    "label":   cur.strftime("%d/%m/%y"),
                    "date":    cur.strftime("%Y-%m-%d"),
                    "failed":  daily_failed.get(cur, 0),
                    "warning": daily_warning.get(cur, 0),
                })

    client_name = ""
    if client_id:
        c = db.query(Client).get(client_id)
        if c:
            client_name = c.name

    period_str = "—"
    if dt_start and dt_end:
        period_str = f"{dt_start.strftime('%d/%m/%Y')} → {dt_end.strftime('%d/%m/%Y')}"

    return dict(kpis=kpis, top_failed_jobs=top_failed_jobs, by_category=by_category,
                daily_series=daily_series, period=period_str, client_name=client_name)


@app.get("/dashboard/offload", response_class=HTMLResponse)
def dashboard_offload_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse("dashboard_offload.html", {"request": request, "clients": clients})


@app.get("/api/offload-report-data", dependencies=[Depends(scope_guard)])
def api_offload_report_data(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    client_id:  Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final.", "code": "dates_required"},
            status_code=400,
        )
    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)
    if not dt_start or not dt_end:
        return JSONResponse({"error": "Datas inválidas.", "code": "dates_invalid"}, status_code=400)

    payload = _offload_report_payload(db, dt_start, dt_end, client_id)
    payload["clients"] = _all_clients_list(db)
    return JSONResponse(payload)


@app.post("/api/export/offload-pptx", dependencies=[Depends(scope_guard)])
async def export_offload_pptx(request: Request, db: Session = Depends(get_db)):
    from fastapi.responses import Response as FastAPIResponse
    from app.services.pptx_service import generate_offload_pptx

    body      = await request.json()
    dt_start  = _parse_date(body.get("date_start")) if body.get("date_start") else None
    dt_end    = _parse_date(body.get("date_end"), end_of_day=True) if body.get("date_end") else None
    client_id = body.get("client_id")

    p       = _offload_report_payload(db, dt_start, dt_end, client_id)
    _brand  = branding_service.get_branding(client_id=client_id)
    pptx_bytes = generate_offload_pptx(
        client_name=p["client_name"], period=p["period"],
        kpis=p["kpis"], top_failed_jobs=p["top_failed_jobs"],
        by_category=p["by_category"], daily_series=p.get("daily_series"),
        branding=_brand, branding_dir=str(BRANDING_DIR),
    )
    safe_name = (p["client_name"] or "todos").replace(" ", "_")
    return FastAPIResponse(
        content=pptx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="Relatorio_Offload_{safe_name}.pptx"'},
    )


# ═══════════════════════════════════════════════════════════════════════════════
# Relatório gerencial de Performance de Backups (espelha o de Offload)
# ═══════════════════════════════════════════════════════════════════════════════

def _backup_report_payload(db: Session, dt_start, dt_end, client_id) -> dict:
    """KPIs, rotinas mais lentas e rotinas com mais falhas (performance de backup)."""
    from sqlalchemy import func, case
    from app.models import BackupVmSession, BackupRoutine

    q = db.query(BackupVmSession)
    if client_id:
        q = q.filter(BackupVmSession.client_id == client_id)
    if dt_start:
        q = q.filter(BackupVmSession.started_at >= dt_start)
    if dt_end:
        q = q.filter(BackupVmSession.started_at <= dt_end)
    sessions = q.all()

    total       = len(sessions)
    failed      = sum(1 for s in sessions if s.result == "Failed")
    warning     = sum(1 for s in sessions if s.result == "Warning")
    success     = sum(1 for s in sessions if s.result == "Success")
    unique_rout = len({s.backup_routine_id for s in sessions})
    failed_pct  = round(failed / total * 100, 1) if total else 0.0
    durs        = [s.duration_seconds for s in sessions if s.duration_seconds is not None]
    avg_dur     = int(sum(durs) / len(durs)) if durs else None

    def _dur_fmt(secs):
        if secs is None:
            return "—"
        secs = int(secs); h, rem = divmod(secs, 3600); m, _ = divmod(rem, 60)
        return f"{h}h{m:02d}m" if h else f"{m}m"

    kpis = dict(total=total, unique_routines=unique_rout, failed=failed,
                warning=warning, success=success, failed_pct=failed_pct,
                avg_duration_fmt=_dur_fmt(avg_dur))

    # ── Rotinas mais LENTAS (mediana de duração na janela) ────────────────────
    med_q = (db.query(
                BackupVmSession.backup_routine_id.label("rid"),
                func.percentile_cont(0.5).within_group(BackupVmSession.duration_seconds.asc()).label("med"),
                func.count().label("n"),
                func.max(BackupVmSession.duration_seconds).label("mx"))
             .filter(BackupVmSession.duration_seconds.isnot(None)))
    if client_id:
        med_q = med_q.filter(BackupVmSession.backup_routine_id.in_(
            [r.id for r in db.query(BackupRoutine.id).filter(BackupRoutine.client_id == client_id)]))
    if dt_start:
        med_q = med_q.filter(BackupVmSession.started_at >= dt_start)
    if dt_end:
        med_q = med_q.filter(BackupVmSession.started_at <= dt_end)
    med_rows = med_q.group_by(BackupVmSession.backup_routine_id).all()

    rmap = {r.id: r for r in db.query(BackupRoutine).filter(
        BackupRoutine.id.in_([m.rid for m in med_rows])).all()} if med_rows else {}

    top_slowest = []
    for m in sorted(med_rows, key=lambda x: -(x.med or 0))[:10]:
        rt = rmap.get(m.rid)
        if not rt:
            continue
        top_slowest.append({
            "vm_name":  rt.vm_name,
            "job_name": rt.job_name,
            "median_fmt": _dur_fmt(m.med),
            "max_fmt":    _dur_fmt(m.mx),
            "executions": m.n,
        })

    # ── Rotinas com mais FALHAS ───────────────────────────────────────────────
    fail_cnt, last_fail = {}, {}
    for s in sessions:
        if s.result != "Failed":
            continue
        fail_cnt[s.backup_routine_id] = fail_cnt.get(s.backup_routine_id, 0) + 1
        if s.started_at and (s.backup_routine_id not in last_fail or s.started_at > last_fail[s.backup_routine_id]):
            last_fail[s.backup_routine_id] = s.started_at
    fmap = {r.id: r for r in db.query(BackupRoutine).filter(
        BackupRoutine.id.in_(list(fail_cnt.keys()))).all()} if fail_cnt else {}
    top_failed = []
    for rid, cnt in sorted(fail_cnt.items(), key=lambda x: -x[1])[:10]:
        rt = fmap.get(rid)
        if not rt:
            continue
        lf = last_fail.get(rid)
        top_failed.append({
            "vm_name": rt.vm_name, "job_name": rt.job_name, "failed_count": cnt,
            "last_failed": lf.strftime("%d/%m/%Y %H:%M") if lf else "—",
        })

    # ── Série diária de falhas (gráfico de onda) ──────────────────────────────
    daily_failed, daily_warning = {}, {}
    for s in sessions:
        if not s.started_at:
            continue
        d = s.started_at.date()
        if s.result == "Failed":
            daily_failed[d] = daily_failed.get(d, 0) + 1
        elif s.result == "Warning":
            daily_warning[d] = daily_warning.get(d, 0) + 1
    d0 = dt_start.date() if dt_start else None
    d1 = dt_end.date()   if dt_end   else None
    if d0 is None or d1 is None:
        present = sorted(set(daily_failed) | set(daily_warning))
        if present:
            d0 = d0 or present[0]; d1 = d1 or present[-1]
    daily_series = []
    if d0 and d1 and d0 <= d1 and (d1 - d0).days <= 400:
        cur = d0
        while cur <= d1:
            daily_series.append({"label": cur.strftime("%d/%m"), "date": cur.strftime("%Y-%m-%d"),
                                 "failed": daily_failed.get(cur, 0), "warning": daily_warning.get(cur, 0)})
            cur += timedelta(days=1)

    client_name = ""
    if client_id:
        c = db.query(Client).get(client_id)
        if c:
            client_name = c.name
    period_str = "—"
    if dt_start and dt_end:
        period_str = f"{dt_start.strftime('%d/%m/%Y')} → {dt_end.strftime('%d/%m/%Y')}"

    return dict(kpis=kpis, top_slowest=top_slowest, top_failed=top_failed,
                daily_series=daily_series, period=period_str, client_name=client_name)


@app.get("/dashboard/backups", response_class=HTMLResponse)
def dashboard_backups_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse("dashboard_backups.html", {"request": request, "clients": clients})


@app.get("/api/backup-report-data", dependencies=[Depends(scope_guard)])
def api_backup_report_data(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    client_id:  Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    if not date_start or not date_end:
        return JSONResponse({"error": "Informe data inicial e data final.", "code": "dates_required"}, status_code=400)
    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)
    if not dt_start or not dt_end:
        return JSONResponse({"error": "Datas inválidas.", "code": "dates_invalid"}, status_code=400)
    payload = _backup_report_payload(db, dt_start, dt_end, client_id)
    payload["clients"] = _all_clients_list(db)
    return JSONResponse(payload)


@app.post("/api/export/backup-pptx", dependencies=[Depends(scope_guard)])
async def export_backup_pptx(request: Request, db: Session = Depends(get_db)):
    from fastapi.responses import Response as FastAPIResponse
    from app.services.pptx_service import generate_backup_pptx

    body      = await request.json()
    dt_start  = _parse_date(body.get("date_start")) if body.get("date_start") else None
    dt_end    = _parse_date(body.get("date_end"), end_of_day=True) if body.get("date_end") else None
    client_id = body.get("client_id")

    p      = _backup_report_payload(db, dt_start, dt_end, client_id)
    _brand = branding_service.get_branding(client_id=client_id)
    pptx_bytes = generate_backup_pptx(
        client_name=p["client_name"], period=p["period"],
        kpis=p["kpis"], top_slowest=p["top_slowest"], top_failed=p["top_failed"],
        daily_series=p.get("daily_series"), branding=_brand, branding_dir=str(BRANDING_DIR),
    )
    safe_name = (p["client_name"] or "todos").replace(" ", "_")
    return FastAPIResponse(
        content=pptx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="Relatorio_Backup_{safe_name}.pptx"'},
    )


@app.get("/dashboard/audit", response_class=HTMLResponse)
def dashboard_audit_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse("dashboard_audit.html", {"request": request, "clients": clients})


@app.get("/api/audit-dashboard", dependencies=[Depends(scope_guard)])
def api_audit_dashboard(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    event_type: str            = Query("all"),
    client_id:  Optional[int] = Query(None),
    job_name:   Optional[str] = Query(None),
    operator:   Optional[str] = Query(None),
    vbr_host:   Optional[str] = Query(None),
    page:       int            = Query(1, ge=1),
    page_size:  int            = Query(20, ge=1, le=200),
    db: Session = Depends(get_db),
):
    from sqlalchemy import func, cast, extract, or_
    from sqlalchemy.types import Date as SADate

    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final para consultar.", "code": "dates_required"},
            status_code=400,
        )

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    # ── Query base com todos os filtros ───────────────────────────────────────
    q_base = db.query(ImportedEvent).filter(ImportedEvent.event_category == "audit")
    if client_id:
        q_base = q_base.filter(ImportedEvent.client_id == client_id)
    if dt_start:
        q_base = q_base.filter(ImportedEvent.time_created >= dt_start)
    if dt_end:
        q_base = q_base.filter(ImportedEvent.time_created <= dt_end)
    if event_type != "all":
        q_base = q_base.filter(ImportedEvent.audit_event_type == event_type)
    if job_name:
        q_base = q_base.filter(ImportedEvent.job_name.ilike(f"%{job_name}%"))
    if operator:
        q_base = q_base.filter(ImportedEvent.operator.ilike(f"%{operator}%"))
    if vbr_host:
        q_base = q_base.filter(ImportedEvent.vbr_hostname.ilike(f"%{vbr_host}%"))

    total_count = q_base.count()

    AUDIT_TYPES = ["job_created", "job_updated", "job_deleted",
                   "objects_added", "objects_changed", "objects_deleted"]
    TYPE_COLORS_HEX = {
        "job_created":     "#7C3AED",
        "job_updated":     "#F97316",
        "job_deleted":     "#DC3545",
        "objects_added":   "#A78BFA",
        "objects_changed": "#FFC107",
        "objects_deleted": "#fd7e14",
    }
    LABELS = {
        "job_created":     "Job Criado",
        "job_updated":     "Config. Alteradas",
        "job_deleted":     "Job Excluído",
        "objects_added":   "Objetos Adicionados",
        "objects_changed": "Objetos Alterados",
        "objects_deleted": "Objetos Removidos",
    }

    # ── KPIs via SQL GROUP BY — correto para qualquer volume ──────────────────
    type_counts: dict = {
        (k or "unknown"): v
        for k, v in (
            q_base
            .with_entities(ImportedEvent.audit_event_type, func.count(ImportedEvent.id))
            .group_by(ImportedEvent.audit_event_type)
            .all()
        )
    }

    operators_count: int = (
        q_base
        .filter(ImportedEvent.operator.isnot(None))
        .with_entities(func.count(func.distinct(ImportedEvent.operator)))
        .scalar() or 0
    )

    kpis = {
        "total":           total_count,
        "job_created":     type_counts.get("job_created",     0),
        "job_updated":     type_counts.get("job_updated",     0),
        "job_deleted":     type_counts.get("job_deleted",     0),
        "objects_added":   type_counts.get("objects_added",   0),
        "objects_changed": type_counts.get("objects_changed", 0),
        "objects_deleted": type_counts.get("objects_deleted", 0),
        "operators_count": operators_count,
    }

    # ── Timeline por dia × tipo via SQL GROUP BY ──────────────────────────────
    day_col = cast(ImportedEvent.time_created, SADate)
    day_type_rows = (
        q_base
        .filter(ImportedEvent.time_created.isnot(None))
        .with_entities(
            day_col.label("day"),
            ImportedEvent.audit_event_type,
            func.count(ImportedEvent.id),
        )
        .group_by(day_col, ImportedEvent.audit_event_type)
        .order_by(day_col)
        .all()
    )

    by_day_type: dict = defaultdict(lambda: defaultdict(int))
    for row in day_type_rows:
        day_str = row[0].strftime("%Y-%m-%d") if row[0] else None
        if day_str:
            by_day_type[day_str][row[1] or "unknown"] += row[2]

    day_keys = sorted(by_day_type.keys())

    timeline_datasets = []
    for t in AUDIT_TYPES:
        counts = [by_day_type[d].get(t, 0) for d in day_keys]
        if any(c > 0 for c in counts):
            timeline_datasets.append({
                "label": LABELS.get(t, t),
                "type":  t,
                "data":  counts,
                "color": TYPE_COLORS_HEX.get(t, "#718096"),
            })

    # ── Operator ranking via SQL GROUP BY ─────────────────────────────────────
    top_operators = (
        q_base
        .filter(ImportedEvent.operator.isnot(None))
        .with_entities(
            ImportedEvent.operator,
            func.count(ImportedEvent.id).label("cnt"),
        )
        .group_by(ImportedEvent.operator)
        .order_by(func.count(ImportedEvent.id).desc())
        .limit(10)
        .all()
    )

    # ── Top jobs via SQL GROUP BY ──────────────────────────────────────────────
    top_jobs = (
        q_base
        .filter(ImportedEvent.job_name.isnot(None))
        .with_entities(
            ImportedEvent.job_name,
            func.count(ImportedEvent.id).label("cnt"),
        )
        .group_by(ImportedEvent.job_name)
        .order_by(func.count(ImportedEvent.id).desc())
        .limit(15)
        .all()
    )

    def _trunc(s: str, n: int = 40) -> str:
        return (s[:n] + "…") if len(s) > n else s

    # ── Cross-reference: jobs com falhas de backup no mesmo período ───────────
    failed_job_norms: set = set()
    if dt_start and dt_end:
        fq = (
            db.query(ImportedEvent.job_name_normalized)
            .filter(
                ImportedEvent.event_category.in_(["backup", "log_backup"]),
                ImportedEvent.job_result == 2,
                ImportedEvent.time_created >= dt_start,
                ImportedEvent.time_created <= dt_end,
            )
        )
        if client_id:
            fq = fq.filter(ImportedEvent.client_id == client_id)
        failed_job_norms = {row[0] for row in fq.distinct().all() if row[0]}

    def _classify_impact(e) -> str:
        after_hours = bool(e.time_created and (e.time_created.hour < 7 or e.time_created.hour >= 19))
        has_failure = bool((e.job_name_normalized or "") in failed_job_norms)
        if e.audit_event_type in ("job_deleted", "objects_deleted") or has_failure:
            return "Alto"
        if e.audit_event_type in ("objects_added", "objects_changed", "job_updated") or after_hours:
            return "Médio"
        return "Baixo"

    def _fmt_event(e) -> dict:
        after_hours = bool(e.time_created and (e.time_created.hour < 7 or e.time_created.hour >= 19))
        return {
            "id":                e.id,
            "time":              e.time_created.strftime("%d/%m/%Y %H:%M") if e.time_created else "—",
            "time_iso":          e.time_created.strftime("%Y-%m-%dT%H:%M")  if e.time_created else "",
            "event_id":          e.event_id,
            "event_type":        e.audit_event_type  or "—",
            "event_label":       e.audit_event_label or "—",
            "job_name":          e.job_name          or "—",
            "details":           (e.audit_details or "")[:200],
            "operator":          e.operator          or "—",
            "vbr_host":          e.vbr_hostname      or "—",
            "after_hours":       after_hours,
            "has_failure_xref":  (e.job_name_normalized or "") in failed_job_norms,
            "impact":            _classify_impact(e),
            "upload_session_id": e.upload_session_id,
        }

    # ── Alterações sensíveis via SQL (max 50) — não usa amostra limitada ──────
    sensitive_conds = [
        ImportedEvent.audit_event_type.in_(["job_deleted", "objects_deleted"]),
        extract("hour", ImportedEvent.time_created) < 7,
        extract("hour", ImportedEvent.time_created) >= 19,
    ]
    if failed_job_norms:
        sensitive_conds.append(ImportedEvent.job_name_normalized.in_(failed_job_norms))

    sensitive_evs = (
        q_base
        .filter(or_(*sensitive_conds))
        .order_by(ImportedEvent.time_created.desc())
        .limit(50)
        .all()
    )
    sensitive_list = [_fmt_event(e) for e in sensitive_evs]

    # ── Paginação server-side ─────────────────────────────────────────────────
    ps          = max(1, page_size)
    total_pages = max(1, (total_count + ps - 1) // ps)
    cur_page    = max(1, min(page, total_pages))
    offset      = (cur_page - 1) * ps

    paged_evs = (
        q_base
        .order_by(ImportedEvent.time_created.desc())
        .offset(offset)
        .limit(ps)
        .all()
    )
    events_list = [_fmt_event(e) for e in paged_evs]

    return JSONResponse({
        "kpis":              kpis,
        "events_by_day": {
            "labels":   [datetime.strptime(d, "%Y-%m-%d").strftime("%d/%m") for d in day_keys],
            "datasets": timeline_datasets,
        },
        "type_distribution": {
            "labels": [LABELS.get(k, k) for k in type_counts],
            "values": list(type_counts.values()),
            "types":  list(type_counts.keys()),
        },
        "top_jobs": {
            "labels": [_trunc(j[0]) for j in top_jobs],
            "counts": [j[1] for j in top_jobs],
        },
        "operator_ranking": {
            "labels": [_trunc(o[0], 30) for o in top_operators],
            "counts": [o[1] for o in top_operators],
        },
        "sensitive_changes": sensitive_list,
        "events":            events_list,
        "pagination": {
            "page":        cur_page,
            "page_size":   ps,
            "total":       total_count,
            "total_pages": total_pages,
        },
        "reports": _all_reports_list(db, client_id=client_id),
        "clients": _all_clients_list(db),
    })


# ── API: disponibilidade de auditoria por cliente ────────────────────────────

@app.get("/api/audit-availability", dependencies=[Depends(scope_guard)])
def api_audit_availability(
    client_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    """
    Retorna o período disponível de eventos de auditoria na base histórica.
    Usado pelo dashboard de auditoria para sugerir datas ao usuário.
    """
    from sqlalchemy import func

    q = db.query(ImportedEvent).filter(ImportedEvent.event_category == "audit")
    if client_id:
        q = q.filter(ImportedEvent.client_id == client_id)

    row = q.with_entities(
        func.count(ImportedEvent.id).label("total"),
        func.min(ImportedEvent.time_created).label("min_date"),
        func.max(ImportedEvent.time_created).label("max_date"),
    ).first()

    if not row or not row[0]:
        return JSONResponse({"total": 0, "min_date": None, "max_date": None,
                             "min_date_fmt": None, "max_date_fmt": None,
                             "type_counts": {}})

    # Contagem por tipo
    type_rows = (
        q
        .with_entities(
            ImportedEvent.audit_event_type,
            func.count(ImportedEvent.id),
        )
        .group_by(ImportedEvent.audit_event_type)
        .all()
    )
    type_counts = {(k or "unknown"): v for k, v in type_rows}

    return JSONResponse({
        "total":        row[0],
        "min_date":     row[1].strftime("%Y-%m-%d") if row[1] else None,
        "max_date":     row[2].strftime("%Y-%m-%d") if row[2] else None,
        "min_date_fmt": row[1].strftime("%d/%m/%Y") if row[1] else None,
        "max_date_fmt": row[2].strftime("%d/%m/%Y") if row[2] else None,
        "type_counts":  type_counts,
    })


# ── Diagnóstico administrativo de auditoria ───────────────────────────────────

@app.get("/admin/audit-diagnostic", response_class=HTMLResponse)
def admin_audit_diagnostic_page(
    request: Request,
    db: Session = Depends(get_db),
):
    from app.auth import require_admin
    require_admin(request)
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse(
        "admin_audit_diagnostic.html",
        {"request": request, "clients": clients},
    )


@app.get("/api/admin/audit-diagnostic", dependencies=[Depends(require_admin)])
def api_admin_audit_diagnostic(
    client_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    """
    Diagnóstico completo de eventos de auditoria no banco.
    Retorna estatísticas por cliente, event_id e upload session.
    """
    from sqlalchemy import func, text

    AUDIT_LABEL = {
        "job_created":     "Job Criado (23010)",
        "job_updated":     "Config. Alteradas (23050)",
        "job_deleted":     "Job Excluído (23090)",
        "objects_added":   "Obj. Adicionados (23110)",
        "objects_changed": "Obj. Alterados (23130)",
        "objects_deleted": "Obj. Removidos (32120)",
    }

    clients = db.query(Client).order_by(Client.name).all()
    result_clients = []

    for c in clients:
        if client_id and c.id != client_id:
            continue

        q = db.query(ImportedEvent).filter(
            ImportedEvent.event_category == "audit",
            ImportedEvent.client_id == c.id,
        )

        stats = q.with_entities(
            func.count(ImportedEvent.id).label("total"),
            func.min(ImportedEvent.time_created).label("min_date"),
            func.max(ImportedEvent.time_created).label("max_date"),
        ).first()

        if not stats or not stats[0]:
            result_clients.append({
                "client_id":   c.id,
                "client_name": c.name,
                "total":       0,
                "min_date":    None,
                "max_date":    None,
                "by_type":     [],
                "by_event_id": [],
                "last_upload": None,
            })
            continue

        # Por tipo
        by_type = [
            {
                "audit_event_type":  k or "unknown",
                "label":             AUDIT_LABEL.get(k or "", k or "unknown"),
                "count":             v,
            }
            for k, v in (
                q.with_entities(
                    ImportedEvent.audit_event_type,
                    func.count(ImportedEvent.id),
                )
                .group_by(ImportedEvent.audit_event_type)
                .order_by(func.count(ImportedEvent.id).desc())
                .all()
            )
        ]

        # Por event_id
        by_eid = [
            {"event_id": k, "count": v}
            for k, v in (
                q.with_entities(
                    ImportedEvent.event_id,
                    func.count(ImportedEvent.id),
                )
                .group_by(ImportedEvent.event_id)
                .order_by(ImportedEvent.event_id)
                .all()
            )
        ]

        # Último upload com auditoria
        last_up = (
            db.query(UploadSession)
            .filter(
                UploadSession.client_id == c.id,
                UploadSession.audit_events_count > 0,
                UploadSession.status == "done",
            )
            .order_by(UploadSession.uploaded_at.desc())
            .first()
        )

        result_clients.append({
            "client_id":   c.id,
            "client_name": c.name,
            "total":       stats[0],
            "min_date":    stats[1].strftime("%d/%m/%Y %H:%M") if stats[1] else None,
            "max_date":    stats[2].strftime("%d/%m/%Y %H:%M") if stats[2] else None,
            "min_date_iso": stats[1].strftime("%Y-%m-%d") if stats[1] else None,
            "max_date_iso": stats[2].strftime("%Y-%m-%d") if stats[2] else None,
            "by_type":     by_type,
            "by_event_id": by_eid,
            "last_upload": {
                "id":            last_up.id,
                "filename":      last_up.original_filename,
                "uploaded_at":   last_up.uploaded_at.strftime("%d/%m/%Y %H:%M"),
                "uploaded_by":   last_up.uploaded_by or "—",
                "audit_count":   last_up.audit_events_count,
                "period_start":  last_up.period_start.strftime("%d/%m/%Y") if last_up.period_start else "—",
                "period_end":    last_up.period_end.strftime("%d/%m/%Y")   if last_up.period_end   else "—",
            } if last_up else None,
        })

    return JSONResponse({"clients": result_clients})


# ── API: detalhe de evento de auditoria (para modal) ──────────────────────────

@app.get("/api/audit-event/{event_db_id}")
def api_audit_event_detail(event_db_id: int, db: Session = Depends(get_db)):
    e = db.query(ImportedEvent).filter(
        ImportedEvent.id == event_db_id,
        ImportedEvent.event_category == "audit",
    ).first()
    if not e:
        return JSONResponse({"error": "Evento não encontrado."}, status_code=404)

    session_info: dict = {}
    if e.upload_session_id:
        us = db.query(UploadSession).get(e.upload_session_id)
        if us:
            session_info = {
                "id":          us.id,
                "filename":    us.original_filename or "—",
                "uploaded_at": us.uploaded_at.strftime("%d/%m/%Y %H:%M") if us.uploaded_at else "—",
                "uploaded_by": us.uploaded_by or "—",
            }

    return JSONResponse({
        "id":               e.id,
        "event_id":         e.event_id,
        "time":             e.time_created.strftime("%d/%m/%Y %H:%M:%S") if e.time_created else "—",
        "event_type":       e.audit_event_type  or "—",
        "event_label":      e.audit_event_label or "—",
        "job_name":         e.job_name          or "—",
        "operator":         e.operator          or "—",
        "vbr_host":         e.vbr_hostname      or "—",
        "details":          e.audit_details     or "—",
        "raw_xml":          e.raw_event_xml     or "",
        "original_message": e.original_message  or "",
        "provider_name":    e.provider_name     or "—",
        "record_id":        e.record_id,
        "source_filename":  e.source_filename   or "—",
        "imported_at":      e.imported_at.strftime("%d/%m/%Y %H:%M") if e.imported_at else "—",
        "upload_session":   session_info,
    })


# ── API: exportar auditoria para Excel ────────────────────────────────────────

@app.get("/api/export/audit-excel", dependencies=[Depends(scope_guard)])
def api_export_audit_excel(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    event_type: str            = Query("all"),
    client_id:  Optional[int] = Query(None),
    job_name:   Optional[str] = Query(None),
    operator:   Optional[str] = Query(None),
    vbr_host:   Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    if not date_start or not date_end:
        return JSONResponse({"error": "Datas obrigatórias."}, status_code=400)

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    q = db.query(ImportedEvent).filter(ImportedEvent.event_category == "audit")
    if client_id:
        q = q.filter(ImportedEvent.client_id == client_id)
    if dt_start:
        q = q.filter(ImportedEvent.time_created >= dt_start)
    if dt_end:
        q = q.filter(ImportedEvent.time_created <= dt_end)
    if event_type != "all":
        q = q.filter(ImportedEvent.audit_event_type == event_type)
    if job_name:
        q = q.filter(ImportedEvent.job_name.ilike(f"%{job_name}%"))
    if operator:
        q = q.filter(ImportedEvent.operator.ilike(f"%{operator}%"))
    if vbr_host:
        q = q.filter(ImportedEvent.vbr_hostname.ilike(f"%{vbr_host}%"))

    events = q.order_by(ImportedEvent.time_created.desc()).limit(5000).all()

    # Cross-reference com falhas
    failed_job_norms: set = set()
    if dt_start and dt_end:
        fq = (
            db.query(ImportedEvent.job_name_normalized)
            .filter(
                ImportedEvent.event_category.in_(["backup", "log_backup"]),
                ImportedEvent.job_result == 2,
                ImportedEvent.time_created >= dt_start,
                ImportedEvent.time_created <= dt_end,
            )
        )
        if client_id:
            fq = fq.filter(ImportedEvent.client_id == client_id)
        failed_job_norms = {row[0] for row in fq.distinct().all() if row[0]}

    client_name = "Todos os clientes"
    if client_id:
        c = db.query(Client).get(client_id)
        if c:
            client_name = c.name

    period_info = {
        "client_name": client_name,
        "date_start":  date_start,
        "date_end":    date_end,
    }

    from app.services.excel_service import generate_audit_excel
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
        xlsx_path = tf.name

    generate_audit_excel(events, period_info, failed_job_norms, xlsx_path)

    fname = f"auditoria_{date_start}_{date_end}.xlsx"
    return FileResponse(
        xlsx_path,
        filename=fname,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ── API: exportar auditoria para PowerPoint ───────────────────────────────────

@app.get("/api/export/audit-pptx", dependencies=[Depends(scope_guard)])
def api_export_audit_pptx(
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    event_type: str            = Query("all"),
    client_id:  Optional[int] = Query(None),
    job_name:   Optional[str] = Query(None),
    operator:   Optional[str] = Query(None),
    vbr_host:   Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    if not date_start or not date_end:
        return JSONResponse({"error": "Datas obrigatórias."}, status_code=400)

    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)

    q = db.query(ImportedEvent).filter(ImportedEvent.event_category == "audit")
    if client_id:
        q = q.filter(ImportedEvent.client_id == client_id)
    if dt_start:
        q = q.filter(ImportedEvent.time_created >= dt_start)
    if dt_end:
        q = q.filter(ImportedEvent.time_created <= dt_end)
    if event_type != "all":
        q = q.filter(ImportedEvent.audit_event_type == event_type)
    if job_name:
        q = q.filter(ImportedEvent.job_name.ilike(f"%{job_name}%"))
    if operator:
        q = q.filter(ImportedEvent.operator.ilike(f"%{operator}%"))
    if vbr_host:
        q = q.filter(ImportedEvent.vbr_hostname.ilike(f"%{vbr_host}%"))

    events = q.order_by(ImportedEvent.time_created.desc()).limit(5000).all()

    # KPIs
    type_counts: dict = defaultdict(int)
    operators_set: set = set()
    for e in events:
        type_counts[e.audit_event_type or "unknown"] += 1
        if e.operator:
            operators_set.add(e.operator)

    kpis = {
        "total":           len(events),
        "job_created":     type_counts["job_created"],
        "job_updated":     type_counts["job_updated"],
        "job_deleted":     type_counts["job_deleted"],
        "objects_added":   type_counts["objects_added"],
        "objects_changed": type_counts["objects_changed"],
        "objects_deleted": type_counts["objects_deleted"],
        "operators_count": len(operators_set),
    }

    # Failed jobs cross-reference
    failed_job_norms: set = set()
    if dt_start and dt_end:
        fq = (
            db.query(ImportedEvent.job_name_normalized)
            .filter(
                ImportedEvent.event_category.in_(["backup", "log_backup"]),
                ImportedEvent.job_result == 2,
                ImportedEvent.time_created >= dt_start,
                ImportedEvent.time_created <= dt_end,
            )
        )
        if client_id:
            fq = fq.filter(ImportedEvent.client_id == client_id)
        failed_job_norms = {row[0] for row in fq.distinct().all() if row[0]}

    def _is_sensitive(e) -> bool:
        if e.audit_event_type in ("job_deleted", "objects_deleted"):
            return True
        if e.time_created and (e.time_created.hour < 7 or e.time_created.hour >= 19):
            return True
        if (e.job_name_normalized or "") in failed_job_norms:
            return True
        return False

    def _classify_impact(e) -> str:
        after_hours = bool(e.time_created and (e.time_created.hour < 7 or e.time_created.hour >= 19))
        has_failure = bool((e.job_name_normalized or "") in failed_job_norms)
        if e.audit_event_type in ("job_deleted", "objects_deleted") or has_failure:
            return "Alto"
        if e.audit_event_type in ("objects_added", "objects_changed", "job_updated") or after_hours:
            return "Médio"
        return "Baixo"

    sensitive_items = [
        {
            "job_name":   e.job_name or "—",
            "event_type": e.audit_event_type or "—",
            "operator":   e.operator or "—",
            "time":       e.time_created.strftime("%d/%m %H:%M") if e.time_created else "—",
            "impact":     _classify_impact(e),
        }
        for e in events if _is_sensitive(e)
    ][:6]

    # Top operators
    op_count: dict = defaultdict(int)
    for e in events:
        if e.operator:
            op_count[e.operator] += 1
    top_operators = sorted(op_count.items(), key=lambda x: x[1], reverse=True)[:5]

    # Top jobs
    job_count: dict = defaultdict(int)
    for e in events:
        if e.job_name:
            job_count[e.job_name] += 1
    top_jobs = sorted(job_count.items(), key=lambda x: x[1], reverse=True)[:6]

    client_name = "Todos os clientes"
    if client_id:
        c = db.query(Client).get(client_id)
        if c:
            client_name = c.name

    period = f"{date_start} → {date_end}"

    from app.services.pptx_service import generate_audit_pptx
    _audit_branding = branding_service.get_branding(client_id=client_id)
    pptx_bytes = generate_audit_pptx(
        client_name    = client_name,
        period         = period,
        kpis           = kpis,
        top_operators  = top_operators,
        sensitive_items= sensitive_items,
        top_jobs       = top_jobs,
        branding       = _audit_branding,
        branding_dir   = str(BRANDING_DIR),
    )

    with tempfile.NamedTemporaryFile(suffix=".pptx", delete=False) as tf:
        tf.write(pptx_bytes)
        pptx_path = tf.name

    fname = f"auditoria_{date_start}_{date_end}.pptx"
    return FileResponse(
        pptx_path,
        filename=fname,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )


# ════════════════════════════════════════════════════════════════════════════════
# DETALHE DA ROTINA — SLA, RPO e drill-down de VMs por job específico
# ════════════════════════════════════════════════════════════════════════════════
#
# Granularidades:
#   • "Rotina" (job)  = job_name nos eventos de SESSÃO de job  (EventID 190–200).
#   • "VM"            = job_name nos eventos por-máquina        (EventID 150).
#
# Linkage job→VM: o EventID 150 carrega, no primeiro <Data> do EventData, o MESMO
# GUID de sessão do EventID 190 correspondente. Isso permite associar cada VM à
# execução da rotina sem heurística de nome.

_ROTINA_JOB_EVENT_MIN = 190
_ROTINA_JOB_EVENT_MAX = 200

# Regex Postgres: primeiro GUID dentro de <Data>…</Data> = GUID da sessão do job.
_SESSION_GUID_RE = (
    r'<Data>([0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-'
    r'[0-9a-fA-F]{4}-[0-9a-fA-F]{12})</Data>'
)


@app.get("/dashboard/rotina", response_class=HTMLResponse)
def dashboard_rotina_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return templates.TemplateResponse(
        "dashboard_rotina.html",
        {"request": request, "clients": clients},
    )


@app.get("/api/rotina-list", dependencies=[Depends(scope_guard)])
def api_rotina_list(
    client_id:  Optional[int] = Query(None),
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    search:     Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    """Lista as rotinas (job_name) de um cliente no período, com stats rápidas."""
    if not client_id:
        return JSONResponse(
            {"error": "Selecione um cliente.", "code": "client_required"}, status_code=400
        )
    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final.", "code": "dates_required"}, status_code=400
        )
    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)
    if not dt_start or not dt_end:
        return JSONResponse({"error": "Datas inválidas.", "code": "dates_invalid"}, status_code=400)

    events = (
        db.query(ImportedEvent)
        .filter(ImportedEvent.client_id == client_id)
        .filter(ImportedEvent.event_id >= _ROTINA_JOB_EVENT_MIN)
        .filter(ImportedEvent.event_id <= _ROTINA_JOB_EVENT_MAX)
        .filter(ImportedEvent.job_name.isnot(None))
        .filter(ImportedEvent.time_created >= dt_start)
        .filter(ImportedEvent.time_created <= dt_end)
        .order_by(ImportedEvent.time_created.desc())
        .all()
    )

    agg: dict = {}
    for e in events:
        jn = e.job_name
        if not jn:
            continue
        a = agg.setdefault(jn, {
            "job_name": jn,
            "job_name_normalized": e.job_name_normalized or jn,
            "job_type": e.event_category or "backup",
            "total": 0, "failed": 0, "warning": 0, "success": 0,
            "last_execution": None, "last_result": None,
        })
        a["total"] += 1
        if   e.job_result == 2: a["failed"]  += 1
        elif e.job_result == 1: a["warning"] += 1
        elif e.job_result == 0: a["success"] += 1
        if a["last_execution"] is None and e.time_created:
            a["last_execution"] = e.time_created.strftime("%d/%m/%Y %H:%M")
            a["last_result"]    = e.job_result

    rotinas = list(agg.values())
    if search:
        s = search.lower()
        rotinas = [r for r in rotinas if s in r["job_name"].lower()]
    for r in rotinas:
        r["success_rate"] = round(r["success"] / r["total"] * 100, 1) if r["total"] else 0.0
    rotinas.sort(key=lambda r: r["job_name"].lower())

    return JSONResponse({"rotinas": rotinas, "total": len(rotinas)})


@app.get("/api/rotina-detail", dependencies=[Depends(scope_guard)])
def api_rotina_detail(
    client_id:  int           = Query(...),
    job_name:   str           = Query(...),
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    """SLA + timeline + ofensores + histórico + gestão (JobAction) de uma rotina."""
    from collections import OrderedDict
    from app.parser.common import normalize_job_name
    from app.services.classification_service import auto_situation

    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final.", "code": "dates_required"}, status_code=400
        )
    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)
    if not dt_start or not dt_end:
        return JSONResponse({"error": "Datas inválidas.", "code": "dates_invalid"}, status_code=400)

    events = (
        db.query(ImportedEvent)
        .filter(ImportedEvent.client_id == client_id)
        .filter(ImportedEvent.event_id >= _ROTINA_JOB_EVENT_MIN)
        .filter(ImportedEvent.event_id <= _ROTINA_JOB_EVENT_MAX)
        .filter(ImportedEvent.job_name == job_name)
        .filter(ImportedEvent.time_created >= dt_start)
        .filter(ImportedEvent.time_created <= dt_end)
        .order_by(ImportedEvent.time_created.asc())
        .all()
    )

    total = len(events)
    if total == 0:
        return JSONResponse({"found": False})

    success = sum(1 for e in events if e.job_result == 0)
    warning = sum(1 for e in events if e.job_result == 1)
    failed  = sum(1 for e in events if e.job_result == 2)
    success_rate = round(success / total * 100, 1) if total else 0.0

    last3   = [e.job_result for e in events[-3:]][::-1]   # mais recente primeiro
    last_ev = events[-1]

    # ── Timeline por dia ──────────────────────────────────────────────────────
    buckets: "OrderedDict[str, dict]" = OrderedDict()
    for e in events:
        if not e.time_created:
            continue
        day = e.time_created.strftime("%Y-%m-%d")
        b   = buckets.setdefault(day, {"success": 0, "warning": 0, "failed": 0})
        if   e.job_result == 2: b["failed"]  += 1
        elif e.job_result == 1: b["warning"] += 1
        elif e.job_result == 0: b["success"] += 1
    timeline = [{"day": d, **v} for d, v in buckets.items()]

    # ── Ofensores das execuções com falha ─────────────────────────────────────
    cat_lookup = {c.id: (c.name, c.color, c.severity) for c in db.query(OffenderCategory).all()}
    off: dict = {}
    for e in events:
        if e.job_result == 2:
            key = e.offender_category_id or 0
            off[key] = off.get(key, 0) + 1
    offenders = []
    for key, count in off.items():
        if key and key in cat_lookup:
            nm, color, sev = cat_lookup[key]
        else:
            nm, color, sev = ("Não classificado", "#6c757d", "low")
        offenders.append({"category": nm, "color": color or "#6c757d",
                          "severity": sev, "count": count})
    offenders.sort(key=lambda x: -x["count"])

    # ── Histórico (mais recente primeiro, limitado) ───────────────────────────
    history = []
    for e in reversed(events):
        msg = (e.original_message or "").strip()
        history.append({
            "time":   e.time_created.strftime("%d/%m/%Y %H:%M") if e.time_created else "—",
            "result": e.job_result,
            "will_be_retried": bool(e.will_be_retried),
            "message": msg[:400],
        })
    history = history[:200]

    # ── Gestão persistente (JobAction) ────────────────────────────────────────
    job_norm = normalize_job_name(job_name)
    ja = (
        db.query(JobAction)
        .filter(JobAction.client_id == client_id,
                JobAction.job_name_normalized == job_norm)
        .first()
    )
    action = {
        "situation":             ja.situation if ja else "",
        "action_ongoing":        ja.action_ongoing if ja else "",
        "technical_observation": ja.technical_observation if ja else "",
        "responsible":           ja.responsible if ja else "",
        "ticket":                ja.ticket if ja else "",
        "updated_at":            ja.updated_at.strftime("%d/%m/%Y %H:%M") if ja and ja.updated_at else None,
        "updated_by":            ja.updated_by if ja else None,
    }

    client = db.query(Client).get(client_id)
    return JSONResponse({
        "found":               True,
        "job_name":            job_name,
        "job_name_normalized": job_norm,
        "job_type":            last_ev.event_category or "backup",
        "client_name":         client.name if client else "",
        "vbr_hostname":        last_ev.vbr_hostname,
        "kpis": {
            "total": total, "success": success, "warning": warning, "failed": failed,
            "success_rate":   success_rate,
            "last_execution": last_ev.time_created.strftime("%d/%m/%Y %H:%M") if last_ev.time_created else "—",
            "last_result":    last_ev.job_result,
            "last3":          last3,
        },
        "timeline":             timeline,
        "offenders":            offenders,
        "history":              history,
        "action":               action,
        "suggested_situation":  auto_situation(failed, total),
    })


@app.get("/api/rotina-vms", dependencies=[Depends(scope_guard)])
def api_rotina_vms(
    client_id:  int           = Query(...),
    job_name:   str           = Query(...),
    date_start: Optional[str] = Query(None),
    date_end:   Optional[str] = Query(None),
    db: Session = Depends(get_db),
):
    """Drill-down: VMs de uma rotina (via GUID de sessão) com RPO observado×esperado."""
    from sqlalchemy import text as sa_text

    if not date_start or not date_end:
        return JSONResponse(
            {"error": "Informe data inicial e data final.", "code": "dates_required"}, status_code=400
        )
    dt_start = _parse_date(date_start)
    dt_end   = _parse_date(date_end, end_of_day=True)
    if not dt_start or not dt_end:
        return JSONResponse({"error": "Datas inválidas.", "code": "dates_invalid"}, status_code=400)

    params = {
        "cid": client_id, "job": job_name,
        "s": dt_start, "e": dt_end,
        "defrpo": DEFAULT_RPO_MINUTES, "re": _SESSION_GUID_RE,
    }
    sql = f"""
    WITH job_sessions AS (
        SELECT DISTINCT (regexp_match(raw_event_xml, :re))[1] AS sguid
        FROM imported_events
        WHERE client_id = :cid
          AND event_id BETWEEN {_ROTINA_JOB_EVENT_MIN} AND {_ROTINA_JOB_EVENT_MAX}
          AND job_name = :job
          AND raw_event_xml IS NOT NULL
          AND time_created BETWEEN :s AND :e
    ),
    vm_events AS (
        SELECT COALESCE(e.job_name_normalized, e.job_name) AS vm_name,
               e.time_created, e.job_result
        FROM imported_events e
        WHERE e.client_id = :cid
          AND e.event_id = 150
          AND e.raw_event_xml IS NOT NULL
          AND e.time_created BETWEEN :s AND :e
          AND (regexp_match(e.raw_event_xml, :re))[1]
              IN (SELECT sguid FROM job_sessions WHERE sguid IS NOT NULL)
    ),
    success_evts AS (
        SELECT vm_name, time_created FROM vm_events WHERE job_result = 0
    ),
    ranked AS (
        SELECT vm_name, time_created,
               LAG(time_created) OVER (PARTITION BY vm_name ORDER BY time_created) AS prev_time
        FROM success_evts
    ),
    intervals AS (
        SELECT vm_name, EXTRACT(EPOCH FROM (time_created - prev_time)) / 3600.0 AS interval_h
        FROM ranked WHERE prev_time IS NOT NULL
    ),
    iagg AS (
        SELECT vm_name, AVG(interval_h) AS avg_rpo_h, MAX(interval_h) AS max_rpo_h
        FROM intervals GROUP BY vm_name
    ),
    agg AS (
        SELECT vm_name,
               COUNT(*)                                        AS total_tasks,
               SUM(CASE WHEN job_result = 0 THEN 1 ELSE 0 END) AS success_count,
               SUM(CASE WHEN job_result = 1 THEN 1 ELSE 0 END) AS warning_count,
               SUM(CASE WHEN job_result = 2 THEN 1 ELSE 0 END) AS failed_count,
               MAX(CASE WHEN job_result = 0 THEN time_created END) AS last_backup
        FROM vm_events GROUP BY vm_name
    )
    SELECT a.vm_name, a.total_tasks, a.success_count, a.warning_count, a.failed_count,
           a.last_backup, i.avg_rpo_h, i.max_rpo_h,
           vra.id AS assignment_id, vra.expected_minutes, vra.criticality,
           rp.name AS policy_name, rp.id AS policy_id,
           COALESCE(vra.expected_minutes, :defrpo) / 60.0 AS expected_h,
           CASE WHEN vra.id IS NOT NULL THEN 'configured' ELSE 'default' END AS rpo_source
    FROM agg a
    LEFT JOIN iagg i ON i.vm_name = a.vm_name
    LEFT JOIN vm_rpo_assignments vra
        ON LOWER(vra.vm_name) = LOWER(a.vm_name)
        AND vra.client_id = :cid AND vra.is_active = true
    LEFT JOIN rpo_policies rp ON rp.id = vra.rpo_policy_id
    ORDER BY a.failed_count DESC, a.vm_name
    """

    rows = db.execute(sa_text(sql), params).fetchall()

    summary = {"ok": 0, "attention": 0, "critical": 0, "no_history": 0, "no_backup": 0}
    vms = []
    for r in rows:
        avg_h = float(r.avg_rpo_h) if r.avg_rpo_h is not None else None
        max_h = float(r.max_rpo_h) if r.max_rpo_h is not None else None
        exp_h = float(r.expected_h) if r.expected_h is not None else DEFAULT_RPO_MINUTES / 60.0
        st    = _rpo_status(r.success_count, avg_h, exp_h)
        summary[st] = summary.get(st, 0) + 1
        since_h = None
        if r.last_backup and dt_end:
            since_h = round((dt_end - r.last_backup).total_seconds() / 3600, 1)
        vms.append({
            "vm_name":       r.vm_name,
            "total_tasks":   r.total_tasks,
            "success_count": r.success_count,
            "warning_count": r.warning_count,
            "failed_count":  r.failed_count,
            "last_backup":   r.last_backup.strftime("%d/%m/%Y %H:%M") if r.last_backup else None,
            "avg_rpo_h":     round(avg_h, 2) if avg_h is not None else None,
            "max_rpo_h":     round(max_h, 2) if max_h is not None else None,
            "expected_minutes": r.expected_minutes,
            "expected_h":    round(exp_h, 2),
            "rpo_source":    r.rpo_source,
            "policy_name":   r.policy_name,
            "criticality":   r.criticality,
            "deviation_pct": _calc_deviation_pct(avg_h, exp_h),
            "since_h":       since_h,
            "status":        st,
        })

    return JSONResponse({"vms": vms, "total": len(vms), "summary": summary})


@app.post("/api/rotina-action", dependencies=[Depends(scope_guard)])
async def api_rotina_action(request: Request, db: Session = Depends(get_db)):
    """Upsert da gestão persistente (JobAction) de uma rotina — mesma chave do relatório."""
    from app.parser.common import normalize_job_name

    data      = await request.json()
    client_id = data.get("client_id")
    job_name  = data.get("job_name")
    if not client_id or not job_name:
        return JSONResponse({"error": "client_id e job_name são obrigatórios."}, status_code=400)

    job_norm = normalize_job_name(job_name)
    ja = (
        db.query(JobAction)
        .filter(JobAction.client_id == client_id,
                JobAction.job_name_normalized == job_norm)
        .first()
    )
    if ja is None:
        ja = JobAction(client_id=client_id, job_name_normalized=job_norm)
        db.add(ja)

    ja.situation             = (data.get("situation") or "Indefinido")
    ja.action_ongoing        = (data.get("action_ongoing") or "")
    ja.technical_observation = (data.get("technical_observation") or "")
    ja.responsible           = (data.get("responsible") or "")
    ja.ticket                = (data.get("ticket") or "")
    ja.updated_at            = datetime.utcnow()
    ja.updated_by            = request.session.get("username")
    db.commit()

    return JSONResponse({
        "ok": True,
        "updated_at": ja.updated_at.strftime("%d/%m/%Y %H:%M"),
        "updated_by": ja.updated_by,
    })


# ── API auxiliar: relatórios por cliente (para dropdowns de filtro) ─────────────

@app.get("/api/reports-by-client", dependencies=[Depends(scope_guard)])
def api_reports_by_client(
    client_id: Optional[int] = Query(None),
    db: Session = Depends(get_db),
):
    """Retorna relatórios de um cliente (ou todos) para popular dropdowns de filtro.
    Endpoint leve — sem consulta de eventos.
    """
    rows = _all_reports_list(db, client_id=client_id)
    return JSONResponse({"reports": rows})


# ── Gestão de clientes ──────────────────────────────────────────────────────────

@app.get("/clients", response_class=HTMLResponse)
def clients_page(request: Request, db: Session = Depends(get_db)):
    clients = db.query(Client).order_by(Client.name).all()
    stats = {}
    for c in clients:
        stats[c.id] = {
            "events":  db.query(ImportedEvent).filter(ImportedEvent.client_id == c.id).count(),
            "uploads": db.query(UploadSession).filter(UploadSession.client_id == c.id).count(),
        }
    return templates.TemplateResponse(
        "clients.html", {"request": request, "clients": clients, "stats": stats}
    )


@app.post("/clients/new", dependencies=[Depends(require_admin)])
async def create_client(
    request: Request,
    name: str = Form(...),
    description: str = Form(""),
    db: Session = Depends(get_db),
):
    name = name.strip()
    if not name:
        clients = db.query(Client).order_by(Client.name).all()
        return templates.TemplateResponse(
            "clients.html",
            {"request": request, "clients": clients, "stats": {}, "error": "O nome do cliente é obrigatório."},
            status_code=400,
        )
    if db.query(Client).filter(Client.name == name).first():
        clients = db.query(Client).order_by(Client.name).all()
        return templates.TemplateResponse(
            "clients.html",
            {"request": request, "clients": clients, "stats": {}, "error": f"Cliente '{name}' já existe."},
            status_code=400,
        )
    db.add(Client(name=name, description=description.strip() or None))
    db.commit()
    return RedirectResponse(url="/clients", status_code=303)


@app.post("/clients/{client_id}/edit", dependencies=[Depends(require_admin)])
async def edit_client(
    client_id: int, request: Request,
    name: str = Form(...), description: str = Form(""),
    db: Session = Depends(get_db),
):
    client = db.query(Client).get(client_id)
    if not client:
        raise HTTPException(status_code=404)
    name = name.strip()
    if not name:
        raise HTTPException(status_code=400, detail="Nome inválido")
    client.name        = name
    client.description = description.strip() or None
    db.commit()
    return RedirectResponse(url="/clients", status_code=303)


@app.get("/api/admin/users/{user_id}/clients", dependencies=[Depends(require_admin)])
def api_user_clients_get(user_id: int, db: Session = Depends(get_db)):
    """Lista clientes com flag 'assigned' para o escopo de um usuário (admin)."""
    assigned = {r[0] for r in db.query(UserClient.client_id).filter(UserClient.user_id == user_id).all()}
    clients = db.query(Client).filter(Client.is_active.is_(True)).order_by(Client.name).all()
    return {
        "user_id": user_id,
        "clients": [{"id": c.id, "name": c.name, "assigned": c.id in assigned} for c in clients],
    }


@app.post("/api/admin/users/{user_id}/clients", dependencies=[Depends(require_admin)])
async def api_user_clients_set(user_id: int, request: Request, db: Session = Depends(get_db)):
    """Define (substitui) o conjunto de clientes do escopo de um usuário (admin)."""
    user = db.query(User).get(user_id)
    if not user:
        raise HTTPException(status_code=404, detail="Usuário não encontrado.")
    data = await request.json()
    valid = {r[0] for r in db.query(Client.id).all()}
    try:
        ids = [int(i) for i in (data.get("client_ids") or []) if int(i) in valid]
    except (TypeError, ValueError):
        raise HTTPException(status_code=400, detail="client_ids inválidos.")
    db.query(UserClient).filter(UserClient.user_id == user_id).delete()
    actor = request.session.get("username")
    for cid in ids:
        db.add(UserClient(user_id=user_id, client_id=cid, created_by=actor))
    db.commit()
    return {"ok": True, "assigned": ids}


@app.post("/clients/{client_id}/toggle", dependencies=[Depends(require_admin)])
def toggle_client(client_id: int, db: Session = Depends(get_db)):
    client = db.query(Client).get(client_id)
    if not client:
        raise HTTPException(status_code=404)
    client.is_active = not client.is_active
    db.commit()
    return RedirectResponse(url="/clients", status_code=303)


# ── Helpers ─────────────────────────────────────────────────────────────────────

def _get_report_or_404(db: Session, report_id: int) -> ReportSession:
    report = db.query(ReportSession).get(report_id)
    if not report:
        raise HTTPException(status_code=404, detail="Relatório não encontrado")
    return report
